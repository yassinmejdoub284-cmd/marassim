"""
gui.py — Interface graphique Tkinter — Marassim (UI/UX redesign v3)
"""
import os
import sqlite3
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import date, timedelta

import database
import rules
import contract_generator
import excel_export
import access_control
import ui_kit as UI
from ui_kit import T as TOK, toast
from calendar_tab import YearCalendarTab
from rules_tab import RulesTab

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "contrats_generes")
EXPORT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports_excel")

MOIS_FR = ["janvier","février","mars","avril","mai","juin","juillet",
           "août","septembre","octobre","novembre","décembre"]

# ═══════════════════════════════════════════════════════════════════
#  PALETTE & THÈME
# ═══════════════════════════════════════════════════════════════════
C = {
    # ── Mappé sur ui_kit.T : les clés restent identiques pour que les
    #    ~5900 lignes existantes continuent de fonctionner telles quelles.
    "bg":         TOK["bg"],
    "surface":    TOK["surface"],
    "sidebar":    TOK["nav"],
    "sidebar_h":  TOK["nav_h"],
    "sidebar_sel":TOK["nav_sel"],
    "primary":    TOK["primary"],
    "primary_d":  TOK["primary_d"],
    "accent":     TOK["accent"],
    "success":    TOK["success"],
    "success_d":  TOK["success_d"],
    "danger":     TOK["danger"],
    "danger_d":   TOK["danger_d"],
    "warning":    TOK["warning"],
    "text":       TOK["text"],
    "text_m":     TOK["text_m"],
    "border":     TOK["border"],
    "stripe":     TOK["surface_2"],
    "card_hdr":   TOK["surface_2"],
    "tag_solde":  TOK["success_l"],
    "tag_part":   TOK["warning_l"],
    "tag_zero":   TOK["danger_l"],
    "tag_rec":    TOK["success_l"],
    "tag_chg":    TOK["danger_l"],
    "tag_tot":    TOK["primary_l"],
    "tag_sld":    TOK["warning_l"],
    # Nouveaux jetons disponibles pour les écrans redessinés
    "surface_2":  TOK["surface_2"],
    "surface_3":  TOK["surface_3"],
    "border_s":   TOK["border_s"],
    "text_d":     TOK["text_d"],
    "primary_l":  TOK["primary_l"],
    "shadow":     TOK["shadow"],
    "violon":     TOK["violon"],
    "violon_l":   TOK["violon_l"],
}

# Police résolue au lancement par ui_kit.install_theme() ; les valeurs
# ci-dessous servent de repli si un module est importé hors application.
FNT = {
    "default":  ("Segoe UI", 10),
    "sm":       ("Segoe UI", 9),
    "xs":       ("Segoe UI", 8),
    "md":       ("Segoe UI", 11),
    "bold":     ("Segoe UI", 10, "bold"),
    "sm_bold":  ("Segoe UI", 9, "bold"),
    "title":    ("Segoe UI", 16, "bold"),
    "section":  ("Segoe UI", 11, "bold"),
    "id":       ("Segoe UI", 19, "bold"),
    "sidebar":  ("Segoe UI", 10),
    "sidebar_b":("Segoe UI", 10, "bold"),
    "mono":     ("Consolas", 10),
    "kpi":      ("Segoe UI", 22, "bold"),
}


def _sync_fonts():
    """Aligne FNT sur la police réellement disponible (résolue par ui_kit)."""
    f = UI.F
    if not f:
        return
    FNT.update({
        "default": f["default"], "sm": f["sm"], "xs": f["xs"], "md": f["md"],
        "bold": f["bold"], "sm_bold": f["sm_bold"], "title": f["title"],
        "section": f["section"], "id": f["id"], "sidebar": f["nav"],
        "sidebar_b": f["nav_b"], "mono": f["mono"], "kpi": f["kpi"],
    })


# ═══════════════════════════════════════════════════════════════════
#  STYLE TTK
# ═══════════════════════════════════════════════════════════════════
def _apply_style():
    """Thème ttk complet — délégué à ui_kit (composants + design tokens)."""
    style = UI.install_theme()
    _sync_fonts()
    return style


# ═══════════════════════════════════════════════════════════════════
#  HELPERS VISUELS
# ═══════════════════════════════════════════════════════════════════
def _divider(parent, padx=0, pady=(6, 6)):
    tk.Frame(parent, bg=C["border"], height=1).pack(fill="x", padx=padx, pady=pady)


def _card(parent, title="", **kw):
    """Carte blanche : bordure fine + ombre douce (style ui_kit).

    Compatible avec l'ancien appel `_card(parent, "Titre")` : l'objet
    retourné accepte pack()/grid() et sert de conteneur direct.
    """
    kw.pop("padding", None)
    card = UI.Card(parent, title=title, padding=14, **kw)
    # Compat : permet de continuer à empiler des widgets dans le retour
    card.pack_propagate(True)
    return card


def _scrollable(parent):
    """Zone scrollable verticale. Retourne (outer, inner)."""
    return UI.scrollable(parent, bg=C["bg"])


def _field(parent, row, col, label, width=18, colspan=1, rowspan=1, bg=None):
    """Bloc « libellé + champ » posé en grille."""
    bg = bg or C["surface"]
    fr = tk.Frame(parent, bg=bg)
    fr.grid(row=row, column=col, columnspan=colspan, rowspan=rowspan,
            sticky="nsew", padx=(4, 10), pady=5)
    fr.columnconfigure(0, weight=1)
    tk.Label(fr, text=label, bg=bg, fg=C["text_m"], font=FNT["sm_bold"],
             anchor="w").grid(row=0, column=0, sticky="w", pady=(0, 3))
    ent = RoundedEntry(fr, bg_canvas=bg)
    ent.grid(row=1, column=0, sticky="ew")
    return ent


def _badge(parent, text, color, bg=C["surface"]):
    """Pastille colorée arrondie."""
    tone = {C["success"]: "success", C["danger"]: "danger",
            C["warning"]: "warning", C["primary"]: "primary"}.get(color)
    if tone:
        return UI.Chip(parent, text, tone=tone, surface=bg)
    fr = tk.Frame(parent, bg=color, padx=9, pady=3)
    tk.Label(fr, text=text, bg=color, fg="white", font=FNT["sm_bold"]).pack()
    return fr


def _section_title(parent, text, icon="", bg=None, accent=None):
    """Titre de section avec barre d'accent."""
    bg = bg or C["surface"]
    f = tk.Frame(parent, bg=bg)
    tk.Frame(f, bg=accent or C["primary"], width=3, height=17).pack(
        side="left", padx=(0, 9))
    tk.Label(f, text=f"{icon}  {text}" if icon else text, bg=bg,
             fg=C["text"], font=FNT["section"]).pack(side="left")
    return f


class RoundedButton(tk.Canvas):
    """Bouton arrondi : ombre douce, survol animé, état désactivé.

    Signature identique à la version précédente — les ~150 appels
    existants dans ce fichier fonctionnent sans modification.
    """

    RADIUS = 10

    # style historique -> (couleur, texte, bordure)
    _STYLES = {
        "primary":   (C["primary"],  "white",      None),
        "secondary": (C["surface"],  C["text"],    C["border_s"]),
        "success":   (C["success"],  "white",      None),
        "danger":    (C["danger"],   "white",      None),
        "warning":   (C["warning"],  "white",      None),
        "ghost":     (None,          C["text_m"],  None),
        "soft":      (C["primary_l"], C["primary_d"], None),
        "sm":        (C["primary"],  "white",      None),
        "sm_danger": (C["danger"],   "white",      None),
    }

    def __init__(self, parent, text, command=None, style="primary",
                 font=None, padx=17, pady=8, **kwargs):
        self._surface = kwargs.pop("bg", None) or UI._parent_bg(parent)
        super().__init__(parent, bg=self._surface, highlightthickness=0, bd=0,
                         cursor="hand2", takefocus=0, **kwargs)

        base, fg, bd = self._STYLES.get(style, self._STYLES["primary"])
        self._style   = style
        self._base    = base or self._surface
        self._flat    = base is None            # variante « ghost »
        self._fg      = fg
        self._bd      = bd
        self._command = command
        self._text    = text
        self._enabled = True
        self._hover   = False
        self._font    = font or (FNT["sm_bold"] if "sm" in style else FNT["bold"])
        self._padx, self._pady = padx, pady
        self._cur     = self._base

        w, h = self._measure()
        if not kwargs.get("width"):
            self.configure(width=w)
        if not kwargs.get("height"):
            self.configure(height=h)

        self.bind("<Configure>",       lambda e: self._draw())
        self.bind("<Enter>",           self._on_enter)
        self.bind("<Leave>",           self._on_leave)
        self.bind("<ButtonPress-1>",   self._on_press)
        self.bind("<ButtonRelease-1>", self._on_release)
        self._draw()

    # ── géométrie ────────────────────────────────────────────────
    def _measure(self):
        import tkinter.font as tkfont
        f = tkfont.Font(font=self._font)
        return (f.measure(self._text) + self._padx * 2,
                f.metrics("linespace") + self._pady * 2 + 3)

    def _hover_color(self):
        if self._style in ("secondary", "ghost"):
            return C["surface_2"]
        if self._style == "soft":
            return TOK["primary_b"]
        return UI.darken(self._base, .12)

    # ── rendu ────────────────────────────────────────────────────
    def _draw(self, fill=None):
        self.delete("all")
        w, h = self.winfo_width(), self.winfo_height()
        if w < 4 or h < 4:
            return
        solid = self._style in ("primary", "success", "danger", "warning",
                                "sm", "sm_danger")
        if not self._enabled:
            body, fg, bd = C["surface_2"], C["text_d"], C["border"]
        else:
            body = fill or self._cur
            fg, bd = self._fg, self._bd
            if self._flat and not self._hover:
                body = self._surface

        # ombre portée
        if solid and self._enabled:
            UI.round_rect(self, 1, 3, w - 1, h, self.RADIUS,
                          fill=UI.lerp(self._surface, self._base, .24), outline="")
        UI.round_rect(self, 0, 0, w - 1, h - 3, self.RADIUS,
                      fill=body, outline=bd or body)
        self.create_text(w / 2, (h - 3) / 2, text=self._text, fill=fg,
                         font=self._font, anchor="center")

    def _to(self, target):
        start = self._cur

        def apply(t):
            self._cur = UI.lerp(start, target, t)
            self._draw()
        UI.anim(self, apply, steps=6, delay=12)

    # ── événements ───────────────────────────────────────────────
    def _on_enter(self, e=None):
        if not self._enabled:
            return
        self._hover = True
        self._to(self._hover_color())

    def _on_leave(self, e=None):
        self._hover = False
        if self._enabled:
            self._to(self._base)

    def _on_press(self, e=None):
        if self._enabled:
            self._cur = UI.darken(self._hover_color(), .10)
            self._draw()

    def _on_release(self, e=None):
        if not self._enabled:
            return
        self._to(self._hover_color() if self._hover else self._base)
        if self._command:
            self._command()

    def invoke(self):
        if self._enabled and self._command:
            self._command()

    # ── API ──────────────────────────────────────────────────────
    def configure(self, **kw):
        redraw = False
        if "text" in kw:
            self._text = kw.pop("text"); redraw = True
        if "command" in kw:
            self._command = kw.pop("command")
        if "fg" in kw:
            self._fg = kw.pop("fg"); redraw = True
        if "state" in kw:
            self._enabled = kw.pop("state") != "disabled"
            try:
                super().configure(cursor="hand2" if self._enabled else "")
            except tk.TclError:
                pass
            redraw = True
        if "style" in kw:
            st = kw.pop("style")
            base, fg, bd = self._STYLES.get(st, self._STYLES["primary"])
            self._style, self._base = st, base or self._surface
            self._flat, self._fg, self._bd = base is None, fg, bd
            self._cur = self._base
            redraw = True
        if kw:
            super().configure(**kw)
        if redraw:
            self._draw()
    config = configure


def _rbtn(parent, text, command=None, style="primary", font=None,
          padx=16, pady=7, **kw):
    """Shortcut: create a RoundedButton, size it, pack it, return it."""
    import tkinter.font as tkfont
    f = tkfont.Font(font=font or (FNT["bold"] if "sm" not in style else FNT["sm"]))
    tw = f.measure(text)
    th = f.metrics("linespace")
    w = tw + padx * 2 + 2
    h = th + pady * 2 + 2
    btn = RoundedButton(parent, text=text, command=command, style=style,
                        font=font, padx=padx, pady=pady,
                        width=w, height=h, **kw)
    btn.pack(side="left", padx=(0, 8), pady=4)
    return btn


# ═══════════════════════════════════════════════════════════════════
#  ROUNDED ENTRY (canvas frame wrapper for smooth look)
# ═══════════════════════════════════════════════════════════════════
class RoundedEntry(tk.Frame):
    """Champ de saisie : anneau de focus animé, curseur épais, état d'erreur.

    Signature et méthodes identiques à la version précédente.
    """

    def __init__(self, parent, bg_canvas=None, width=None, justify="left",
                 font=None, state="normal", textvariable=None, **kw):
        self._surface = bg_canvas or UI._parent_bg(parent)
        super().__init__(parent, bg=C["border"], padx=1, pady=1)
        self._inner = tk.Frame(self, bg=C["surface"])
        self._inner.pack(fill="both", expand=True)
        self._entry = tk.Entry(
            self._inner, relief="flat", bg=C["surface"], fg=C["text"],
            insertbackground=C["primary"], insertwidth=2,
            font=font or FNT["default"], justify=justify,
            highlightthickness=0, bd=0,
            selectbackground=TOK["primary_b"], selectforeground=C["text"])
        if textvariable:
            self._entry.configure(textvariable=textvariable)
        if width:
            self._entry.configure(width=width)
        self._entry.pack(fill="both", expand=True, padx=8, pady=6)
        self._error = False
        self._entry.bind("<FocusIn>",  self._focus_in,  add="+")
        self._entry.bind("<FocusOut>", self._focus_out, add="+")
        self._apply_state(state)

    # ── anneau de focus animé ────────────────────────────────────
    def _ring_to(self, target):
        start = self.cget("bg")
        UI.anim(self, lambda t: tk.Frame.configure(
            self, bg=UI.lerp(start, target, t)), steps=5, delay=11)

    def _focus_in(self, _=None):
        self._ring_to(C["danger"] if self._error else C["primary"])

    def _focus_out(self, _=None):
        self._ring_to(C["danger"] if self._error else C["border"])

    def set_error(self, on=True):
        """Passe le champ en rouge (validation échouée)."""
        self._error = bool(on)
        tk.Frame.configure(self, bg=C["danger"] if on else C["border"])
        bg = TOK["danger_l"] if on else C["surface"]
        self._inner.configure(bg=bg)
        self._entry.configure(bg=bg)

    def _apply_state(self, state):
        if state == "readonly":
            self._entry.configure(state="readonly",
                                  readonlybackground=C["stripe"])
            self._inner.configure(bg=C["stripe"])
        elif state == "disabled":
            self._entry.configure(state="disabled",
                                  disabledbackground=C["stripe"],
                                  disabledforeground=C["text_d"])
            self._inner.configure(bg=C["stripe"])
        else:
            self._entry.configure(state="normal", bg=C["surface"])
            self._inner.configure(bg=C["surface"])

    # ── proxy Entry ──────────────────────────────────────────────
    def get(self):            return self._entry.get()
    def delete(self, *a):     self._entry.delete(*a)
    def insert(self, *a):     self._entry.insert(*a)
    def bind(self, *a, **kw): return self._entry.bind(*a, **kw)
    def focus_set(self):      self._entry.focus_set()

    @property
    def entry(self):          return self._entry

    def configure(self, **kw):
        entry_keys = {"font", "justify", "textvariable", "show",
                      "disabledforeground", "readonlybackground"}
        if "state" in kw:
            self._apply_state(kw.pop("state"))
        ekw = {k: kw.pop(k) for k in list(kw) if k in entry_keys}
        if ekw:
            self._entry.configure(**ekw)
        if kw:
            super().configure(**kw)
    config = configure


# ═══════════════════════════════════════════════════════════════════
#  APPLICATION PRINCIPALE — sidebar + content panes
# ═══════════════════════════════════════════════════════════════════
# Libellé affiché dans la sidebar quand le nom du module est trop long.
# La clé reste le nom du module — c'est lui qui porte les permissions.
NAV_SHORT = {
    "Règles de réservation": "Règles",
    "Modifier réservation":  "Modifier",
    "Ajouter un acompte":    "Acomptes",
    "Journal Caisse Omar":   "Caisse Omar",
    "Journal de Caisse":     "Journal caisse",
    "Centre de Réception":   "Centre réception",
}

NAV_ITEMS = [
    ("Nouvelle réservation", "＋"),
    ("Réservations",         "☰"),
    ("Calendrier",           "📅"),
    ("Ajouter un acompte",   "💵"),
    ("Charges",              "📦"),
    ("Charges Omar",         "📦"),
    ("Journal de Caisse",    "📒"),
    ("Journal Caisse Omar",  "📒"),
    ("Modifier réservation", "✏"),
    ("Export Excel",         "📊"),
    ("Ajouter employé",      "👤"),
    ("Liste employés",       "👥"),
    ("Pointage",             "⏰"),
    ("Centre de Réception",  "📥"),
    ("Règles de réservation", "⚖"),
]



class _NavItem(tk.Frame):
    """Entrée de navigation : pastille active, survol animé, indicateur."""

    def __init__(self, parent, icon, label, command=None):
        super().__init__(parent, bg=C["sidebar"], highlightthickness=0, bd=0)
        self._cmd = command
        self._active = False
        self._bg = C["sidebar"]

        self._bar = tk.Frame(self, bg=C["sidebar"], width=3)
        self._bar.pack(side="left", fill="y")
        self._row = tk.Frame(self, bg=C["sidebar"])
        self._row.pack(side="left", fill="both", expand=True)
        self._icon = tk.Label(self._row, text=icon, bg=C["sidebar"],
                              fg=TOK["nav_txt"], font=(UI.FAMILY, 11),
                              width=2, anchor="center")
        self._icon.pack(side="left", padx=(10, 0), pady=9)
        self._lbl = tk.Label(self._row, text=label, bg=C["sidebar"],
                             fg=TOK["nav_txt"], font=FNT["sidebar"], anchor="w")
        self._lbl.pack(side="left", padx=(9, 10), pady=9, fill="x", expand=True)

        for w in (self, self._row, self._icon, self._lbl):
            w.configure(cursor="hand2")
            w.bind("<Button-1>", self._click)
            w.bind("<Enter>", self._enter)
            w.bind("<Leave>", self._leave)

    def _paint(self, bg, fg, bar):
        for w in (self, self._row, self._icon, self._lbl):
            w.configure(bg=bg)
        self._icon.configure(fg=fg)
        self._lbl.configure(fg=fg)
        self._bar.configure(bg=bar)

    def _animate_bg(self, target):
        start = self._bg

        def apply(t):
            self._bg = UI.lerp(start, target, t)
            for w in (self, self._row, self._icon, self._lbl):
                w.configure(bg=self._bg)
        UI.anim(self, apply, steps=5, delay=11)

    def _click(self, _=None):
        if self._cmd:
            self._cmd()

    def _enter(self, _=None):
        if not self._active:
            self._animate_bg(C["sidebar_h"])
            self._icon.configure(fg="white")
            self._lbl.configure(fg="white")

    def _leave(self, _=None):
        if not self._active:
            self._animate_bg(C["sidebar"])
            self._icon.configure(fg=TOK["nav_txt"])
            self._lbl.configure(fg=TOK["nav_txt"])

    def set_active(self, on):
        self._active = bool(on)
        if on:
            self._bg = C["sidebar_sel"]
            self._paint(C["sidebar_sel"], "white", "white")
            self._lbl.configure(font=FNT["sidebar_b"])
        else:
            self._bg = C["sidebar"]
            self._paint(C["sidebar"], TOK["nav_txt"], C["sidebar"])
            self._lbl.configure(font=FNT["sidebar"])

    def bind_all_wheel(self, canvas):
        def wheel(e):
            canvas.yview_scroll(int(-1 * (e.delta / 120)), "units")
        for w in (self, self._row, self._icon, self._lbl):
            w.bind("<MouseWheel>", wheel)

    def configure(self, **kw):
        kw.pop("font", None); kw.pop("fg", None)
        super().configure(**kw)


class MarassimApp(tk.Tk):
    def __init__(self, user):
        super().__init__()
        self.title("Marassim — Gestion des réservations")
        # Screen-relative size: 90% wide, 88% tall, centered
        self.update_idletasks()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        w = max(1000, int(sw * 0.90))
        h = max(700,  int(sh * 0.88))
        x = (sw - w) // 2
        y = (sh - h) // 2
        self.geometry(f"{w}x{h}+{x}+{y}")
        self.minsize(900, 700)
        self.configure(bg=C["sidebar"])
        _apply_style()
        database.init_db()
        
        # Stocker l'utilisateur connecté
        self.current_user = user
        self.user_modules = access_control.get_user_modules(user["id"])

        self._frames = {}
        self._nav_btns = {}
        self._current = None
        self._sidebar_visible = True

        UI.init_toasts(self)
        self._build_shell()
        self._build_tabs()
        self._select_first_available_tab()

    # ── Shell : sidebar + content ─────────────────────────────────
    def _build_shell(self):
        # ── Sidebar ───────────────────────────────────────────────
        self.sidebar = tk.Frame(self, bg=C["sidebar"], width=258)
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)

        # Logo / brand
        brand = tk.Frame(self.sidebar, bg=C["sidebar"], pady=18, padx=18)
        brand.pack(fill="x")
        logo_row = tk.Frame(brand, bg=C["sidebar"])
        logo_row.pack(fill="x")
        mark = tk.Frame(logo_row, bg=C["primary"], width=38, height=38)
        mark.pack(side="left"); mark.pack_propagate(False)
        tk.Label(mark, text="M", bg=C["primary"], fg="white",
                 font=(UI.FAMILY, 17, "bold")).place(relx=.5, rely=.5, anchor="center")
        txt = tk.Frame(logo_row, bg=C["sidebar"])
        txt.pack(side="left", padx=(11, 0))
        tk.Label(txt, text="MARASSIM", bg=C["sidebar"], fg="white",
                 font=(UI.FAMILY, 13, "bold"), anchor="w").pack(anchor="w")
        tk.Label(txt, text="Complexe de réception", bg=C["sidebar"],
                 fg=TOK["nav_txt"], font=(UI.FAMILY, 8), anchor="w").pack(anchor="w")

        # User info
        user_info = tk.Frame(self.sidebar, bg=C["nav_h"] if "nav_h" in C else C["sidebar_h"],
                             pady=9, padx=11)
        user_info.pack(fill="x", padx=14, pady=(2, 10))
        user_name = f"{self.current_user.get('prenom', '')} {self.current_user.get('nom', '')}".strip()
        initials = "".join(w[0].upper() for w in user_name.split()[:2]) or "?"
        av = tk.Frame(user_info, bg=C["primary"], width=30, height=30)
        av.pack(side="left"); av.pack_propagate(False)
        tk.Label(av, text=initials, bg=C["primary"], fg="white",
                 font=(UI.FAMILY, 10, "bold")).place(relx=.5, rely=.5, anchor="center")
        ub = tk.Frame(user_info, bg=C["sidebar_h"])
        ub.pack(side="left", padx=(9, 0), fill="x", expand=True)
        shown = user_name or "Utilisateur"
        if len(shown) > 19:
            shown = shown[:18] + "…"
        tk.Label(ub, text=shown, bg=C["sidebar_h"], fg="white",
                 font=(UI.FAMILY, 9, "bold"), anchor="w").pack(anchor="w")
        tk.Label(ub, text=self.current_user.get('role', '').capitalize(),
                 bg=C["sidebar_h"], fg=TOK["nav_txt"],
                 font=(UI.FAMILY, 8), anchor="w").pack(anchor="w")

        # Nav buttons - wrapped in scrollable frame
        nav_container = tk.Frame(self.sidebar, bg=C["sidebar"])
        nav_container.pack(fill="both", expand=True)
        
        # Create canvas for scrolling
        self.nav_canvas = tk.Canvas(nav_container, bg=C["sidebar"], 
                                     highlightthickness=0, borderwidth=0)
        self.nav_scrollbar = tk.Scrollbar(nav_container, orient="vertical",
                                          command=self.nav_canvas.yview,
                                          bg=C["sidebar_h"], troughcolor=C["sidebar"],
                                          activebackground=C["primary"],
                                          highlightthickness=0, bd=0,
                                          relief="flat", width=4)
        self.nav_canvas.configure(yscrollcommand=self.nav_scrollbar.set)
        
        # Pack scrollbar and canvas
        self.nav_scrollbar.pack(side="right", fill="y")
        self.nav_canvas.pack(side="left", fill="both", expand=True)
        
        # Create frame inside canvas for nav buttons
        self.nav_frame = tk.Frame(self.nav_canvas, bg=C["sidebar"])
        self.nav_canvas_window = self.nav_canvas.create_window(
            (0, 0), window=self.nav_frame, anchor="nw")
        
        # Update canvas window width when canvas resizes
        def _on_canvas_resize(event):
            self.nav_canvas.itemconfig(self.nav_canvas_window, width=event.width)
        self.nav_canvas.bind("<Configure>", _on_canvas_resize)
        
        # Bind mousewheel to canvas and all child widgets
        def _on_nav_mousewheel(event):
            self.nav_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        def _bind_nav_scroll(widget):
            widget.bind("<MouseWheel>", _on_nav_mousewheel)
            for child in widget.winfo_children():
                _bind_nav_scroll(child)

        self.nav_canvas.bind("<MouseWheel>", _on_nav_mousewheel)
        self.nav_frame.bind("<MouseWheel>", _on_nav_mousewheel)
        self._bind_nav_scroll = _bind_nav_scroll
        
        # Configure scroll region
        def _configure_scroll_region(event=None):
            self.nav_canvas.configure(scrollregion=self.nav_canvas.bbox("all"))
        
        self.nav_frame.bind("<Configure>", _configure_scroll_region)

        # Filtrer les modules selon les accès de l'utilisateur
        self._visible_nav_items = []
        for i, (label, icon) in enumerate(NAV_ITEMS):
            # Vérifier si l'utilisateur a accès à ce module
            if label not in self.user_modules:
                continue  # Ne pas afficher ce module
            
            self._visible_nav_items.append((i, label, icon))
            
            btn = _NavItem(self.nav_frame, icon, NAV_SHORT.get(label, label),
                           command=lambda idx=i: self._select_tab(idx))
            btn.pack(fill="x", padx=10, pady=1)
            btn.bind_all_wheel(self.nav_canvas)
            self._nav_btns[i] = btn

        # Force scroll region update after all buttons are packed
        self.nav_frame.update_idletasks()
        self.nav_canvas.configure(scrollregion=self.nav_canvas.bbox("all"))
        
        # Bottom fixed buttons (outside scroll area)
        bottom_section = tk.Frame(self.sidebar, bg=C["sidebar"])
        bottom_section.pack(side="bottom", fill="x")
        
        # Footer date
        tk.Label(bottom_section,
                 text=date.today().strftime("%A, %d %B %Y").capitalize(),
                 bg=C["sidebar"], fg="#64748B", font=("Segoe UI", 8)
                 ).pack(side="bottom", pady=10)
        tk.Frame(bottom_section, bg=C["sidebar_h"], height=1).pack(
            side="bottom", fill="x", padx=12)
        
        # Déconnexion
        logout_btn = tk.Button(
            bottom_section,
            text="  🚪  Déconnexion",
            anchor="w",
            bg=C["sidebar"], fg="#F87171",
            activebackground=C["sidebar_h"], activeforeground="#FCA5A5",
            relief="flat", bd=0, cursor="hand2", highlightthickness=0,
            font=FNT["sidebar"],
            padx=18, pady=11,
            command=self._logout,
        )
        logout_btn.pack(side="bottom", fill="x")
        logout_btn.bind("<Enter>", lambda e: logout_btn.configure(
            bg=C["sidebar_h"]))
        logout_btn.bind("<Leave>", lambda e: logout_btn.configure(
            bg=C["sidebar"]))
        
        tk.Frame(bottom_section, bg=C["sidebar_h"], height=1).pack(
            side="bottom", fill="x", padx=12)
        
        # Gestion des accès (admin only)
        if self.current_user.get('role') == 'admin':
            admin_btn = tk.Button(
                bottom_section,
                text="  ⚙  Gestion des accès",
                anchor="w",
                bg=C["sidebar"], fg=TOK["nav_txt"],
                activebackground=C["sidebar_h"], activeforeground="white",
                relief="flat", bd=0, cursor="hand2", highlightthickness=0,
                font=FNT["sidebar"],
                padx=18, pady=11,
                command=self._open_access_management,
            )
            admin_btn.pack(side="bottom", fill="x")
            admin_btn.bind("<Enter>", lambda e: admin_btn.configure(
                bg=C["sidebar_h"], fg="white"))
            admin_btn.bind("<Leave>", lambda e: admin_btn.configure(
                bg=C["sidebar"], fg="#CBD5E1"))

        # ── Toggle button strip (always visible, between sidebar & content) ──
        self.toggle_strip = tk.Frame(self, bg=C["sidebar"], width=14)
        self.toggle_strip.pack(side="left", fill="y")
        self.toggle_strip.pack_propagate(False)
        self._toggle_btn = tk.Button(
            self.toggle_strip, text="‹",
            bg=C["sidebar_h"], fg="white",
            relief="flat", bd=0, cursor="hand2",
            font=("Segoe UI", 12, "bold"),
            command=self._toggle_sidebar,
            activebackground=C["sidebar_sel"],
            activeforeground="white",
        )
        self._toggle_btn.place(relx=0.5, rely=0.5, anchor="center")

        # ── Content area ──────────────────────────────────────────
        self.content = tk.Frame(self, bg=C["bg"])
        self.content.pack(side="left", fill="both", expand=True)

    def _build_tabs(self):
        self.tab_new              = NewReservationTab(self.content, self)
        self.tab_list             = ListTab(self.content, self)
        self.tab_calendar         = YearCalendarTab(self.content, self)
        self.tab_acompte          = AcompteTab(self.content, self)
        self.tab_charges          = ChargesTab(self.content, self)
        self.tab_charges_omar     = ChargesOmarTab(self.content, self)
        self.tab_journal          = JournalCaisseTab(self.content, self)
        self.tab_journal_omar     = JournalCaisseOmarTab(self.content, self)
        self.tab_modifier         = ModifierReservationTab(self.content, self)
        self.tab_export           = ExportTab(self.content, self)
        self.tab_employee         = EmployeeTab(self.content, self)
        self.tab_employee_list    = EmployeeListTab(self.content, self)
        self.tab_pointage         = PointageTab(self.content, self)
        self.tab_centre_reception = CentreReceptionTab(self.content, self)
        self.tab_rules            = RulesTab(self.content, self)

        self._frames = {
            0:  self.tab_new,
            1:  self.tab_list,
            2:  self.tab_calendar,
            3:  self.tab_acompte,
            4:  self.tab_charges,
            5:  self.tab_charges_omar,
            6:  self.tab_journal,
            7:  self.tab_journal_omar,
            8:  self.tab_modifier,
            9:  self.tab_export,
            10: self.tab_employee,
            11: self.tab_employee_list,
            12: self.tab_pointage,
            13: self.tab_centre_reception,
            14: self.tab_rules,
        }

    def _select_tab(self, idx):
        # Hide all
        for f in self._frames.values():
            f.place_forget()
        if self._current is not None and self._current in self._nav_btns:
            self._nav_btns[self._current].set_active(False)
        self._current = idx
        self._frames[idx].place(relx=0, rely=0, relwidth=1, relheight=1)
        if idx in self._nav_btns:
            self._nav_btns[idx].set_active(True)

    def _toggle_sidebar(self):
        if self._sidebar_visible:
            self.sidebar.pack_forget()
            self._toggle_btn.configure(text="›")
        else:
            # Re-insert sidebar before the toggle strip
            self.sidebar.pack(side="left", fill="y", before=self.toggle_strip)
            self._toggle_btn.configure(text="‹")
        self._sidebar_visible = not self._sidebar_visible

    def refresh_list(self):
        self.tab_list.refresh()
        self.tab_calendar.render_year()
        if hasattr(self, "tab_acompte") and self.tab_acompte.selected_id:
            self.tab_acompte._reload_selected()
    
    def _select_first_available_tab(self):
        """Sélectionne le premier onglet disponible pour l'utilisateur."""
        if self._visible_nav_items:
            first_idx = self._visible_nav_items[0][0]
            self._select_tab(first_idx)
    
    def _logout(self):
        """Déconnexion : ferme l'application et relance le login."""
        if messagebox.askyesno("Déconnexion", 
                               "Voulez-vous vraiment vous déconnecter ?"):
            self.destroy()
            # Relancer le processus de login
            from login_window import show_login
            user = show_login()
            if user:
                run(user)
    
    def _open_access_management(self):
        """Ouvre la fenêtre de gestion des accès (admin only)."""
        AccessManagementWindow(self)


# ═══════════════════════════════════════════════════════════════════
#  TAB — NOUVELLE RÉSERVATION
# ═══════════════════════════════════════════════════════════════════
class NewReservationTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._build()

    def _build(self):
        # ── Page header ──────────────────────────────────────────
        hdr = tk.Frame(self, bg=C["surface"], pady=0)
        hdr.pack(fill="x")
        inner_hdr = tk.Frame(hdr, bg=C["surface"])
        inner_hdr.pack(fill="x", padx=20, pady=14)

        left = tk.Frame(inner_hdr, bg=C["surface"])
        left.pack(side="left")
        tk.Label(left, text="Nouvelle réservation", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(anchor="w")
        tk.Label(left, text="Remplissez le formulaire pour créer un nouveau contrat",
                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"]).pack(anchor="w")

        right = tk.Frame(inner_hdr, bg=C["surface"])
        right.pack(side="right")
        tk.Label(right, text="N° prévu", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(anchor="e")
        self.id_label = tk.Label(right, text=self._next_id_preview(),
                                  bg=C["surface"], fg=C["primary"], font=FNT["id"])
        self.id_label.pack(anchor="e")

        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # ── Scrollable body ───────────────────────────────────────
        outer, sf = _scrollable(self)
        outer.pack(fill="both", expand=True)

        # ── SECTION : Créneau ────────────────────────────────────
        sec1 = self._section(sf, "Créneau & Salle")
        sec1_body = tk.Frame(sec1, bg=C["surface"])
        sec1_body.pack(fill="x", padx=16, pady=(0, 14))
        for c in (1, 3, 5, 7):
            sec1_body.columnconfigure(c, weight=1)

        # Salle selector (visual pills — multi-select)
        tk.Label(sec1_body, text="Salle de fête", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(
            row=0, column=0, sticky="w", padx=(0, 8), pady=(0, 4))
        self._salles_selected = set()   # replaces salle_var (multi-select)
        self.salle_var = tk.StringVar(value="Almes")  # kept for compat (first selected)
        salle_frame = tk.Frame(sec1_body, bg=C["surface"])
        salle_frame.grid(row=1, column=0, columnspan=2, sticky="w", padx=(0, 20), pady=(0, 10))
        self._salle_btns = {}
        SALLE_COLORS_UI = {"Almes": "#2563EB", "Chichkhane": "#DC2626", "Rayhane": "#16A34A"}
        for salle in rules.SALLES:
            clr = SALLE_COLORS_UI[salle]
            b = tk.Button(salle_frame, text=salle,
                          bg=C["stripe"], fg=C["text"], relief="flat",
                          font=FNT["sm"], padx=14, pady=6, cursor="hand2",
                          command=lambda s=salle: self._toggle_salle(s))
            b.pack(side="left", padx=(0, 6))
            self._salle_btns[salle] = (b, clr)
        # "Toute la complexe" button
        self._btn_complexe = tk.Button(
            salle_frame, text="🏛 Toute la complexe",
            bg="#7C3AED", fg="white", relief="flat",
            font=FNT["sm"], padx=14, pady=6, cursor="hand2",
            command=self._select_all_salles)
        self._btn_complexe.pack(side="left", padx=(12, 0))
        self._toggle_salle("Almes")  # default selection

        tk.Label(sec1_body, text="Date événement", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(
            row=0, column=2, sticky="w", padx=(0, 8))
        self.date_evt_entry = RoundedEntry(sec1_body, width=13)
        self.date_evt_entry.insert(0, date.today().strftime("%d/%m/%Y"))
        self.date_evt_entry.grid(row=1, column=2, sticky="ew", padx=(0, 12))

        tk.Label(sec1_body, text="Heure début", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=4, sticky="w")
        self.heure_debut = RoundedEntry(sec1_body, width=9)
        self.heure_debut.insert(0, "15:00")
        self.heure_debut.grid(row=1, column=4, sticky="ew", padx=(0, 8))

        tk.Label(sec1_body, text="Heure fin", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=6, sticky="w")
        self.heure_fin = RoundedEntry(sec1_body, width=9)
        self.heure_fin.insert(0, "18:00")
        self.heure_fin.grid(row=1, column=6, sticky="ew", padx=(0, 12))

        btn_dispo = RoundedButton(sec1_body, text="Vérifier disponibilité",
                                  command=self.refresh_availability, style="secondary",
                                  padx=16, pady=7)
        import tkinter.font as _tf; _f = _tf.Font(font=FNT["bold"])
        btn_dispo.configure(width=_f.measure("Vérifier disponibilité")+34,
                            height=_f.metrics("linespace")+16)
        btn_dispo.grid(row=1, column=8, sticky="w", padx=(8, 0))

        self.avail_status = tk.Label(sec1_body, text="", bg=C["surface"],
                                     fg=C["text_m"], font=FNT["sm"])
        self.avail_status.grid(row=2, column=0, columnspan=9, sticky="w",
                               padx=0, pady=(6, 0))

        # Availability mini-table
        avail_wrap = tk.Frame(sec1, bg=C["surface"])
        avail_wrap.pack(fill="x", padx=16, pady=(0, 14))
        cols_a = ("salle", "reserve_a", "temp", "gazon")
        self.avail_tree = ttk.Treeview(avail_wrap, columns=cols_a, show="headings", height=4)
        for c, lbl, w in (("salle","Salle",100),("reserve_a","Client",200),
                           ("temp","Créneau",140),("gazon","Gazon (Rayhane)",180)):
            self.avail_tree.heading(c, text=lbl)
            self.avail_tree.column(c, width=w, minwidth=50, stretch=True, anchor="center")
        av_sb = ttk.Scrollbar(avail_wrap, orient="vertical", command=self.avail_tree.yview)
        self.avail_tree.configure(yscrollcommand=av_sb.set)
        av_sb.pack(side="right", fill="y")
        self.avail_tree.pack(fill="x")

        # ── SECTION : Client ─────────────────────────────────────
        sec2 = self._section(sf, "Informations client")
        g2 = tk.Frame(sec2, bg=C["surface"])
        g2.pack(fill="x", padx=16, pady=(0, 14))
        for c in range(4):
            g2.columnconfigure(c, weight=1)

        self.nom_client = _field(g2, 0, 0, "Nom et prénom", colspan=2, bg=C["surface"])
        self.tel1       = _field(g2, 0, 2, "GSM 1", bg=C["surface"])
        self.tel2       = _field(g2, 0, 3, "GSM 2", bg=C["surface"])
        self.cin        = _field(g2, 1, 0, "CIN", bg=C["surface"])
        self.date_cin   = _field(g2, 1, 1, "Date CIN (JJ/MM/AAAA)", bg=C["surface"])
        self.lieu_cin   = _field(g2, 1, 2, "Lieu CIN", bg=C["surface"])

        # Blocage checkbox in its own styled frame
        blk_fr = tk.Frame(g2, bg=C["surface"])
        blk_fr.grid(row=1, column=3, sticky="nsew", padx=(4, 8), pady=4)
        blk_fr.columnconfigure(0, weight=1)
        tk.Label(blk_fr, text=" ", bg=C["surface"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.blocage_var = tk.IntVar(value=0)
        cb = tk.Checkbutton(blk_fr, text="Blocage toute la journée",
                            variable=self.blocage_var,
                            bg=C["surface"], fg=C["text"],
                            activebackground=C["surface"],
                            selectcolor=C["surface"], font=FNT["default"])
        cb.grid(row=1, column=0, sticky="w")

        self.nb_places = _field(g2, 2, 0, "Nombre de places", bg=C["surface"])
        self.type_evt  = _field(g2, 2, 1, "Type d'événement", bg=C["surface"])

        # ── SECTION : Contrat & Paiement ─────────────────────────
        sec3 = self._section(sf, "Contrat & Paiement")
        g3 = tk.Frame(sec3, bg=C["surface"])
        g3.pack(fill="x", padx=16, pady=(0, 14))
        for c in range(4):
            g3.columnconfigure(c, weight=1)

        self.date_contrat = _field(g3, 0, 0, "Date de contrat (JJ/MM/AAAA)", bg=C["surface"])
        self.forfait      = _field(g3, 0, 1, "Forfait (DT)", bg=C["surface"])
        self.forfait.bind("<FocusOut>", lambda e: self._recalc_reste())

        # Type combo
        ft_fr = tk.Frame(g3, bg=C["surface"])
        ft_fr.grid(row=0, column=2, sticky="nsew", padx=(4, 8), pady=4)
        ft_fr.columnconfigure(0, weight=1)
        tk.Label(ft_fr, text="Type forfait", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.forfait_type_var = tk.StringVar(value="Forfait")
        ttk.Combobox(ft_fr, textvariable=self.forfait_type_var,
                     values=["Forfait", "Extra"], state="readonly"
                     ).grid(row=1, column=0, sticky="ew", ipady=2)

        self.n_facture = _field(g3, 0, 3, "N° Facture", bg=C["surface"])

        # Payment row
        mp_fr = tk.Frame(g3, bg=C["surface"])
        mp_fr.grid(row=1, column=0, sticky="nsew", padx=(4, 8), pady=4)
        mp_fr.columnconfigure(0, weight=1)
        tk.Label(mp_fr, text="Méthode de paiement", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.methode_paiement_var = tk.StringVar(value="Espèce")
        mp_cb = ttk.Combobox(mp_fr, textvariable=self.methode_paiement_var,
                             values=["Espèce", "Chèque", "Virement", "Effet"], state="readonly")
        mp_cb.grid(row=1, column=0, sticky="ew", ipady=2)
        mp_cb.bind("<<ComboboxSelected>>", lambda e: self._toggle_payment_details())

        self.acompte1 = _field(g3, 1, 1, "Acompte (DT)", bg=C["surface"])
        self.acompte1.bind("<FocusOut>", lambda e: self._recalc_reste())
        self.caisse    = _field(g3, 1, 2, "Caisse Tawfik", bg=C["surface"])
        self.tresorerie= _field(g3, 1, 3, "Caisse Omar", bg=C["surface"])

        # Chèque/Virement details (collapsible)
        self.payment_details_frame = tk.LabelFrame(
            sec3, text="Détails du paiement (chèque / virement / effet)",
            bg=C["surface"], fg=C["primary"], font=FNT["bold"],
            relief="solid", bd=1)
        for _c in (1, 3, 5):
            self.payment_details_frame.columnconfigure(_c, weight=1)
        self.payment_details_frame.pack_forget()

        tk.Label(self.payment_details_frame, text="N° Chèque/Virement",
                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"]).grid(
            row=0, column=0, sticky="w", padx=(12, 4), pady=8)
        self.num_cheque_virement = RoundedEntry(self.payment_details_frame)
        self.num_cheque_virement.grid(row=0, column=1, sticky="ew", padx=(0, 16))

        tk.Label(self.payment_details_frame, text="Banque",
                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"]).grid(
            row=0, column=2, sticky="w", padx=(0, 4))
        self.banque = RoundedEntry(self.payment_details_frame)
        self.banque.grid(row=0, column=3, sticky="ew", padx=(0, 16))

        tk.Label(self.payment_details_frame, text="Date (JJ/MM/AAAA)",
                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"]).grid(
            row=0, column=4, sticky="w", padx=(0, 4))
        self.date_cheque_virement = RoundedEntry(self.payment_details_frame, width=14)
        self.date_cheque_virement.grid(row=0, column=5, sticky="ew",
                                       padx=(0, 12), pady=8)

        # Reste d'acompte row
        reste_sec = tk.Frame(sec3, bg=C["surface"])
        reste_sec.pack(fill="x", padx=16, pady=(4, 14))
        for c in (1, 3):
            reste_sec.columnconfigure(c, weight=1)

        tk.Label(reste_sec, text="Reste d'acompte (DT)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.reste_acompte = RoundedEntry(reste_sec, width=13, state="readonly")
        self.reste_acompte.grid(row=0, column=1, sticky="ew", padx=(6, 24), ipady=2)

        tk.Label(reste_sec, text="Date du reste (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=2, sticky="w")
        self.date_reste = RoundedEntry(reste_sec, width=14)
        self.date_reste.grid(row=0, column=3, sticky="ew", padx=(6, 0), ipady=2)

        self.reste_warning = tk.Label(reste_sec, text="", bg=C["surface"],
                                      fg=C["warning"], font=FNT["sm"])
        self.reste_warning.grid(row=1, column=0, columnspan=4, sticky="w", pady=(4, 0))

        # Détails extra
        det_fr = tk.Frame(sec3, bg=C["surface"])
        det_fr.pack(fill="x", padx=16, pady=(0, 14))
        tk.Label(det_fr, text="Détails extra / notes", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(anchor="w")
        self.details_extra = tk.Text(det_fr, height=3, font=FNT["default"],
                                     relief="solid", bd=1, bg=C["surface"],
                                     insertbackground=C["text"])
        self.details_extra.pack(fill="x", pady=(4, 0))

        # Espace Gazon option
        gazon_fr = tk.Frame(sec3, bg=C["surface"])
        gazon_fr.pack(fill="x", padx=16, pady=(0, 10))
        self.opt_gazon_var = tk.IntVar(value=0)
        tk.Checkbutton(gazon_fr, text="Option Espace Gazon",
                       variable=self.opt_gazon_var,
                       bg=C["surface"], fg=C["text"],
                       activebackground=C["surface"],
                       selectcolor=C["surface"], font=FNT["default"]).pack(side="left")

        # ── SECTION : Violoniste ─────────────────────────────────
        sec4 = self._section(sf, "Option Violoniste — Espace Rayhane (1h gratuite)")
        v_body = tk.Frame(sec4, bg=C["surface"])
        v_body.pack(fill="x", padx=16, pady=(0, 14))
        for c in (1, 3):
            v_body.columnconfigure(c, weight=1)

        self.with_violoniste_var = tk.IntVar(value=0)
        tk.Checkbutton(v_body, text="Inclure l'option violoniste",
                       variable=self.with_violoniste_var,
                       command=self._toggle_violon,
                       bg=C["surface"], fg=C["text"],
                       activebackground=C["surface"],
                       selectcolor=C["surface"], font=FNT["bold"]).grid(
            row=0, column=0, columnspan=5, sticky="w", pady=(0, 8))

        tk.Label(v_body, text="Date salle Rayhane", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=1, column=0, sticky="w")
        self.violon_date = RoundedEntry(v_body, width=13, state="disabled")
        self.violon_date.grid(row=2, column=0, sticky="ew", padx=(0, 12), ipady=2)

        tk.Label(v_body, text="Heure début", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=1, column=2, sticky="w")
        self.violon_debut = RoundedEntry(v_body, width=9, state="disabled")
        self.violon_debut.insert(0, "19:00")
        self.violon_debut.grid(row=2, column=2, sticky="ew", padx=(0, 8), ipady=2)

        tk.Label(v_body, text="Heure fin", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=1, column=4, sticky="w")
        self.violon_fin = RoundedEntry(v_body, width=9, state="disabled")
        self.violon_fin.insert(0, "20:00")
        self.violon_fin.grid(row=2, column=4, sticky="ew", padx=(0, 12), ipady=2)

        btn_rayhane = RoundedButton(v_body, text="Vérifier dispo Rayhane",
                                    command=self._check_rayhane_dispo, style="secondary",
                                    padx=16, pady=7)
        import tkinter.font as _tf2; _f2 = _tf2.Font(font=FNT["bold"])
        btn_rayhane.configure(width=_f2.measure("Vérifier dispo Rayhane")+34,
                              height=_f2.metrics("linespace")+16)
        btn_rayhane.grid(row=2, column=6, sticky="w", padx=(8, 0))

        self.violon_dispo_lbl = tk.Label(v_body, text="", bg=C["surface"],
                                         fg=C["text_m"], font=FNT["sm"])
        self.violon_dispo_lbl.grid(row=3, column=0, columnspan=7, sticky="w", pady=(6, 0))

        self.with_sono_var = tk.IntVar(value=0)
        tk.Checkbutton(v_body, text="Sans sono (requis pour réception Rayhane)",
                       variable=self.with_sono_var,
                       bg=C["surface"], fg=C["text"],
                       activebackground=C["surface"],
                       selectcolor=C["surface"], font=FNT["default"]).grid(
            row=4, column=0, columnspan=5, sticky="w", pady=(8, 0))

        # ── SECTION : Infos contrat arabe ────────────────────────
        sec5 = self._section(sf, "معلومات العقد — Informations du contrat (arabe)")
        ar_g = tk.Frame(sec5, bg=C["surface"])
        ar_g.pack(fill="x", padx=16, pady=(0, 14))
        for c in range(4):
            ar_g.columnconfigure(c, weight=1)

        def ar_lbl(parent, text, row, col):
            tk.Label(parent, text=text, bg=C["surface"],
                     fg=C["text_m"], font=FNT["sm"], anchor="e"
                     ).grid(row=row*2, column=col, sticky="e", padx=(8, 4))
        def ar_ent(parent, row, col):
            e = RoundedEntry(parent, font=("Segoe UI", 10), justify="right")
            e.grid(row=row*2+1, column=col, sticky="ew", padx=(0, 16), pady=(0, 10), ipady=2)
            return e

        ar_lbl(ar_g, "الاسم و اللقب", 0, 0)
        self.ar_nom = ar_ent(ar_g, 0, 0)
        ar_lbl(ar_g, "مكان بطاقة التعريف", 0, 1)
        self.ar_lieu_cin = ar_ent(ar_g, 0, 1)
        ar_lbl(ar_g, "قاعة الأفراح", 0, 2)
        self.ar_salle = ar_ent(ar_g, 0, 2)
        ar_lbl(ar_g, "تاريخ التحرير", 0, 3)
        self.ar_date = ar_ent(ar_g, 0, 3)
        self.ar_date.insert(0, date.today().strftime("%d/%m/%Y"))
        ar_lbl(ar_g, "عدد العقد", 1, 0)
        self.ar_num_contrat = ar_ent(ar_g, 1, 0)
        self.ar_num_contrat.insert(0, f"{date.today().year}/")

        self._refresh_arabic_info()

        # ── Action buttons ───────────────────────────────────────
        btn_bar = tk.Frame(sf, bg=C["bg"])
        btn_bar.pack(fill="x", padx=20, pady=(12, 20))

        _rbtn(btn_bar, "✔  Enregistrer la réservation", command=self.save_reservation, style="primary", padx=16, pady=7)
        _rbtn(btn_bar, "📄  Contrat (FR)", command=self.generate_contract_from_form, style="secondary", padx=16, pady=7)
        _rbtn(btn_bar, "📜  Contrat (عربي)", command=self.generate_arabic_contract_from_form, style="success", padx=16, pady=7)
        _rbtn(btn_bar, "🔄 Actualiser infos arabe", command=self._refresh_arabic_info, style="secondary", padx=16, pady=7)

        # Big Temporaire button — distinct yellow, right-aligned
        tmp_btn = tk.Button(
            btn_bar, text="⏳  TEMPORAIRE",
            bg="#D97706", fg="white",
            activebackground="#B45309", activeforeground="white",
            relief="flat", bd=0, cursor="hand2",
            font=("Segoe UI", 11, "bold"),
            padx=22, pady=9,
            command=self.save_temporaire,
        )
        tmp_btn.pack(side="right", padx=(12, 0))

        # Status bar
        self.status_label = tk.Label(sf, text="", bg=C["bg"],
                                     fg=C["danger"], font=FNT["sm"])
        self.status_label.pack(padx=20, pady=(0, 10), anchor="w")

    # ── Section helper ────────────────────────────────────────────
    def _section(self, parent, title):
        """Carte blanche titrée : bordure fine, ombre douce, barre d'accent."""
        wrapper = tk.Frame(parent, bg=C["bg"])
        wrapper.pack(fill="x", padx=22, pady=(0, 14))

        shadow = tk.Frame(wrapper, bg=C["shadow"])
        shadow.pack(fill="x", padx=(1, 0), pady=(2, 0))
        shell = tk.Frame(shadow, bg=C["border"])
        shell.pack(fill="x", padx=(0, 1), pady=(0, 2))
        card = tk.Frame(shell, bg=C["surface"])
        card.pack(fill="x", padx=1, pady=1)

        hdr = tk.Frame(card, bg=C["surface"])
        hdr.pack(fill="x", padx=18, pady=(14, 0))
        tk.Frame(hdr, bg=C["primary"], width=3, height=16).pack(
            side="left", padx=(0, 10))
        tk.Label(hdr, text=title, bg=C["surface"], fg=C["text"],
                 font=FNT["section"]).pack(side="left")
        tk.Frame(card, bg=C["border"], height=1).pack(
            fill="x", padx=18, pady=(12, 0))
        return card

    # ── Salle pill selector (multi-select) ───────────────────────
    def _toggle_salle(self, salle):
        """Toggle a single salle on/off.  At least one must stay selected."""
        if salle in self._salles_selected:
            if len(self._salles_selected) > 1:          # don't deselect the last one
                self._salles_selected.discard(salle)
        else:
            self._salles_selected.add(salle)
        self._refresh_salle_buttons()
        self._refresh_arabic_info()

    def _select_all_salles(self):
        """Select all salles (Toute la complexe).  If already all selected, deselect to just Almes."""
        all_salles = set(rules.SALLES)
        if self._salles_selected == all_salles:
            # toggle off: revert to default single selection
            self._salles_selected = {"Almes"}
        else:
            self._salles_selected = all_salles.copy()
        self._refresh_salle_buttons()
        self._refresh_arabic_info()

    def _refresh_salle_buttons(self):
        """Repaint pill buttons to reflect current _salles_selected set."""
        SALLE_COLORS_UI = {"Almes": "#2563EB", "Chichkhane": "#DC2626", "Rayhane": "#16A34A"}
        all_salles = set(rules.SALLES)
        for s, (btn, clr) in self._salle_btns.items():
            if s in self._salles_selected:
                btn.configure(bg=clr, fg="white", font=FNT["bold"])
            else:
                btn.configure(bg=C["stripe"], fg=C["text"], font=FNT["sm"])
        # Highlight "Toute la complexe" button when all 3 are selected
        if self._salles_selected == all_salles:
            self._btn_complexe.configure(bg="#5B21B6", font=FNT["bold"])
        else:
            self._btn_complexe.configure(bg="#7C3AED", font=FNT["sm"])
        # Keep salle_var in sync with "primary" salle (first in list that is selected)
        for s in rules.SALLES:
            if s in self._salles_selected:
                self.salle_var.set(s)
                break

    # ── Helpers ───────────────────────────────────────────────────
    def _next_id_preview(self):
        conn = database.get_connection()
        row = conn.execute(
            "SELECT COALESCE(MAX(id),0)+1 AS next_id FROM reservations").fetchone()
        conn.close()
        return str(row["next_id"])

    def _refresh_arabic_info(self):
        # Guard: arabic widgets may not be built yet during initial _build()
        if not hasattr(self, "ar_nom"):
            return
        nom = self.nom_client.get().strip() if hasattr(self, "nom_client") else ""
        self.ar_nom.delete(0, "end"); self.ar_nom.insert(0, nom)
        lieu = self.lieu_cin.get().strip() if hasattr(self, "lieu_cin") else ""
        self.ar_lieu_cin.delete(0, "end"); self.ar_lieu_cin.insert(0, lieu)
        # Show "Toute la complexe" Arabic label when all 3 salles are selected
        salles_sel = self._salles_selected if hasattr(self, "_salles_selected") else {self.salle_var.get()}
        if salles_sel == set(rules.SALLES):
            salle_ar = contract_generator.SALLE_AR.get("Toute la complexe", "كامل الفضاء")
        else:
            salle_val = self.salle_var.get()
            salle_ar = contract_generator.SALLE_AR.get(salle_val, salle_val)
        self.ar_salle.delete(0, "end"); self.ar_salle.insert(0, salle_ar)
        if not self.ar_date.get().strip():
            self.ar_date.insert(0, date.today().strftime("%d/%m/%Y"))

    def _toggle_payment_details(self):
        methode = self.methode_paiement_var.get()
        if methode in ("Chèque", "Virement", "Effet"):
            self.payment_details_frame.pack(fill="x", padx=16, pady=(0, 10))
        else:
            self.payment_details_frame.pack_forget()

    def _recalc_reste(self):
        try:
            forfait = float(self.forfait.get().strip() or 0)
            acompte = float(self.acompte1.get().strip() or 0)
        except ValueError:
            return
        reste = forfait - acompte
        self.reste_acompte.configure(state="normal")
        self.reste_acompte.delete(0, "end")
        self.reste_acompte.insert(0, f"{reste:.3f}" if reste > 0 else "0.000")
        self.reste_acompte.configure(state="readonly")
        if forfait > 0 and reste >= forfait * 0.5:
            try:
                iso = self._to_iso_date(self.date_evt_entry.get().strip())
                d_evt = date.fromisoformat(iso)
                d_reste = d_evt - timedelta(days=15)
                self.date_reste.delete(0, "end")
                self.date_reste.insert(0, d_reste.strftime("%d/%m/%Y"))
                self.reste_warning.configure(
                    text=f"⚠  Reste ≥ 50% du forfait — échéance recommandée : {d_reste.strftime('%d/%m/%Y')}")
            except Exception:
                self.reste_warning.configure(text="")
        else:
            self.reste_warning.configure(text="")

    def _toggle_violon(self):
        state = "normal" if self.with_violoniste_var.get() else "disabled"
        for w in (self.violon_date, self.violon_debut, self.violon_fin):
            w.configure(state=state)

    def _check_rayhane_dispo(self):
        self.violon_dispo_lbl.configure(text="", fg=C["text_m"])
        raw = self.violon_date.get().strip() or self.date_evt_entry.get().strip()
        try:
            iso = self._to_iso_date(raw)
        except Exception:
            self.violon_dispo_lbl.configure(text="⚠  Date invalide.", fg=C["warning"])
            return
        hd, hf = self.violon_debut.get().strip(), self.violon_fin.get().strip()
        if not hd or not hf:
            self.violon_dispo_lbl.configure(text="⚠  Saisissez heure début et fin.", fg=C["warning"])
            return
        ok, msg = rules.check_violoniste_slot(None, iso, hd, hf, self.with_sono_var.get())
        self.violon_dispo_lbl.configure(
            text=f"✔  Rayhane disponible {raw} de {hd} à {hf}" if ok else f"✘  {msg}",
            fg=C["success"] if ok else C["danger"])

    @staticmethod
    def _to_iso_date(ddmmyyyy):
        d, m, y = ddmmyyyy.split("/")
        return f"{y}-{m}-{d}"

    def refresh_availability(self):
        for row in self.avail_tree.get_children():
            self.avail_tree.delete(row)
        self.avail_status.configure(text="", fg=C["text_m"])
        raw = self.date_evt_entry.get().strip()
        if not raw:
            return
        try:
            iso = self._to_iso_date(raw)
        except Exception:
            self.avail_status.configure(text="Date invalide (JJ/MM/AAAA).", fg=C["danger"])
            return
        all_res = database.get_reservations_for_date(iso)
        gazon_res = [r for r in all_res if r["salle"] == "Rayhane"]
        gazon_txt = " | ".join(
            f"{r['heure_debut']}-{r['heure_fin']}" for r in gazon_res) if gazon_res else "Libre"
        for res in all_res:
            g = gazon_txt if res["salle"] != "Rayhane" else "—"
            self.avail_tree.insert("", "end", values=(
                res["salle"], res.get("nom_client") or "",
                f"{res['heure_debut']} → {res['heure_fin']}", g))
        if not all_res:
            self.avail_tree.insert("", "end", values=("—", "Aucune réservation", "—", gazon_txt))
        hd = self.heure_debut.get().strip()
        hf = self.heure_fin.get().strip()
        salles_sel = self._salles_selected if hasattr(self, "_salles_selected") else {self.salle_var.get()}
        if hd and hf:
            is_complexe = (salles_sel == set(rules.SALLES))
            conflicts = []
            ok_salles = []
            for salle in rules.SALLES:
                if salle not in salles_sel:
                    continue
                v_on = bool(self.with_violoniste_var.get()) if hasattr(self, "with_violoniste_var") else False
                v_deb = self.violon_debut.get().strip() if v_on else None
                v_fin = self.violon_fin.get().strip() if v_on else None
                try:
                    v_date = (self._to_iso_date(self.violon_date.get().strip())
                              if v_on and self.violon_date.get().strip() else None)
                except Exception:
                    v_date = None
                test = {"salle": salle, "date_evenement": iso, "heure_debut": hd,
                        "heure_fin": hf,
                        "with_violoniste": 1 if (v_on and v_deb and v_fin) else 0,
                        "violoniste_date": v_date,
                        "violoniste_heure_debut": v_deb,
                        "violoniste_heure_fin": v_fin,
                        "with_sono": self.with_sono_var.get() if hasattr(self, "with_sono_var") else 0,
                        "type_evenement": ""}
                ok, msg = rules.validate_reservation(test, complexe=is_complexe)
                if ok:
                    ok_salles.append(salle)
                else:
                    conflicts.append(f"{salle}: {msg}")
            # Règle applicable (module « Règles de réservation »)
            applied = []
            for salle in ok_salles or sorted(salles_sel):
                r = rules.find_matching_rule(salle, iso, hd, hf)
                if r:
                    label = f"{salle} → « {r['nom']} »"
                    if r.get("violon_gratuit"):
                        label += " 🎻"
                    applied.append(label)
            regle_txt = ("   ⚖ " + " ; ".join(applied)) if applied else ""

            if conflicts:
                self.avail_status.configure(
                    text="✘  " + " | ".join(conflicts) + regle_txt, fg=C["danger"])
            else:
                self.avail_status.configure(
                    text=f"✔  Créneau {hd}-{hf} disponible pour {', '.join(ok_salles)}"
                         + regle_txt,
                    fg=C["success"])

    def generate_contract_from_form(self):
        try:
            res = self._collect_form_data(draft=True)
        except Exception as ex:
            messagebox.showerror("Erreur", f"Données invalides : {ex}")
            return
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        client = res["nom_client"] or "brouillon"
        filename = f"Contrat_{client}_{res['date_evenement']}.docx".replace("/", "-").replace(" ", "_")
        output_path = os.path.join(OUTPUT_DIR, filename)
        try:
            contract_generator.generate_contract(res, output_path)
            os.startfile(output_path)
        except Exception as ex:
            messagebox.showerror("Erreur", f"Impossible de générer le contrat : {ex}")

    def generate_arabic_contract_from_form(self):
        self._refresh_arabic_info()
        try:
            res = self._collect_form_data(draft=True)
            ar_info = {
                "ar_nom":      self.ar_nom.get().strip(),
                "ar_lieu_cin": self.ar_lieu_cin.get().strip(),
                "ar_salle":    self.ar_salle.get().strip(),
                "ar_date":     self.ar_date.get().strip(),
                "num_contrat": self.ar_num_contrat.get().strip(),
            }
        except Exception as ex:
            messagebox.showerror("Erreur", f"Données invalides : {ex}")
            return
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        client = res["nom_client"] or ar_info["ar_nom"] or "brouillon"
        filename = (f"Contrat_AR_{client}_{res['date_evenement']}.docx"
                    .replace("/", "-").replace(" ", "_"))
        output_path = os.path.join(OUTPUT_DIR, filename)
        try:
            contract_generator.generate_arabic_contract(res, ar_info, output_path)
            os.startfile(output_path)
        except Exception as ex:
            messagebox.showerror("Erreur", f"Impossible de générer : {ex}")

    def _collect_form_data(self, draft=False):
        salle = self.salle_var.get()  # primary salle (for single-contract preview)
        date_str = self._to_iso_date(self.date_evt_entry.get().strip())
        methode = self.methode_paiement_var.get()
        forfait_val = float(self.forfait.get()) if self.forfait.get().strip() else None
        acompte_val = float(self.acompte1.get()) if self.acompte1.get().strip() else None
        reste_val = float(self.reste_acompte.get()) if self.reste_acompte.get().strip() else None
        return {
            "id": "BROUILLON" if draft else None,
            "num_bon": "BROUILLON" if draft else None,
            "salle": salle,
            "date_evenement": date_str,
            "heure_debut": self.heure_debut.get().strip(),
            "heure_fin": self.heure_fin.get().strip(),
            "type_evenement": self.type_evt.get().strip(),
            "nb_places": self.nb_places.get().strip() or None,
            "forfait": forfait_val,
            "nom_client": self.nom_client.get().strip(),
            "cin": self.cin.get().strip(),
            "date_cin": self._to_iso_date(self.date_cin.get()) if self.date_cin.get().strip() else None,
            "lieu_cin": self.lieu_cin.get().strip(),
            "tel1": self.tel1.get().strip(),
            "tel2": self.tel2.get().strip(),
            "lieu_contrat": "",
            "date_contrat": self._to_iso_date(self.date_contrat.get()) if self.date_contrat.get().strip() else None,
            "heure_contrat": "",
            "dossier_traite_par": f"{self.app.current_user.get('prenom','').strip()} {self.app.current_user.get('nom','').strip()}".strip(),
            "with_violoniste": 1 if self.with_violoniste_var.get() else 0,
            "violoniste_date": self._to_iso_date(self.violon_date.get()) if self.with_violoniste_var.get() and self.violon_date.get().strip() else None,
            "violoniste_heure_debut": self.violon_debut.get() if self.with_violoniste_var.get() else None,
            "violoniste_heure_fin": self.violon_fin.get() if self.with_violoniste_var.get() else None,
            "acompte1": acompte_val,
            "date_acompte1": date.today().isoformat() if acompte_val else None,
            "fn1": None,
            "methode_acompte1": methode if acompte_val else None,
            "acompte2": None, "date_acompte2": None, "fn2": None, "methode_acompte2": None,
            "acompte3": None, "date_acompte3": None, "fn3": None, "methode_acompte3": None,
            "reste_acompte": reste_val,
            "date_reste": self._to_iso_date(self.date_reste.get()) if self.date_reste.get().strip() else None,
        }

    def save_temporaire(self):
        """Enregistre une réservation temporaire minimale (nom, salle, date, heures).
        Aucune validation de créneau ni de paiement. Apparaît en jaune sur le calendrier."""
        nom = self.nom_client.get().strip() if hasattr(self, "nom_client") else ""
        if not nom:
            messagebox.showwarning("Champ manquant", "Saisissez au moins le nom du client.")
            return
        try:
            date_str = self._to_iso_date(self.date_evt_entry.get().strip())
        except Exception:
            messagebox.showerror("Erreur", "Date événement invalide (JJ/MM/AAAA).")
            return
        heure_debut = self.heure_debut.get().strip()
        heure_fin   = self.heure_fin.get().strip()
        if not heure_debut or not heure_fin:
            messagebox.showwarning("Champ manquant", "Saisissez l'heure de début et de fin.")
            return

        salles_to_save = sorted(
            self._salles_selected if hasattr(self, "_salles_selected") else {self.salle_var.get()},
            key=lambda s: rules.SALLES.index(s),
        )

        new_ids = []
        conn = database.get_connection()
        try:
            for salle in salles_to_save:
                res = {
                    "salle":          salle,
                    "date_evenement": date_str,
                    "heure_debut":    heure_debut,
                    "heure_fin":      heure_fin,
                    "nom_client":     nom,
                    "is_temporaire":  1,
                    "num_bon":        None,
                    "date_signature": date.today().isoformat(),
                }
                new_id = database.insert_reservation(res)
                conn.execute("UPDATE reservations SET num_bon=? WHERE id=?",
                             (f"TMP-{date.today().year}-{new_id:04d}", new_id))
                new_ids.append(new_id)
            conn.commit()
        finally:
            conn.close()

        ids_str   = ", ".join(f"#{i}" for i in new_ids)
        salles_str = ", ".join(salles_to_save)
        self.status_label.configure(
            text=f"⏳  Temporaire enregistrée ({salles_str}) — ID {ids_str}",
            fg="#D97706")
        messagebox.showinfo("Temporaire enregistrée",
            f"Réservation temporaire créée.\nSalles : {salles_str}\nID : {ids_str}")
        self.id_label.configure(text=self._next_id_preview())
        self.refresh_availability()
        self.app.refresh_list()

    def save_reservation(self):
        try:
            salles_to_save = sorted(
                self._salles_selected if hasattr(self, "_salles_selected") else {self.salle_var.get()},
                key=lambda s: rules.SALLES.index(s)
            )
            salles_to_save_set = set(salles_to_save)
            date_str = self._to_iso_date(self.date_evt_entry.get().strip())
            heure_debut = self.heure_debut.get().strip()
            heure_fin = self.heure_fin.get().strip()
            methode = self.methode_paiement_var.get()
            if methode in ("Chèque", "Virement", "Effet"):
                if not self.num_cheque_virement.get().strip():
                    raise ValueError(f"N° {methode} obligatoire.")
                if not self.banque.get().strip():
                    raise ValueError("Banque obligatoire.")
                if not self.date_cheque_virement.get().strip():
                    raise ValueError(f"Date du {methode} obligatoire.")
            acompte_val = float(self.acompte1.get()) if self.acompte1.get().strip() else None
            forfait_val = float(self.forfait.get()) if self.forfait.get().strip() else None
            reste_val = float(self.reste_acompte.get()) if self.reste_acompte.get().strip() else None
            if forfait_val and reste_val and reste_val >= forfait_val * 0.5:
                if not self.date_reste.get().strip():
                    raise ValueError("Date du reste obligatoire (reste ≥ 50% du forfait).")
                d_reste = date.fromisoformat(self._to_iso_date(self.date_reste.get().strip()))
                d_evt = date.fromisoformat(date_str)
                if d_reste > d_evt - timedelta(days=15):
                    raise ValueError(
                        f"Date du reste doit être ≥ 15j avant l'événement "
                        f"(au plus tard le {(d_evt - timedelta(days=15)).strftime('%d/%m/%Y')}).")

            # Build base reservation dict (salle will be overwritten per iteration)
            base_res = {
                "salle": None,  # set per salle below
                "date_evenement": date_str,
                "heure_debut": heure_debut, "heure_fin": heure_fin,
                "type_evenement": self.type_evt.get().strip(),
                "nb_places": int(self.nb_places.get()) if self.nb_places.get().strip() else None,
                "forfait": forfait_val,
                "nom_client": self.nom_client.get().strip(),
                "cin": self.cin.get().strip(),
                "date_cin": self._to_iso_date(self.date_cin.get()) if self.date_cin.get().strip() else None,
                "lieu_cin": self.lieu_cin.get().strip(),
                "tel1": self.tel1.get().strip(),
                "tel2": self.tel2.get().strip(),
                "lieu_contrat": "", "heure_contrat": "", "dossier_traite_par": f"{self.app.current_user.get('prenom','').strip()} {self.app.current_user.get('nom','').strip()}".strip(),
                "date_contrat": self._to_iso_date(self.date_contrat.get()) if self.date_contrat.get().strip() else None,
                "with_violoniste": 1 if self.with_violoniste_var.get() else 0,
                "violoniste_date": self._to_iso_date(self.violon_date.get()) if self.with_violoniste_var.get() and self.violon_date.get().strip() else None,
                "violoniste_heure_debut": self.violon_debut.get() if self.with_violoniste_var.get() else None,
                "violoniste_heure_fin": self.violon_fin.get() if self.with_violoniste_var.get() else None,
                "with_sono": 1 if self.with_sono_var.get() else 0,
                "date_signature": date.today().isoformat(),
                "statut_paiement": None,
                "blocage_journee": 1 if self.blocage_var.get() else 0,
                "forfait_type": self.forfait_type_var.get(),
                "methode_paiement": methode,
                "caisse": self.caisse.get().strip(),
                "tresorerie": self.tresorerie.get().strip(),
                "reste_acompte": reste_val,
                "date_reste": self._to_iso_date(self.date_reste.get()) if self.date_reste.get().strip() else None,
                "details_extra": self.details_extra.get("1.0", "end").strip(),
                "opt_espace_gazon": 1 if self.opt_gazon_var.get() else 0,
                "opt_tourniquet": 0, "opt_logiciel": 0, "opt_invitations": 0,
                "n_facture": self.n_facture.get().strip(),
                "num_cheque_virement": self.num_cheque_virement.get().strip() if methode in ("Chèque", "Virement", "Effet") else None,
                "banque": self.banque.get().strip() if methode in ("Chèque", "Virement", "Effet") else None,
                "date_cheque_virement": self._to_iso_date(self.date_cheque_virement.get()) if methode in ("Chèque", "Virement", "Effet") and self.date_cheque_virement.get().strip() else None,
                "acompte1": acompte_val,
                "date_acompte1": date.today().isoformat() if acompte_val else None,
                "fn1": self.num_cheque_virement.get().strip() if methode in ("Chèque", "Virement", "Effet") else None,
                "methode_acompte1": methode if acompte_val else None,
                "num_caisse1": None,
            }
            for i in range(2, 4):
                base_res[f"acompte{i}"] = None; base_res[f"date_acompte{i}"] = None
                base_res[f"fn{i}"] = None; base_res[f"methode_acompte{i}"] = None
                base_res[f"num_caisse{i}"] = None
        except Exception as ex:
            self.status_label.configure(text=f"Erreur : {ex}", fg=C["danger"])
            messagebox.showerror("Erreur de saisie", str(ex))
            return

        # Validate each selected salle before inserting any
        is_complexe = (salles_to_save_set == set(rules.SALLES))
        for salle in salles_to_save:
            check = dict(base_res); check["salle"] = salle
            ok, msg = rules.validate_reservation(check, complexe=is_complexe)
            if not ok:
                self.status_label.configure(text=msg, fg=C["danger"])
                messagebox.showerror(f"Réservation refusée — {salle}", msg)
                return

        # Insert one reservation per selected salle
        new_ids = []
        conn = database.get_connection()
        try:
            for salle in salles_to_save:
                res = dict(base_res); res["salle"] = salle
                res["num_bon"] = None
                new_id = database.insert_reservation(res)
                conn.execute("UPDATE reservations SET num_bon=? WHERE id=?",
                             (f"BON-{date.today().year}-{new_id:04d}", new_id))
                new_ids.append(new_id)
            conn.commit()
        finally:
            conn.close()

        ids_str = ", ".join(f"#{i}" for i in new_ids)
        salles_str = ", ".join(salles_to_save)
        self.status_label.configure(
            text=f"✔  {len(new_ids)} réservation(s) enregistrée(s) ({salles_str}) — ID {ids_str}",
            fg=C["success"])
        toast(f"{len(new_ids)} réservation(s) — {salles_str} (ID {ids_str})",
              "success", "Réservation enregistrée")
        self.id_label.configure(text=self._next_id_preview())
        self.refresh_availability()
        self.app.refresh_list()


# ═══════════════════════════════════════════════════════════════════
#  TAB — LISTE DES RÉSERVATIONS
# ═══════════════════════════════════════════════════════════════════
class ListTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._sort_col = "date"
        self._sort_rev = False
        self._build()
        self.refresh()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        h_inner = tk.Frame(hdr, bg=C["surface"])
        h_inner.pack(fill="x", padx=20, pady=14)
        tk.Label(h_inner, text="Réservations", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(side="left", anchor="w")
        self.count_label = tk.Label(h_inner, text="", bg=C["surface"],
                                    fg=C["text_m"], font=FNT["sm"])
        self.count_label.pack(side="left", padx=16, anchor="w")
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Search bar
        sb = tk.Frame(self, bg=C["surface"], pady=0)
        sb.pack(fill="x")
        sb_inner = tk.Frame(sb, bg=C["surface"])
        sb_inner.pack(fill="x", padx=20, pady=12)

        tk.Label(sb_inner, text="🔍", bg=C["surface"], fg=C["text_m"],
                 font=("Segoe UI", 12)).pack(side="left", padx=(0, 6))
        self.search_var = tk.StringVar()
        se = RoundedEntry(sb_inner, textvariable=self.search_var, width=36)
        se.pack(side="left", ipady=3)
        se.bind("<Return>", lambda e: self._search())
        tk.Label(sb_inner, text="(Tél / CIN / Date / N° Bon / Facture)",
                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"]).pack(side="left", padx=8)
        _rbtn(sb_inner, "Rechercher", command=self._search, style="sm", padx=10, pady=5)
        _rbtn(sb_inner, "Tout afficher", command=self.refresh, style="secondary", padx=16, pady=7)

        # Action toolbar
        tb = tk.Frame(self, bg=C["bg"])
        tb.pack(fill="x", padx=20, pady=(8, 0))
        _rbtn(tb, "🔄  Actualiser", command=self.refresh, style="secondary", padx=16, pady=7)
        _rbtn(tb, "📄  Ouvrir bon de reçu", command=self.open_bon, style="secondary", padx=16, pady=7)
        _rbtn(tb, "📝  Générer contrat Word", command=self.generate_contract, style="secondary", padx=16, pady=7)
        _rbtn(tb, "🗑  Supprimer", command=self.delete_selected, style="sm_danger", padx=10, pady=5)

        # Treeview
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=10)

        columns = ("id", "num_bon", "salle", "date", "heure",
                   "client", "type", "forfait", "paye", "reste", "statut")
        self.tree = ttk.Treeview(tree_wrap, columns=columns, show="headings")
        headings = {"id": "ID", "num_bon": "N° Bon", "salle": "Salle",
                    "date": "Date", "heure": "Créneau", "client": "Client",
                    "type": "Type", "forfait": "Forfait DT",
                    "paye": "Payé DT", "reste": "Reste DT", "statut": "Statut"}
        widths = {"id": 40, "num_bon": 115, "salle": 90, "date": 90,
                  "heure": 110, "client": 160, "type": 95,
                  "forfait": 85, "paye": 85, "reste": 85, "statut": 165}
        for c in columns:
            self.tree.heading(c, text=headings[c],
                              command=lambda _c=c: self._sort_by(_c))
            self.tree.column(c, width=widths[c], minwidth=40,
                             stretch=True, anchor="center")

        self.tree.tag_configure("solde",   background=C["tag_solde"])
        self.tree.tag_configure("partiel", background=C["tag_part"])
        self.tree.tag_configure("zero",    background=C["tag_zero"])

        vsb = ttk.Scrollbar(tree_wrap, orient="vertical",   command=self.tree.yview)
        hsb = ttk.Scrollbar(tree_wrap, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        vsb.pack(side="right", fill="y")
        hsb.pack(side="bottom", fill="x")
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<Double-1>", lambda e: self.open_bon())

    def _sort_by(self, col):
        self._sort_rev = (col == self._sort_col) and not self._sort_rev
        self._sort_col = col
        rows = [(self.tree.set(k, col), k) for k in self.tree.get_children("")]
        rows.sort(reverse=self._sort_rev)
        for i, (_, k) in enumerate(rows):
            self.tree.move(k, "", i)

    def _populate(self, reservations):
        for row in self.tree.get_children():
            self.tree.delete(row)
        for r in reservations:
            forfait = float(r.get("forfait") or 0)
            paye    = sum(float(r.get(f"acompte{i}") or 0) for i in range(1, 4))
            reste   = forfait - paye
            statut  = r.get("statut_paiement") or (
                "SOLDÉ" if reste <= 0.001 and forfait > 0 else "En attente")
            tag = "solde" if "SOLDE" in statut.upper() else (
                "partiel" if paye > 0 else "zero")
            self.tree.insert("", "end", iid=str(r["id"]), tags=(tag,), values=(
                r["id"], r.get("num_bon") or "", r.get("salle") or "",
                r.get("date_evenement") or "",
                f"{r.get('heure_debut','')}-{r.get('heure_fin','')}",
                r.get("nom_client") or "", r.get("type_evenement") or "",
                f"{forfait:,.3f}" if forfait else "",
                f"{paye:,.3f}" if paye else "",
                f"{reste:,.3f}" if forfait else "", statut,
            ))
        self.count_label.configure(text=f"{len(reservations)} réservation(s)")

    def refresh(self):
        self.search_var.set("")
        self._populate(database.get_all_reservations())

    def _search(self):
        q = self.search_var.get().strip()
        if not q:
            self.refresh(); return
        self._populate(database.search_reservations(q))

    def _selected_id(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Sélection requise",
                                   "Sélectionne une réservation dans la liste.")
            return None
        return int(sel[0])

    def open_bon(self):
        res_id = self._selected_id()
        if res_id is None: return
        res = database.get_reservation(res_id)
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        output_path = os.path.join(OUTPUT_DIR,
            f"Bon_Reservation_{res.get('num_bon') or res_id}.docx".replace("/", "-"))
        try:
            contract_generator.generate_contract(res, output_path)
            os.startfile(output_path)
        except Exception as e:
            messagebox.showerror("Erreur", str(e))

    def generate_contract(self):
        res_id = self._selected_id()
        if res_id is None: return
        res = database.get_reservation(res_id)
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        output_path = os.path.join(OUTPUT_DIR,
            f"Bon_Reservation_{res['num_bon'] or res_id}.docx".replace("/", "-"))
        contract_generator.generate_contract(res, output_path)
        toast(os.path.basename(output_path), "success", "Contrat généré")

    def delete_selected(self):
        res_id = self._selected_id()
        if res_id is None: return
        if messagebox.askyesno("Confirmer", f"Supprimer la réservation #{res_id} ?"):
            database.delete_reservation(res_id)
            self.refresh()
            self.app.tab_calendar.render_year()


# ═══════════════════════════════════════════════════════════════════
#  TAB — ACOMPTE
# ═══════════════════════════════════════════════════════════════════
class AcompteTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self.selected_id = None
        self._build()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        tk.Label(hdr, text="Ajouter un acompte", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Search
        sf = tk.Frame(self, bg=C["surface"])
        sf.pack(fill="x")
        sf_in = tk.Frame(sf, bg=C["surface"])
        sf_in.pack(fill="x", padx=20, pady=12)
        tk.Label(sf_in, text="🔍", bg=C["surface"], fg=C["text_m"],
                 font=("Segoe UI", 12)).pack(side="left", padx=(0, 6))
        self.search_var = tk.StringVar()
        se = RoundedEntry(sf_in, textvariable=self.search_var, width=36)
        se.pack(side="left", ipady=3)
        se.bind("<Return>", lambda e: self._search())
        _rbtn(sf_in, "Rechercher", command=self._search, style="sm", padx=10, pady=5)
        _rbtn(sf_in, "Tout afficher", command=self._show_all, style="secondary", padx=16, pady=7)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Results tree
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="x", padx=20, pady=10)
        cols = ("id", "num_bon", "client", "salle", "date", "forfait", "total_paye", "statut")
        self.result_tree = ttk.Treeview(tree_wrap, columns=cols, show="headings", height=5)
        heads = {"id":"ID","num_bon":"N° Bon","client":"Client","salle":"Salle",
                 "date":"Date évén.","forfait":"Forfait DT","total_paye":"Total payé DT","statut":"Statut"}
        widths = {"id":40,"num_bon":115,"client":165,"salle":90,
                  "date":90,"forfait":85,"total_paye":95,"statut":185}
        for c in cols:
            self.result_tree.heading(c, text=heads[c])
            self.result_tree.column(c, width=widths[c], anchor="center")
        rt_vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.result_tree.yview)
        self.result_tree.configure(yscrollcommand=rt_vsb.set)
        rt_vsb.pack(side="right", fill="y")
        self.result_tree.pack(fill="x")
        self.result_tree.bind("<<TreeviewSelect>>", self._on_select)

        # Detail banner
        self.detail_frame = tk.Frame(self, bg=C["stripe"], pady=0)
        self.detail_frame.pack(fill="x", padx=20)
        tk.Frame(self.detail_frame, bg=C["primary"], width=4).pack(side="left", fill="y")
        self.detail_label = tk.Label(self.detail_frame, text="Aucune sélection.",
                                     bg=C["stripe"], fg=C["text_m"], font=FNT["sm"],
                                     anchor="w", padx=12, pady=8)
        self.detail_label.pack(side="left", fill="x", expand=True)

        # ── Middle container: footer pinned bottom, form fills rest ──
        mid = tk.Frame(self, bg=C["bg"])
        mid.pack(fill="both", expand=True)

        # Footer — packed FIRST with side="bottom" to reserve space
        footer = tk.Frame(mid, bg=C["surface"])
        tk.Frame(footer, bg=C["border"], height=1).pack(fill="x")
        btn_row = tk.Frame(footer, bg=C["surface"])
        btn_row.pack(fill="x", padx=20, pady=10)
        _rbtn(btn_row, "✔  Enregistrer l'acompte", command=self._save_acompte, style="primary", padx=16, pady=7)
        _rbtn(btn_row, "📄  Ouvrir bon de reçu", command=self._open_bon, style="secondary", padx=16, pady=7)
        self.status_label = tk.Label(footer, text="", bg=C["surface"],
                                     fg=C["danger"], font=FNT["sm"])
        self.status_label.pack(padx=20, pady=(0, 8), anchor="w")
        footer.pack(fill="x", side="bottom")

        # Payment form card — scrollable area above footer
        outer_form, sf_form = _scrollable(mid)
        outer_form.pack(fill="both", expand=True)

        card = tk.Frame(sf_form, bg=C["surface"])
        card.pack(fill="x", padx=20, pady=12)
        tk.Frame(card, bg=C["accent"], height=2).pack(fill="x")
        tk.Label(card, text="Nouvel acompte", bg=C["card_hdr"],
                 fg=C["primary"], font=FNT["bold"]).pack(
            fill="x", anchor="w", padx=16, pady=10)
        tk.Frame(card, bg=C["border"], height=1).pack(fill="x")

        form = tk.Frame(card, bg=C["surface"])
        form.pack(fill="x", padx=16, pady=14)
        for c in range(6):
            form.columnconfigure(c, weight=1)

        # Row 0
        tk.Label(form, text="Montant (DT)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w", padx=(0, 8))
        self.montant_entry = RoundedEntry(form, width=14)
        self.montant_entry.grid(row=1, column=0, sticky="ew", padx=(0, 16), ipady=3)

        tk.Label(form, text="Méthode", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=1, sticky="w")
        self.methode_var = tk.StringVar(value="Espèce")
        mp = ttk.Combobox(form, textvariable=self.methode_var,
                          values=["Espèce", "Chèque", "Virement", "Effet"],
                          state="readonly", width=14)
        mp.grid(row=1, column=1, sticky="ew", padx=(0, 16), ipady=3)
        mp.bind("<<ComboboxSelected>>", lambda e: self._toggle_ref_fields())

        tk.Label(form, text="Caisse Tawfik", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=2, sticky="w")
        self.num_caisse_entry = RoundedEntry(form, width=12)
        self.num_caisse_entry.grid(row=1, column=2, sticky="ew", padx=(0, 16), ipady=3)

        tk.Label(form, text="Caisse Omar", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=3, sticky="w")
        self.num_caisse_omar_entry = RoundedEntry(form, width=12)
        self.num_caisse_omar_entry.grid(row=1, column=3, sticky="ew", padx=(0, 16), ipady=3)

        # Ref fields (hidden by default)
        self.ref_frame = tk.Frame(form, bg=C["surface"])
        self.ref_frame.grid(row=2, column=0, columnspan=6, sticky="w", pady=(8, 0))
        self.ref_frame.grid_remove()

        tk.Label(self.ref_frame, text="N° Réf", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(side="left", padx=(0, 4))
        self.num_ref_entry = RoundedEntry(self.ref_frame, width=16)
        self.num_ref_entry.pack(side="left", padx=(0, 16))
        tk.Label(self.ref_frame, text="Banque", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(side="left", padx=(0, 4))
        self.banque_entry = RoundedEntry(self.ref_frame, width=16)
        self.banque_entry.pack(side="left", padx=(0, 16))
        tk.Label(self.ref_frame, text="Date (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(side="left", padx=(0, 4))
        self.date_ref_entry = RoundedEntry(self.ref_frame, width=13)
        self.date_ref_entry.pack(side="left")

    def _toggle_ref_fields(self):
        if self.methode_var.get() in ("Chèque", "Virement", "Effet"):
            self.ref_frame.grid()
        else:
            self.ref_frame.grid_remove()

    def _populate_results(self, rows):
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        for r in rows:
            total = sum(float(r.get(f"acompte{i}") or 0) for i in range(1, 4))
            statut = r.get("statut_paiement") or ""
            self.result_tree.insert("", "end", iid=str(r["id"]), values=(
                r["id"], r.get("num_bon") or "", r.get("nom_client") or "",
                r.get("salle") or "", r.get("date_evenement") or "",
                f"{r.get('forfait') or 0:.3f}", f"{total:.3f}", statut,
            ))

    def _search(self):
        q = self.search_var.get().strip()
        rows = database.search_reservations(q) if q else database.get_all_reservations()
        self._populate_results(rows)
        self.status_label.configure(text=f"{len(rows)} résultat(s).", fg=C["primary"])

    def _show_all(self):
        self._populate_results(database.get_all_reservations())
        self.status_label.configure(text="", fg=C["danger"])

    def _on_select(self, event=None):
        sel = self.result_tree.selection()
        if not sel: return
        self.selected_id = int(sel[0])
        self._reload_selected()

    def _reload_selected(self):
        if not self.selected_id: return
        res = database.get_reservation(self.selected_id)
        if not res: return
        forfait = float(res.get("forfait") or 0)
        total   = sum(float(res.get(f"acompte{i}") or 0) for i in range(1, 4))
        reste   = forfait - total
        statut  = res.get("statut_paiement") or "En cours"
        self.detail_label.configure(
            text=(f"  {res.get('nom_client','')}   |   {res.get('salle','')}   |   "
                  f"{res.get('date_evenement','')}   |   Forfait : {forfait:.3f} DT   |   "
                  f"Payé : {total:.3f} DT   |   Reste : {reste:.3f} DT   |   {statut}"),
            fg=C["text"])

    def _save_acompte(self):
        self.status_label.configure(text="", fg=C["danger"])
        if not self.selected_id:
            self.status_label.configure(text="Sélectionne d'abord une réservation.")
            return
        try:
            montant = float(self.montant_entry.get().strip())
            if montant <= 0: raise ValueError("Montant positif requis.")
        except ValueError as e:
            self.status_label.configure(text=f"Montant invalide : {e}")
            return
        methode = self.methode_var.get()
        num_ref = banque = date_ref = ""
        if methode in ("Chèque", "Virement", "Effet"):
            num_ref = self.num_ref_entry.get().strip()
            banque  = self.banque_entry.get().strip()
            date_ref= self.date_ref_entry.get().strip()
            if not num_ref or not banque or not date_ref:
                self.status_label.configure(
                    text="N° réf, banque et date sont obligatoires.")
                return
        try:
            updated = database.add_acompte(
                self.selected_id, montant, methode, num_ref, banque, date_ref,
                num_caisse=self.num_caisse_entry.get().strip(),
                num_caisse_omar=self.num_caisse_omar_entry.get().strip())
        except ValueError as e:
            self.status_label.configure(text=str(e))
            return
        self._reload_selected()
        self._search()
        self.app.refresh_list()
        statut = updated.get("statut_paiement") or ""
        msg = f"Acompte de {montant:.3f} DT enregistré."
        if statut.startswith("SOLDE"):
            msg += f"\n✔ Réservation soldée ({statut})."
        self.status_label.configure(text=msg, fg=C["success"])
        toast(msg, "success", "Acompte enregistré")

    def _open_bon(self):
        if not self.selected_id:
            messagebox.showwarning("Sélection requise", "Sélectionne d'abord une réservation.")
            return
        res = database.get_reservation(self.selected_id)
        if not res:
            messagebox.showerror("Erreur", "Réservation introuvable.")
            return
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        output_path = os.path.join(OUTPUT_DIR,
            f"Bon_Reservation_{res.get('num_bon') or self.selected_id}.docx".replace("/", "-"))
        try:
            contract_generator.generate_contract(res, output_path)
            os.startfile(output_path)
        except Exception as e:
            messagebox.showerror("Erreur", str(e))


# ═══════════════════════════════════════════════════════════════════
#  TAB — CHARGES
# ═══════════════════════════════════════════════════════════════════
class ChargesTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._build()
        self._refresh_list()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        h_in = tk.Frame(hdr, bg=C["surface"])
        h_in.pack(fill="x", padx=20, pady=14)
        tk.Label(h_in, text="Charges", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(side="left")
        self.total_label = tk.Label(h_in, text="Total : 0.000 DT",
                                    bg=C["surface"], fg=C["danger"],
                                    font=FNT["bold"])
        self.total_label.pack(side="right")
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Toolbar
        tb = tk.Frame(self, bg=C["bg"])
        tb.pack(fill="x", padx=20, pady=10)
        _rbtn(tb, "🔄  Actualiser", command=self._refresh_list, style="secondary", padx=16, pady=7)
        _rbtn(tb, "🗑  Supprimer la sélection", command=self._delete_selected, style="sm_danger", padx=10, pady=5)

        # Tree
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        cols = ("id", "date", "designation", "montant")
        self.tree = ttk.Treeview(tree_wrap, columns=cols, show="headings")
        heads = {"id":"ID","date":"Date d'encaissement","designation":"Désignation","montant":"Montant (DT)"}
        widths = {"id":45,"date":140,"designation":0,"montant":130}
        for c in cols:
            self.tree.heading(c, text=heads[c])
            w = widths[c]
            self.tree.column(c, width=w if w else 400, minwidth=40,
                             stretch=True, anchor="w" if c == "designation" else "center")
        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)

        # Add form card
        add_wrap = tk.Frame(self, bg=C["bg"])
        add_wrap.pack(fill="x", padx=20, pady=(0, 16))
        card = tk.Frame(add_wrap, bg=C["surface"])
        card.pack(fill="x")
        tk.Frame(card, bg=C["danger"], height=2).pack(fill="x")
        tk.Label(card, text="Ajouter une charge", bg=C["card_hdr"],
                 fg=C["primary"], font=FNT["bold"]).pack(
            fill="x", anchor="w", padx=16, pady=10)
        tk.Frame(card, bg=C["border"], height=1).pack(fill="x")

        form = tk.Frame(card, bg=C["surface"])
        form.pack(fill="x", padx=16, pady=14)
        for c in range(7):
            form.columnconfigure(c, weight=1 if c % 2 else 0)

        tk.Label(form, text="Date (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.date_entry = RoundedEntry(form, width=14)
        self.date_entry.insert(0, date.today().strftime("%d/%m/%Y"))
        self.date_entry.grid(row=1, column=0, sticky="ew", padx=(0, 16), ipady=3)

        tk.Label(form, text="Désignation", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=1, sticky="w")
        self.desig_entry = RoundedEntry(form, width=44)
        self.desig_entry.grid(row=1, column=1, columnspan=3, sticky="ew",
                              padx=(0, 16), ipady=3)

        tk.Label(form, text="Montant (DT)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=4, sticky="w")
        self.montant_entry = RoundedEntry(form, width=14)
        self.montant_entry.grid(row=1, column=4, sticky="ew", padx=(0, 16), ipady=3)

        save_btn = RoundedButton(form, text="✔  Enregistrer", command=self._save,
                                 style="primary", padx=16, pady=7)
        import tkinter.font as _tf4; _f4 = _tf4.Font(font=FNT["bold"])
        save_btn.configure(width=_f4.measure("✔  Enregistrer")+34,
                           height=_f4.metrics("linespace")+16)
        save_btn.grid(row=1, column=6, sticky="w", padx=(8, 0))

        self.status_label = tk.Label(card, text="", bg=C["surface"],
                                     fg=C["danger"], font=FNT["sm"])
        self.status_label.pack(padx=16, pady=(0, 10), anchor="w")

    @staticmethod
    def _to_iso(ddmmyyyy):
        d, m, y = ddmmyyyy.strip().split("/")
        return f"{y}-{m}-{d}"

    def _refresh_list(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        charges = database.get_all_charges()
        total = 0.0
        for c in charges:
            total += float(c.get("montant") or 0)
            d_fmt = c["date_encaissement"]
            try:
                from datetime import datetime
                d_fmt = datetime.strptime(d_fmt, "%Y-%m-%d").strftime("%d/%m/%Y")
            except Exception:
                pass
            self.tree.insert("", "end", iid=str(c["id"]), values=(
                c["id"], d_fmt, c["designation"], f"{float(c['montant']):.3f}"))
        self.total_label.configure(text=f"Total : {total:.3f} DT")

    def _delete_selected(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Sélection requise", "Sélectionne une charge.")
            return
        if messagebox.askyesno("Confirmer", f"Supprimer {len(sel)} charge(s) ?"):
            for iid in sel:
                database.delete_charge(int(iid))
            self._refresh_list()

    def _save(self):
        self.status_label.configure(text="", fg=C["danger"])
        try:
            iso_date = self._to_iso(self.date_entry.get())
        except Exception:
            self.status_label.configure(text="Date invalide (JJ/MM/AAAA).")
            return
        desig = self.desig_entry.get().strip()
        if not desig:
            self.status_label.configure(text="Désignation obligatoire.")
            return
        try:
            montant = float(self.montant_entry.get().strip())
            if montant <= 0: raise ValueError
        except ValueError:
            self.status_label.configure(text="Montant invalide.")
            return
        database.insert_charge(iso_date, desig, montant)
        self.status_label.configure(
            text=f"Charge de {montant:.3f} DT enregistrée.", fg=C["success"])
        self.desig_entry.delete(0, "end")
        self.montant_entry.delete(0, "end")
        self._refresh_list()


# ═══════════════════════════════════════════════════════════════════
#  TAB — CHARGES OMAR
# ═══════════════════════════════════════════════════════════════════
class ChargesOmarTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._build()
        self._refresh_list()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        h_in = tk.Frame(hdr, bg=C["surface"])
        h_in.pack(fill="x", padx=20, pady=14)
        tk.Label(h_in, text="Charges Omar", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(side="left")
        self.total_label = tk.Label(h_in, text="Total : 0.000 DT",
                                    bg=C["surface"], fg=C["danger"],
                                    font=FNT["bold"])
        self.total_label.pack(side="right")
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Toolbar
        tb = tk.Frame(self, bg=C["bg"])
        tb.pack(fill="x", padx=20, pady=10)
        _rbtn(tb, "Actualiser", command=self._refresh_list, style="secondary", padx=16, pady=7)
        _rbtn(tb, "Supprimer la selection", command=self._delete_selected, style="sm_danger", padx=10, pady=5)

        # Tree
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        cols = ("id", "date", "designation", "montant")
        self.tree = ttk.Treeview(tree_wrap, columns=cols, show="headings")
        heads = {"id": "ID", "date": "Date d'encaissement",
                 "designation": "Designation", "montant": "Montant (DT)"}
        widths = {"id": 45, "date": 140, "designation": 0, "montant": 130}
        for c in cols:
            self.tree.heading(c, text=heads[c])
            w = widths[c]
            self.tree.column(c, width=w if w else 400, minwidth=40,
                             stretch=True,
                             anchor="w" if c == "designation" else "center")
        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)

        # Add-form card
        add_wrap = tk.Frame(self, bg=C["bg"])
        add_wrap.pack(fill="x", padx=20, pady=(0, 16))
        card = tk.Frame(add_wrap, bg=C["surface"])
        card.pack(fill="x")
        tk.Frame(card, bg=C["danger"], height=2).pack(fill="x")
        tk.Label(card, text="Ajouter une charge Omar", bg=C["card_hdr"],
                 fg=C["primary"], font=FNT["bold"]).pack(
            fill="x", anchor="w", padx=16, pady=10)
        tk.Frame(card, bg=C["border"], height=1).pack(fill="x")

        form = tk.Frame(card, bg=C["surface"])
        form.pack(fill="x", padx=16, pady=14)
        for c in range(7):
            form.columnconfigure(c, weight=1 if c % 2 else 0)

        tk.Label(form, text="Date (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.date_entry = RoundedEntry(form, width=14)
        self.date_entry.insert(0, date.today().strftime("%d/%m/%Y"))
        self.date_entry.grid(row=1, column=0, sticky="ew", padx=(0, 16), ipady=3)

        tk.Label(form, text="Designation", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=1, sticky="w")
        self.desig_entry = RoundedEntry(form, width=44)
        self.desig_entry.grid(row=1, column=1, columnspan=3, sticky="ew",
                              padx=(0, 16), ipady=3)

        tk.Label(form, text="Montant (DT)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=4, sticky="w")
        self.montant_entry = RoundedEntry(form, width=14)
        self.montant_entry.grid(row=1, column=4, sticky="ew", padx=(0, 16), ipady=3)

        save_btn = RoundedButton(form, text="Enregistrer", command=self._save,
                                 style="primary", padx=16, pady=7)
        import tkinter.font as _tf4o
        _f4o = _tf4o.Font(font=FNT["bold"])
        save_btn.configure(width=_f4o.measure("Enregistrer") + 34,
                           height=_f4o.metrics("linespace") + 16)
        save_btn.grid(row=1, column=6, sticky="w", padx=(8, 0))

        self.status_label = tk.Label(card, text="", bg=C["surface"],
                                     fg=C["danger"], font=FNT["sm"])
        self.status_label.pack(padx=16, pady=(0, 10), anchor="w")

    @staticmethod
    def _to_iso(ddmmyyyy):
        d, m, y = ddmmyyyy.strip().split("/")
        return f"{y}-{m}-{d}"

    def _refresh_list(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        charges = database.get_all_charges_omar()
        total = 0.0
        for c in charges:
            total += float(c.get("montant") or 0)
            d_fmt = c["date_encaissement"]
            try:
                from datetime import datetime
                d_fmt = datetime.strptime(d_fmt, "%Y-%m-%d").strftime("%d/%m/%Y")
            except Exception:
                pass
            self.tree.insert("", "end", iid=str(c["id"]), values=(
                c["id"], d_fmt, c["designation"],
                f"{float(c['montant']):.3f}"))
        self.total_label.configure(text=f"Total : {total:.3f} DT")

    def _delete_selected(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Selection requise", "Selectionnez une charge.")
            return
        if messagebox.askyesno("Confirmer", f"Supprimer {len(sel)} charge(s) ?"):
            for iid in sel:
                database.delete_charge_omar(int(iid))
            self._refresh_list()

    def _save(self):
        self.status_label.configure(text="", fg=C["danger"])
        try:
            iso_date = self._to_iso(self.date_entry.get())
        except Exception:
            self.status_label.configure(text="Date invalide (JJ/MM/AAAA).")
            return
        desig = self.desig_entry.get().strip()
        if not desig:
            self.status_label.configure(text="Designation obligatoire.")
            return
        try:
            montant = float(self.montant_entry.get().strip())
            if montant <= 0:
                raise ValueError
        except ValueError:
            self.status_label.configure(text="Montant invalide.")
            return
        database.insert_charge_omar(iso_date, desig, montant)
        self.status_label.configure(
            text=f"Charge de {montant:.3f} DT enregistree.", fg=C["success"])
        self.desig_entry.delete(0, "end")
        self.montant_entry.delete(0, "end")
        self._refresh_list()


# ═══════════════════════════════════════════════════════════════════
#  TAB — JOURNAL DE CAISSE
# ═══════════════════════════════════════════════════════════════════
class JournalCaisseTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._last_data = None
        self._build()
        # Pre-fill with current month range
        import calendar as _cal
        today = date.today()
        last_day = _cal.monthrange(today.year, today.month)[1]
        self.date_debut_entry.insert(0, f"01/{today.month:02d}/{today.year}")
        self.date_fin_entry.insert(0, f"{last_day:02d}/{today.month:02d}/{today.year}")

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        tk.Label(hdr, text="Journal de Caisse", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Period card
        period_wrap = tk.Frame(self, bg=C["bg"])
        period_wrap.pack(fill="x", padx=20, pady=12)
        card = tk.Frame(period_wrap, bg=C["surface"])
        card.pack(fill="x")
        tk.Frame(card, bg=C["primary"], height=2).pack(fill="x")
        form = tk.Frame(card, bg=C["surface"])
        form.pack(fill="x", padx=16, pady=14)
        for c in (1, 3, 5):
            form.columnconfigure(c, weight=1)

        tk.Label(form, text="N° Journal", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.num_journal_entry = RoundedEntry(form, width=8)
        self.num_journal_entry.grid(row=0, column=1, sticky="w", padx=(0, 20), ipady=3)

        tk.Label(form, text="Du (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=2, sticky="w")
        self.date_debut_entry = RoundedEntry(form, width=13)
        self.date_debut_entry.grid(row=0, column=3, sticky="w", padx=(0, 20), ipady=3)

        tk.Label(form, text="Au (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=4, sticky="w")
        self.date_fin_entry = RoundedEntry(form, width=13)
        self.date_fin_entry.grid(row=0, column=5, sticky="w", padx=(0, 20), ipady=3)

        btn_row = tk.Frame(card, bg=C["surface"])
        btn_row.pack(fill="x", padx=16, pady=(0, 14))
        _rbtn(btn_row, "📋  Générer le journal", command=self._generate, style="primary", padx=16, pady=7)
        _rbtn(btn_row, "🖨  Exporter Excel", command=self._export_excel, style="success", padx=16, pady=7)

        self.header_label = tk.Label(self, text="", bg=C["bg"],
                                     fg=C["primary"], font=FNT["bold"])
        self.header_label.pack(padx=20, pady=(4, 0), anchor="w")

        # Tree
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=8)
        cols = ("date", "designation", "fn", "num_caisse", "recette", "charge")
        self.tree = ttk.Treeview(tree_wrap, columns=cols, show="headings")
        heads = {"date":"Date","designation":"Désignation","fn":"FN°",
                 "num_caisse":"N° Caisse","recette":"Recette (DT)","charge":"Charge (DT)"}
        widths = {"date":95,"designation":0,"fn":85,"num_caisse":85,"recette":120,"charge":120}
        for c in cols:
            self.tree.heading(c, text=heads[c])
            w = widths[c]
            self.tree.column(c, width=w if w else 450, stretch=True,
                             anchor="e" if c in ("recette","charge") else "w")
        self.tree.tag_configure("recette", background=C["tag_rec"])
        self.tree.tag_configure("charge",  background=C["tag_chg"])
        self.tree.tag_configure("total",   background=C["tag_tot"],
                                font=FNT["bold"])
        self.tree.tag_configure("solde",   background=C["tag_sld"],
                                font=("Segoe UI", 11, "bold"))
        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)

        # Footer totals
        footer = tk.Frame(self, bg=C["surface"])
        footer.pack(fill="x")
        tk.Frame(footer, bg=C["border"], height=1).pack(fill="x")
        f_in = tk.Frame(footer, bg=C["surface"])
        f_in.pack(fill="x", padx=20, pady=12)
        self.lbl_total_recettes = tk.Label(f_in, text="Recettes : —",
                                           bg=C["surface"], fg=C["success"],
                                           font=FNT["bold"])
        self.lbl_total_recettes.pack(side="left", padx=(0, 30))
        self.lbl_total_charges = tk.Label(f_in, text="Charges : —",
                                          bg=C["surface"], fg=C["danger"],
                                          font=FNT["bold"])
        self.lbl_total_charges.pack(side="left", padx=(0, 30))
        self.lbl_solde = tk.Label(f_in, text="Solde net : —",
                                  bg=C["surface"], fg=C["primary"],
                                  font=("Segoe UI", 12, "bold"))
        self.lbl_solde.pack(side="left")

    @staticmethod
    def _to_iso(ddmmyyyy):
        d, m, y = ddmmyyyy.strip().split("/")
        return f"{y}-{m}-{d}"

    @staticmethod
    def _fmt_date(iso):
        try:
            from datetime import datetime
            return datetime.strptime(iso, "%Y-%m-%d").strftime("%d/%m/%Y")
        except Exception:
            return iso

    @staticmethod
    def _day_name(iso):
        try:
            d = date.fromisoformat(iso)
            return ["LUNDI","MARDI","MERCREDI","JEUDI","VENDREDI","SAMEDI","DIMANCHE"][d.weekday()]
        except Exception:
            return ""

    def _generate(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        try:
            iso_debut = self._to_iso(self.date_debut_entry.get())
            iso_fin   = self._to_iso(self.date_fin_entry.get())
        except Exception:
            messagebox.showerror("Erreur", "Dates invalides (JJ/MM/AAAA).")
            return
        num = self.num_journal_entry.get().strip() or "?"
        header = (f"Journal de Caisse N°{num}  —  "
                  f"{self._day_name(iso_debut)} {self._fmt_date(iso_debut)}"
                  f"  au  {self._day_name(iso_fin)} {self._fmt_date(iso_fin)}")
        self.header_label.configure(text=header)
        data = database.get_journal_caisse(iso_debut, iso_fin)
        self._last_data = {"header": header, "debut": iso_debut, "fin": iso_fin, **data}
        total_r = total_c = 0.0
        # Recettes en ordre de saisie (id ASC, slot ASC), puis charges
        for row in data["recettes"]:
            total_r += row["montant"]
            fn_d = f"FN°{row['fn']}" if row["fn"] and not row["fn"].startswith("FN") else row["fn"]
            self.tree.insert("", "end", values=(
                self._fmt_date(row["date"]), row["designation"],
                fn_d or "", row["num_caisse"] or "",
                f"{row['montant']:,.3f}", ""), tags=("recette",))
        for row in data["charges"]:
            total_c += row["montant"]
            self.tree.insert("", "end", values=(
                row["date_fmt"], row["designation"], "", "",
                "", f"{row['montant']:,.3f}"), tags=("charge",))
        solde = total_r - total_c
        self.tree.insert("", "end", values=(
            "", "", "", "TOTAUX", f"{total_r:,.3f}", f"{total_c:,.3f}"), tags=("total",))
        self.tree.insert("", "end", values=(
            "", "", "", "SOLDE NET", f"{solde:,.3f}", ""), tags=("solde",))
        self.lbl_total_recettes.configure(text=f"Recettes : {total_r:,.3f} DT")
        self.lbl_total_charges.configure(text=f"Charges : {total_c:,.3f} DT")
        fg = C["success"] if solde >= 0 else C["danger"]
        self.lbl_solde.configure(text=f"Solde net : {solde:,.3f} DT", fg=fg)

    def _export_excel(self):
        if not self._last_data:
            messagebox.showwarning("Aucune donnée", "Génère d'abord le journal.")
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            messagebox.showerror("Erreur", "openpyxl non installé.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"Journal_Caisse_{self._last_data['debut']}_{self._last_data['fin']}.xlsx")
        if not path: return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Journal de Caisse"
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.merge_cells("A1:F1")
        c = ws["A1"]
        c.value = self._last_data["header"]
        c.font = Font(bold=True, size=12, color="1A2B45")
        c.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 24
        hdrs = ["Date","Désignation","FN°","N° Caisse","Recette (DT)","Charge (DT)"]
        fill_hdr = PatternFill("solid", fgColor="1A2B45")
        for ci, h in enumerate(hdrs, 1):
            cell = ws.cell(row=2, column=ci, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill_hdr
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
        ri = 3
        for row in self._last_data["recettes"]:
            fn_d = f"FN°{row['fn']}" if row["fn"] and not row["fn"].startswith("FN") else row["fn"]
            for ci, v in enumerate([self._fmt_date(row["date"]), row["designation"],
                                    fn_d or "", row["num_caisse"] or "", row["montant"], ""], 1):
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.fill = PatternFill("solid", fgColor="ECFDF5")
                cell.border = border
                if ci >= 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
            ri += 1
        for row in self._last_data["charges"]:
            for ci, v in enumerate([row["date_fmt"], row["designation"], "", "", "", row["montant"]], 1):
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.fill = PatternFill("solid", fgColor="FFF1F2")
                cell.border = border
                if ci >= 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
            ri += 1
        tr = sum(r["montant"] for r in self._last_data["recettes"])
        tc = sum(r["montant"] for r in self._last_data["charges"])
        for ci, v in enumerate(["","","","TOTAUX", tr, tc], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill = PatternFill("solid", fgColor="EEF2FF")
            cell.font = Font(bold=True)
            cell.border = border
            if ci >= 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
        ri += 1
        for ci, v in enumerate(["","","","SOLDE NET", tr-tc, ""], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill = PatternFill("solid", fgColor="FFFBEB")
            cell.font = Font(bold=True, size=11)
            cell.border = border
            if ci == 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 58
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 10
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 16
        ws.freeze_panes = "A3"
        wb.save(path)
        toast(os.path.basename(path), "success", "Journal exporté")
        try:
            os.startfile(path)
        except Exception:
            pass


# ═══════════════════════════════════════════════════════════════════
#  TAB — JOURNAL DE CAISSE OMAR
# ═══════════════════════════════════════════════════════════════════
class JournalCaisseOmarTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._last_data = None
        self._build()
        # Pre-fill with current month range
        today = date.today()
        import calendar as _cal
        last_day = _cal.monthrange(today.year, today.month)[1]
        self.date_debut_entry.insert(0, f"01/{today.month:02d}/{today.year}")
        self.date_fin_entry.insert(0, f"{last_day:02d}/{today.month:02d}/{today.year}")

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        tk.Label(hdr, text="Journal de Caisse Omar", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Period card
        period_wrap = tk.Frame(self, bg=C["bg"])
        period_wrap.pack(fill="x", padx=20, pady=12)
        card = tk.Frame(period_wrap, bg=C["surface"])
        card.pack(fill="x")
        tk.Frame(card, bg=C["primary"], height=2).pack(fill="x")
        form = tk.Frame(card, bg=C["surface"])
        form.pack(fill="x", padx=16, pady=14)
        for c in (1, 3, 5):
            form.columnconfigure(c, weight=1)

        tk.Label(form, text="N° Journal", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.num_journal_entry = RoundedEntry(form, width=8)
        self.num_journal_entry.grid(row=0, column=1, sticky="w", padx=(0, 20), ipady=3)

        tk.Label(form, text="Du (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=2, sticky="w")
        self.date_debut_entry = RoundedEntry(form, width=13)
        self.date_debut_entry.grid(row=0, column=3, sticky="w", padx=(0, 20), ipady=3)

        tk.Label(form, text="Au (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=4, sticky="w")
        self.date_fin_entry = RoundedEntry(form, width=13)
        self.date_fin_entry.grid(row=0, column=5, sticky="w", padx=(0, 20), ipady=3)

        btn_row = tk.Frame(card, bg=C["surface"])
        btn_row.pack(fill="x", padx=16, pady=(0, 14))
        _rbtn(btn_row, "Generer le journal", command=self._generate,
              style="primary", padx=16, pady=7)
        _rbtn(btn_row, "Exporter Excel", command=self._export_excel,
              style="success", padx=16, pady=7)
        _rbtn(btn_row, "➕  Ajouter recette", command=self._add_recette_dialog,
              style="secondary", padx=16, pady=7)

        self.header_label = tk.Label(self, text="", bg=C["bg"],
                                     fg=C["primary"], font=FNT["bold"])
        self.header_label.pack(padx=20, pady=(4, 0), anchor="w")

        # Tree
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=8)

        cols = ("date", "designation", "recette", "sortie")
        self.tree = ttk.Treeview(tree_wrap, columns=cols, show="headings")
        heads = {
            "date":        "Date",
            "designation": "Designation",
            "recette":     "Recette (DT)",
            "sortie":      "Sortie (DT)",
        }
        widths = {
            "date":        95,
            "designation": 0,
            "recette":     130,
            "sortie":      130,
        }
        for c in cols:
            self.tree.heading(c, text=heads[c])
            self.tree.column(c, width=widths[c] if widths[c] else 450,
                             stretch=(widths[c] == 0),
                             anchor="e" if c in ("recette", "sortie") else "w")

        self.tree.tag_configure("recette", background=C["tag_rec"])
        self.tree.tag_configure("sortie_charge",  background=C["tag_chg"])
        self.tree.tag_configure("sortie_employe", background="#FFF0F9")
        self.tree.tag_configure("total",  background=C["tag_tot"], font=FNT["bold"])
        self.tree.tag_configure("solde",  background=C["tag_sld"],
                                font=("Segoe UI", 11, "bold"))
        self.tree.tag_configure("date_sep", background=C["sidebar"],
                                foreground="white", font=FNT["bold"])

        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)

        # Hint overlay shown before first generate
        self._hint_lbl = tk.Label(
            tree_wrap,
            text=(
                "Entrez une periode (Du / Au) et cliquez  'Generer le journal'\n\n"
                "Entrees (Recettes) :\n"
                "  - Acomptes encaisses dans la Caisse Omar (champ 'Caisse Omar' dans Ajouter un acompte)\n\n"
                "Sorties :\n"
                "  - Charges Omar (module Charges Omar)\n"
                "  - Paiements des employes de nuit (enregistres automatiquement au pointage)\n"
                "  - Paiements employes manuels (module Liste employes > Fiche)"
            ),
            bg=C["surface"], fg=C["text_m"],
            font=FNT["default"], justify="left",
            anchor="nw", padx=24, pady=24,
        )
        self._hint_lbl.place(relx=0, rely=0, relwidth=1, relheight=1)

        # Footer totals
        footer = tk.Frame(self, bg=C["surface"])
        footer.pack(fill="x")
        tk.Frame(footer, bg=C["border"], height=1).pack(fill="x")
        f_in = tk.Frame(footer, bg=C["surface"])
        f_in.pack(fill="x", padx=20, pady=12)

        self.lbl_total_recettes = tk.Label(f_in, text="Recettes : —",
                                           bg=C["surface"], fg=C["success"],
                                           font=FNT["bold"])
        self.lbl_total_recettes.pack(side="left", padx=(0, 30))
        self.lbl_total_sorties = tk.Label(f_in, text="Sorties : —",
                                          bg=C["surface"], fg=C["danger"],
                                          font=FNT["bold"])
        self.lbl_total_sorties.pack(side="left", padx=(0, 30))
        self.lbl_solde = tk.Label(f_in, text="Solde net : —",
                                  bg=C["surface"], fg=C["primary"],
                                  font=("Segoe UI", 12, "bold"))
        self.lbl_solde.pack(side="left")

    # ── Helpers ───────────────────────────────────────────────────
    @staticmethod
    def _to_iso(ddmmyyyy):
        d, m, y = ddmmyyyy.strip().split("/")
        return f"{y}-{m}-{d}"

    @staticmethod
    def _fmt_date(iso):
        try:
            from datetime import datetime
            return datetime.strptime(iso, "%Y-%m-%d").strftime("%d/%m/%Y")
        except Exception:
            return iso

    @staticmethod
    def _day_name(iso):
        try:
            d = date.fromisoformat(iso)
            return ["LUNDI","MARDI","MERCREDI","JEUDI","VENDREDI","SAMEDI","DIMANCHE"][d.weekday()]
        except Exception:
            return ""

    def _add_recette_dialog(self):
        """Ouvre une fenêtre modale pour saisir une recette manuelle."""
        dlg = tk.Toplevel(self)
        dlg.title("Ajouter une recette — Caisse Omar")
        dlg.resizable(False, False)
        dlg.grab_set()
        dlg.configure(bg=C["surface"])

        pad = {"padx": 12, "pady": 6}

        tk.Label(dlg, text="Date (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w", **pad)
        date_entry = RoundedEntry(dlg, width=14)
        date_entry.insert(0, date.today().strftime("%d/%m/%Y"))
        date_entry.grid(row=0, column=1, sticky="ew", **pad)

        tk.Label(dlg, text="Désignation", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=1, column=0, sticky="w", **pad)
        desig_entry = RoundedEntry(dlg, width=34)
        desig_entry.grid(row=1, column=1, sticky="ew", **pad)

        tk.Label(dlg, text="Montant (DT)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=2, column=0, sticky="w", **pad)
        montant_entry = RoundedEntry(dlg, width=14)
        montant_entry.grid(row=2, column=1, sticky="ew", **pad)

        status_lbl = tk.Label(dlg, text="", bg=C["surface"],
                              fg=C["danger"], font=FNT["sm"])
        status_lbl.grid(row=3, column=0, columnspan=2, sticky="w", padx=12)

        def _save():
            # Validate date
            try:
                d_raw = date_entry.get().strip()
                dd, mm, yy = d_raw.split("/")
                iso_date = f"{yy}-{mm}-{dd}"
                date.fromisoformat(iso_date)          # raises if invalid
            except Exception:
                status_lbl.configure(text="Date invalide (JJ/MM/AAAA).")
                return
            desig = desig_entry.get().strip()
            if not desig:
                status_lbl.configure(text="Désignation obligatoire.")
                return
            try:
                montant = float(montant_entry.get().strip().replace(",", "."))
                if montant <= 0:
                    raise ValueError
            except ValueError:
                status_lbl.configure(text="Montant invalide (nombre positif).")
                return

            database.insert_recette_omar_extra(iso_date, desig, montant)
            dlg.destroy()
            # Refresh journal if already generated
            if self._last_data:
                self._generate()

        btn_frame = tk.Frame(dlg, bg=C["surface"])
        btn_frame.grid(row=4, column=0, columnspan=2, pady=(4, 12))
        _rbtn(btn_frame, "✔  Enregistrer", command=_save,
              style="primary", padx=14, pady=6)
        _rbtn(btn_frame, "Annuler", command=dlg.destroy,
              style="secondary", padx=14, pady=6)

        dlg.columnconfigure(1, weight=1)
        dlg.update_idletasks()
        # Centre on parent
        x = self.winfo_rootx() + (self.winfo_width() - dlg.winfo_width()) // 2
        y = self.winfo_rooty() + (self.winfo_height() - dlg.winfo_height()) // 2
        dlg.geometry(f"+{x}+{y}")

    def _generate(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        # Hide the hint overlay once the user generates
        try:
            self._hint_lbl.place_forget()
        except Exception:
            pass
        try:
            iso_debut = self._to_iso(self.date_debut_entry.get())
            iso_fin   = self._to_iso(self.date_fin_entry.get())
        except Exception:
            messagebox.showerror("Erreur", "Dates invalides (JJ/MM/AAAA).")
            return

        num = self.num_journal_entry.get().strip() or "?"
        header = (
            f"Journal de Caisse Omar N°{num}  —  "
            f"{self._day_name(iso_debut)} {self._fmt_date(iso_debut)}"
            f"  au  {self._day_name(iso_fin)} {self._fmt_date(iso_fin)}"
        )
        self.header_label.configure(text=header)

        data = database.get_journal_caisse_omar(iso_debut, iso_fin)
        self._last_data = {
            "header": header,
            "debut":  iso_debut,
            "fin":    iso_fin,
            **data,
        }

        total_r = total_s = 0.0

        # Merge recettes + sorties and sort by date (ISO key for correct ordering)
        all_rows = []
        for row in data["recettes"]:
            all_rows.append(("recette", row["date"], row))
        for row in data["sorties"]:
            all_rows.append(("sortie", row["date"], row))
        all_rows.sort(key=lambda x: x[1])  # sort by ISO date ascending

        current_date = None
        for kind, _date_iso, row in all_rows:
            # Insert a date separator row when the date changes
            if _date_iso != current_date:
                current_date = _date_iso
                day_lbl = f"── {self._day_name(_date_iso)}  {self._fmt_date(_date_iso)} ──"
                self.tree.insert("", "end", tags=("date_sep",), values=(
                    "", day_lbl, "", "",
                ))
            if kind == "recette":
                total_r += row["montant"]
                self.tree.insert("", "end", tags=("recette",), values=(
                    self._fmt_date(row["date"]),
                    row["designation"],
                    f"{row['montant']:,.3f}",
                    "",
                ))
            else:
                total_s += row["montant"]
                tag = "sortie_employe" if row["type"] == "employe" else "sortie_charge"
                self.tree.insert("", "end", tags=(tag,), values=(
                    row["date_fmt"],
                    row["designation"],
                    "",
                    f"{row['montant']:,.3f}",
                ))

        solde = total_r - total_s

        self.tree.insert("", "end", tags=("total",), values=(
            "", "TOTAUX",
            f"{total_r:,.3f}", f"{total_s:,.3f}",
        ))
        self.tree.insert("", "end", tags=("solde",), values=(
            "", "SOLDE NET",
            f"{solde:,.3f}", "",
        ))

        self.lbl_total_recettes.configure(text=f"Recettes : {total_r:,.3f} DT")
        self.lbl_total_sorties.configure(text=f"Sorties  : {total_s:,.3f} DT")
        fg = C["success"] if solde >= 0 else C["danger"]
        self.lbl_solde.configure(text=f"Solde net : {solde:,.3f} DT", fg=fg)

    def _export_excel(self):
        if not self._last_data:
            messagebox.showwarning("Aucune donnee", "Generez d'abord le journal.")
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            messagebox.showerror("Erreur", "openpyxl non installe.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=(
                f"Journal_Caisse_Omar_"
                f"{self._last_data['debut']}_{self._last_data['fin']}.xlsx"
            ),
        )
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Journal Caisse Omar"

        thin   = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Title row
        ws.merge_cells("A1:D1")
        c = ws["A1"]
        c.value = self._last_data["header"]
        c.font  = Font(bold=True, size=12, color="1A2B45")
        c.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 24

        # Header row
        hdrs = ["Date", "Designation", "Recette (DT)", "Sortie (DT)"]
        fill_hdr = PatternFill("solid", fgColor="1A2B45")
        for ci, h in enumerate(hdrs, 1):
            cell = ws.cell(row=2, column=ci, value=h)
            cell.font      = Font(bold=True, color="FFFFFF")
            cell.fill      = fill_hdr
            cell.alignment = Alignment(horizontal="center")
            cell.border    = border

        ri = 3
        # Recettes
        for row in self._last_data["recettes"]:
            for ci, v in enumerate([
                self._fmt_date(row["date"]),
                row["designation"],
                row["montant"],
                "",
            ], 1):
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.fill   = PatternFill("solid", fgColor="ECFDF5")
                cell.border = border
                if ci >= 3:
                    cell.number_format = "#,##0.000"
                    cell.alignment = Alignment(horizontal="right")
            ri += 1

        # Sorties
        for row in self._last_data["sorties"]:
            fg_color = "FFF1F2" if row["type"] == "charge" else "FFF0F9"
            for ci, v in enumerate([
                row["date_fmt"],
                row["designation"],
                "",
                row["montant"],
            ], 1):
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.fill   = PatternFill("solid", fgColor=fg_color)
                cell.border = border
                if ci >= 3:
                    cell.number_format = "#,##0.000"
                    cell.alignment = Alignment(horizontal="right")
            ri += 1

        # Totals
        tr = sum(r["montant"] for r in self._last_data["recettes"])
        ts = sum(r["montant"] for r in self._last_data["sorties"])
        for ci, v in enumerate(["", "TOTAUX", tr, ts], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill   = PatternFill("solid", fgColor="EEF2FF")
            cell.font   = Font(bold=True)
            cell.border = border
            if ci >= 3:
                cell.number_format = "#,##0.000"
                cell.alignment = Alignment(horizontal="right")
        ri += 1
        for ci, v in enumerate(["", "SOLDE NET", tr - ts, ""], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill   = PatternFill("solid", fgColor="FFFBEB")
            cell.font   = Font(bold=True, size=11)
            cell.border = border
            if ci == 3:
                cell.number_format = "#,##0.000"
                cell.alignment = Alignment(horizontal="right")

        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 60
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 16
        ws.freeze_panes = "A3"

        wb.save(path)
        toast(os.path.basename(path), "success", "Journal exporté")
        try:
            os.startfile(path)
        except Exception:
            pass


# ═══════════════════════════════════════════════════════════════════
#  TAB — MODIFIER RÉSERVATION
# ═══════════════════════════════════════════════════════════════════
class ModifierReservationTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._res_id = None
        self._build()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        tk.Label(hdr, text="Modifier une réservation", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Search
        sf = tk.Frame(self, bg=C["surface"])
        sf.pack(fill="x")
        sf_in = tk.Frame(sf, bg=C["surface"])
        sf_in.pack(fill="x", padx=20, pady=12)
        tk.Label(sf_in, text="🔍", bg=C["surface"], fg=C["text_m"],
                 font=("Segoe UI", 12)).pack(side="left", padx=(0, 6))
        self.search_var = tk.StringVar()
        e = RoundedEntry(sf_in, textvariable=self.search_var, width=36)
        e.pack(side="left", ipady=3)
        e.bind("<Return>", lambda _: self._search())
        _rbtn(sf_in, "Rechercher", command=self._search, style="sm", padx=10, pady=5)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Results tree
        rt_wrap = tk.Frame(self, bg=C["bg"])
        rt_wrap.pack(fill="x", padx=20, pady=10)
        cols = ("id", "num_bon", "client", "salle", "date", "heure", "forfait")
        self.result_tree = ttk.Treeview(rt_wrap, columns=cols, show="headings", height=4)
        heads = {"id":"ID","num_bon":"N° Bon","client":"Client",
                 "salle":"Salle","date":"Date","heure":"Créneau","forfait":"Forfait"}
        widths = {"id":40,"num_bon":115,"client":185,"salle":90,
                  "date":90,"heure":115,"forfait":85}
        for c in cols:
            self.result_tree.heading(c, text=heads[c])
            self.result_tree.column(c, width=widths[c], anchor="center")
        rt_vsb = ttk.Scrollbar(rt_wrap, orient="vertical", command=self.result_tree.yview)
        self.result_tree.configure(yscrollcommand=rt_vsb.set)
        rt_vsb.pack(side="right", fill="y")
        self.result_tree.pack(fill="x")
        self.result_tree.bind("<<TreeviewSelect>>", self._on_select)

        # Selected banner
        self.select_lbl = tk.Label(self, text="Sélectionne une réservation ci-dessus.",
                                   bg=C["stripe"], fg=C["text_m"], font=FNT["sm"],
                                   anchor="w", padx=20, pady=8)
        self.select_lbl.pack(fill="x", padx=20)

        # ── Middle container: footer pinned at bottom, scrollable fills rest ──
        mid = tk.Frame(self, bg=C["bg"])
        mid.pack(fill="both", expand=True)

        # Footer — packed first with side="bottom" so it reserves space before scrollable
        footer = tk.Frame(mid, bg=C["surface"])
        tk.Frame(footer, bg=C["border"], height=1).pack(fill="x")
        btn_bar_outer = tk.Frame(footer, bg=C["surface"])
        btn_bar_outer.pack(fill="x", padx=20, pady=10)
        _rbtn(btn_bar_outer, "💾  Enregistrer", command=self._save, style="primary", padx=16, pady=7)
        _rbtn(btn_bar_outer, "📄  Contrat (FR)", command=self._generate, style="secondary", padx=16, pady=7)
        _rbtn(btn_bar_outer, "📜  Contrat (عربي)", command=self._generate_arabic, style="success", padx=16, pady=7)
        _rbtn(btn_bar_outer, "🔄 Actualiser infos arabe", command=self._m_refresh_arabic_info, style="secondary", padx=16, pady=7)
        self.status_lbl = tk.Label(footer, text="", bg=C["surface"],
                                   fg=C["danger"], font=FNT["sm"])
        self.status_lbl.pack(padx=20, pady=(0, 8), anchor="w")
        footer.pack(fill="x", side="bottom")

        # Scrollable form — fills remaining space above footer
        outer, sf2 = _scrollable(mid)
        outer.pack(fill="both", expand=True)

        # ── Form card ──────────────────────────────────────────────
        def _sec(title):
            wr = tk.Frame(sf2, bg=C["bg"])
            wr.pack(fill="x", padx=20, pady=(0, 10))
            card = tk.Frame(wr, bg=C["surface"])
            card.pack(fill="x")
            tk.Frame(card, bg=C["primary"], height=2).pack(fill="x")
            tk.Label(card, text=title, bg=C["card_hdr"],
                     fg=C["primary"], font=FNT["bold"]).pack(
                fill="x", anchor="w", padx=16, pady=8)
            tk.Frame(card, bg=C["border"], height=1).pack(fill="x")
            body = tk.Frame(card, bg=C["surface"])
            body.pack(fill="x", padx=16, pady=12)
            for cc in range(4):
                body.columnconfigure(cc, weight=1)
            return body

        g1 = _sec("Créneau & Salle")
        tk.Label(g1, text="Salle", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.m_salle = tk.StringVar(value="Almes")
        ttk.Combobox(g1, textvariable=self.m_salle, values=rules.SALLES,
                     state="readonly").grid(row=1, column=0, sticky="ew",
                                            padx=(0, 16), ipady=2)
        self.m_date_evt   = _field(g1, 1, 1, "Date (JJ/MM/AAAA)", bg=C["surface"])
        self.m_heure_debut= _field(g1, 1, 2, "Heure début", bg=C["surface"])
        self.m_heure_fin  = _field(g1, 1, 3, "Heure fin", bg=C["surface"])

        g2 = _sec("Client")
        self.m_nom      = _field(g2, 0, 0, "Nom et prénom", colspan=2, bg=C["surface"])
        self.m_tel1     = _field(g2, 0, 2, "Tél 1", bg=C["surface"])
        self.m_tel2     = _field(g2, 0, 3, "Tél 2", bg=C["surface"])
        self.m_cin      = _field(g2, 1, 0, "CIN", bg=C["surface"])
        self.m_date_cin = _field(g2, 1, 1, "Date CIN (JJ/MM/AAAA)", bg=C["surface"])
        self.m_lieu_cin = _field(g2, 1, 2, "Lieu CIN", bg=C["surface"])
        self.m_nb_places= _field(g2, 2, 0, "Nb places", bg=C["surface"])
        self.m_type_evt = _field(g2, 2, 1, "Type d'événement", colspan=2, bg=C["surface"])

        g3 = _sec("Contrat & Paiement")
        self.m_forfait      = _field(g3, 0, 0, "Forfait (DT)", bg=C["surface"])
        self.m_n_facture    = _field(g3, 0, 1, "N° Facture", bg=C["surface"])
        self.m_date_contrat = _field(g3, 0, 2, "Date contrat (JJ/MM/AAAA)",
                                     colspan=2, bg=C["surface"])
        tk.Label(g3, text="Méthode paiement", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=2, column=0, sticky="w",
                                                       pady=(8, 0))
        self.m_methode = tk.StringVar(value="Espèce")
        ttk.Combobox(g3, textvariable=self.m_methode,
                     values=["Espèce","Chèque","Virement","Effet"],
                     state="readonly").grid(row=3, column=0, sticky="ew",
                                            padx=(0, 16), ipady=2)
        self.m_caisse    = _field(g3, 2, 1, "Caisse Tawfik", bg=C["surface"])
        self.m_tresorerie= _field(g3, 2, 2, "Caisse Omar", bg=C["surface"])
        tk.Label(g3, text="Détails extra", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=4, column=0, sticky="w",
                                                       pady=(8, 0))
        self.m_details = tk.Text(g3, height=3, font=FNT["default"],
                                 relief="solid", bd=1, bg=C["surface"])
        self.m_details.grid(row=5, column=0, columnspan=4, sticky="ew",
                            padx=(0, 8), ipady=2)

        # Violoniste section
        g4 = _sec("Option Violoniste")
        self.m_with_violoniste = tk.IntVar(value=0)
        tk.Checkbutton(g4, text="Inclure l'option violoniste",
                       variable=self.m_with_violoniste,
                       command=self._m_toggle_violon,
                       bg=C["surface"], fg=C["text"],
                       activebackground=C["surface"],
                       selectcolor=C["surface"], font=FNT["bold"]).grid(
            row=0, column=0, columnspan=4, sticky="w", pady=(0, 8))
        self.m_violon_date  = _field(g4, 1, 0, "Date Rayhane (JJ/MM/AAAA)", bg=C["surface"])
        self.m_violon_debut = _field(g4, 1, 1, "Heure début", bg=C["surface"])
        self.m_violon_fin   = _field(g4, 1, 2, "Heure fin", bg=C["surface"])
        btn_ray2 = RoundedButton(g4, text="Vérifier dispo Rayhane",
                                 command=self._m_check_rayhane, style="secondary",
                                 padx=16, pady=7)
        import tkinter.font as _tf3; _f3 = _tf3.Font(font=FNT["bold"])
        btn_ray2.configure(width=_f3.measure("Vérifier dispo Rayhane")+34,
                           height=_f3.metrics("linespace")+16)
        btn_ray2.grid(row=1, column=3, sticky="w", padx=(8, 0))
        self.m_violon_dispo_lbl = tk.Label(g4, text="", bg=C["surface"],
                                           fg=C["text_m"], font=FNT["sm"])
        self.m_violon_dispo_lbl.grid(row=2, column=1, columnspan=3, sticky="w",
                                     pady=(8, 0))
        self.m_with_sono = tk.IntVar(value=0)
        tk.Checkbutton(g4, text="Sans sono",
                       variable=self.m_with_sono,
                       bg=C["surface"], fg=C["text"],
                       activebackground=C["surface"],
                       selectcolor=C["surface"], font=FNT["default"]).grid(
            row=3, column=0, columnspan=4, sticky="w", pady=(8, 0))

        # Arabic info block
        ar_wr = tk.Frame(sf2, bg=C["bg"])
        ar_wr.pack(fill="x", padx=20, pady=(0, 10))
        ar_card = tk.Frame(ar_wr, bg=C["surface"])
        ar_card.pack(fill="x")
        tk.Frame(ar_card, bg=C["accent"], height=2).pack(fill="x")
        tk.Label(ar_card, text="معلومات العقد — Informations du contrat (arabe)",
                 bg=C["card_hdr"], fg=C["primary"], font=FNT["bold"]).pack(
            fill="x", anchor="w", padx=16, pady=8)
        tk.Frame(ar_card, bg=C["border"], height=1).pack(fill="x")
        ar_g = tk.Frame(ar_card, bg=C["surface"])
        ar_g.pack(fill="x", padx=16, pady=12)
        for cc in range(4): ar_g.columnconfigure(cc, weight=1)
        def _arl(text, row, col):
            tk.Label(ar_g, text=text, bg=C["surface"],
                     fg=C["text_m"], font=FNT["sm"], anchor="e").grid(
                row=row*2, column=col, sticky="e", padx=(8, 4))
        def _are(row, col):
            e = RoundedEntry(ar_g, font=("Segoe UI", 10), justify="right")
            e.grid(row=row*2+1, column=col, sticky="ew",
                   padx=(0, 16), pady=(0, 10), ipady=2)
            return e
        _arl("الاسم و اللقب", 0, 0); self.m_ar_nom = _are(0, 0)
        _arl("مكان بطاقة التعريف", 0, 1); self.m_ar_lieu_cin = _are(0, 1)
        _arl("قاعة الأفراح", 0, 2); self.m_ar_salle = _are(0, 2)
        _arl("تاريخ التحرير", 0, 3); self.m_ar_date = _are(0, 3)
        self.m_ar_date.insert(0, date.today().strftime("%d/%m/%Y"))
        _arl("عدد العقد", 1, 0); self.m_ar_num_contrat = _are(1, 0)
        self.m_ar_num_contrat.insert(0, f"{date.today().year}/")

        # Disable violon fields initially
        self._violon_fields = [self.m_violon_date, self.m_violon_debut, self.m_violon_fin]
        self._m_toggle_violon()

    # ── Helpers ───────────────────────────────────────────────────
    @staticmethod
    def _to_iso(ddmmyyyy):
        d, m, y = ddmmyyyy.strip().split("/")
        return f"{y}-{m}-{d}"

    @staticmethod
    def _from_iso(iso):
        if not iso: return ""
        try:
            from datetime import datetime
            return datetime.strptime(iso, "%Y-%m-%d").strftime("%d/%m/%Y")
        except Exception:
            return iso

    def _search(self):
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        q = self.search_var.get().strip()
        rows = database.search_reservations(q) if q else database.get_all_reservations()
        for r in rows:
            self.result_tree.insert("", "end", iid=str(r["id"]), values=(
                r["id"], r.get("num_bon") or "", r.get("nom_client") or "",
                r.get("salle") or "", r.get("date_evenement") or "",
                f"{r.get('heure_debut','')}-{r.get('heure_fin','')}",
                f"{r.get('forfait') or 0:.3f}"))

    def _on_select(self, _=None):
        sel = self.result_tree.selection()
        if not sel: return
        self._res_id = int(sel[0])
        res = database.get_reservation(self._res_id)
        if not res: return
        self.select_lbl.configure(
            text=f"  ✏  Modification :  {res.get('nom_client','')}   |   N° {res.get('num_bon','')}",
            fg=C["primary"])
        # Fill fields
        self.m_salle.set(res.get("salle", "Almes"))
        for attr, key, conv in [
            ("m_date_evt",   "date_evenement", self._from_iso),
            ("m_heure_debut","heure_debut",     str),
            ("m_heure_fin",  "heure_fin",       str),
            ("m_nom",        "nom_client",      str),
            ("m_tel1",       "tel1",            str),
            ("m_tel2",       "tel2",            str),
            ("m_cin",        "cin",             str),
            ("m_date_cin",   "date_cin",        self._from_iso),
            ("m_lieu_cin",   "lieu_cin",        str),
            ("m_nb_places",  "nb_places",       lambda v: str(v) if v else ""),
            ("m_type_evt",   "type_evenement",  str),
            ("m_forfait",    "forfait",         lambda v: str(v) if v else ""),
            ("m_n_facture",  "n_facture",       str),
            ("m_date_contrat","date_contrat",   self._from_iso),
        ]:
            widget = getattr(self, attr)
            val = conv(res.get(key, "") or "")
            widget.delete(0, "end")
            widget.insert(0, val)
        self.m_methode.set(res.get("methode_paiement", "Espèce") or "Espèce")
        for attr, key in [("m_caisse", "caisse"), ("m_tresorerie", "tresorerie")]:
            w = getattr(self, attr)
            w.delete(0, "end")
            w.insert(0, res.get(key, "") or "")
        self.m_details.delete("1.0", "end")
        self.m_details.insert("1.0", res.get("details_extra", "") or "")
        with_v = int(res.get("with_violoniste") or 0)
        self.m_with_violoniste.set(with_v)
        self._m_toggle_violon()
        if with_v:
            for attr, key, conv in [
                ("m_violon_date",  "violoniste_date",        self._from_iso),
                ("m_violon_debut", "violoniste_heure_debut", str),
                ("m_violon_fin",   "violoniste_heure_fin",   str),
            ]:
                w = getattr(self, attr)
                w.delete(0, "end")
                w.insert(0, conv(res.get(key, "") or ""))
        self.m_with_sono.set(int(res.get("with_sono") or 0))
        self.m_violon_dispo_lbl.configure(text="")
        self.status_lbl.configure(text="")
        self._m_refresh_arabic_info()

    def _m_toggle_violon(self):
        state = "normal" if self.m_with_violoniste.get() else "disabled"
        for w in self._violon_fields:
            w.configure(state=state)

    def _m_check_rayhane(self):
        self.m_violon_dispo_lbl.configure(text="", fg=C["text_m"])
        raw = self.m_violon_date.get().strip() or self.m_date_evt.get().strip()
        try:
            iso = self._to_iso(raw)
        except Exception:
            self.m_violon_dispo_lbl.configure(text="⚠  Date invalide.", fg=C["warning"])
            return
        hd = self.m_violon_debut.get().strip()
        hf = self.m_violon_fin.get().strip()
        if not hd or not hf:
            self.m_violon_dispo_lbl.configure(
                text="⚠  Heure début et fin requises.", fg=C["warning"])
            return
        ok, msg = rules.check_violoniste_slot(
            None, iso, hd, hf, self.m_with_sono.get(), exclude_id=self._res_id)
        self.m_violon_dispo_lbl.configure(
            text=f"✔  Disponible {raw} {hd}→{hf}" if ok else f"✘  {msg}",
            fg=C["success"] if ok else C["danger"])

    def _collect(self):
        if not self._res_id:
            raise ValueError("Aucune réservation sélectionnée.")
        return {
            "salle":           self.m_salle.get(),
            "date_evenement":  self._to_iso(self.m_date_evt.get()),
            "heure_debut":     self.m_heure_debut.get().strip(),
            "heure_fin":       self.m_heure_fin.get().strip(),
            "nom_client":      self.m_nom.get().strip(),
            "tel1":            self.m_tel1.get().strip(),
            "tel2":            self.m_tel2.get().strip(),
            "cin":             self.m_cin.get().strip(),
            "date_cin":        self._to_iso(self.m_date_cin.get()) if self.m_date_cin.get().strip() else None,
            "lieu_cin":        self.m_lieu_cin.get().strip(),
            "nb_places":       int(self.m_nb_places.get()) if self.m_nb_places.get().strip() else None,
            "type_evenement":  self.m_type_evt.get().strip(),
            "forfait":         float(self.m_forfait.get()) if self.m_forfait.get().strip() else None,
            "n_facture":       self.m_n_facture.get().strip(),
            "date_contrat":    self._to_iso(self.m_date_contrat.get()) if self.m_date_contrat.get().strip() else None,
            "methode_paiement":self.m_methode.get(),
            "caisse":          self.m_caisse.get().strip(),
            "tresorerie":      self.m_tresorerie.get().strip(),
            "details_extra":   self.m_details.get("1.0","end").strip(),
            "with_violoniste": 1 if self.m_with_violoniste.get() else 0,
            "violoniste_date": self._to_iso(self.m_violon_date.get()) if self.m_with_violoniste.get() and self.m_violon_date.get().strip() else None,
            "violoniste_heure_debut": self.m_violon_debut.get().strip() if self.m_with_violoniste.get() else None,
            "violoniste_heure_fin":   self.m_violon_fin.get().strip() if self.m_with_violoniste.get() else None,
            "with_sono": 1 if self.m_with_sono.get() else 0,
        }

    def _save(self):
        self.status_lbl.configure(text="", fg=C["danger"])
        try:
            data = self._collect()
        except Exception as ex:
            self.status_lbl.configure(text=str(ex))
            return
        conn = database.get_connection()
        sets = ", ".join(f"{k}=?" for k in data)
        conn.execute(f"UPDATE reservations SET {sets} WHERE id=?",
                     list(data.values()) + [self._res_id])
        conn.commit(); conn.close()
        self.status_lbl.configure(
            text=f"✔  Réservation #{self._res_id} mise à jour.", fg=C["success"])
        self.app.refresh_list()
        self._search()

    def _generate(self):
        self.status_lbl.configure(text="", fg=C["danger"])
        if not self._res_id:
            self.status_lbl.configure(text="Aucune réservation sélectionnée.")
            return
        try:
            data = self._collect()
            conn = database.get_connection()
            sets = ", ".join(f"{k}=?" for k in data)
            conn.execute(f"UPDATE reservations SET {sets} WHERE id=?",
                         list(data.values()) + [self._res_id])
            conn.commit(); conn.close()
        except Exception as ex:
            self.status_lbl.configure(text=f"Erreur sauvegarde : {ex}")
            return
        res = database.get_reservation(self._res_id)
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        output_path = os.path.join(OUTPUT_DIR,
            f"Bon_Reservation_{res.get('num_bon') or self._res_id}.docx".replace("/", "-"))
        try:
            contract_generator.generate_contract(res, output_path)
            os.startfile(output_path)
            self.status_lbl.configure(text=f"✔  Contrat généré.", fg=C["success"])
        except Exception as ex:
            self.status_lbl.configure(text=f"Erreur : {ex}")

    def _m_refresh_arabic_info(self):
        try:
            self.m_ar_nom.delete(0, "end")
            self.m_ar_nom.insert(0, self.m_nom.get().strip())
            self.m_ar_lieu_cin.delete(0, "end")
            self.m_ar_lieu_cin.insert(0, self.m_lieu_cin.get().strip())
            salle_ar = contract_generator.SALLE_AR.get(self.m_salle.get(), self.m_salle.get())
            self.m_ar_salle.delete(0, "end")
            self.m_ar_salle.insert(0, salle_ar)
        except Exception:
            pass

    def _generate_arabic(self):
        self.status_lbl.configure(text="", fg=C["danger"])
        if not self._res_id:
            self.status_lbl.configure(text="Aucune réservation sélectionnée.")
            return
        self._m_refresh_arabic_info()
        try:
            data = self._collect()
            conn = database.get_connection()
            sets = ", ".join(f"{k}=?" for k in data)
            conn.execute(f"UPDATE reservations SET {sets} WHERE id=?",
                         list(data.values()) + [self._res_id])
            conn.commit(); conn.close()
        except Exception as ex:
            self.status_lbl.configure(text=f"Erreur sauvegarde : {ex}")
            return
        res = database.get_reservation(self._res_id)
        ar_info = {
            "ar_nom":      self.m_ar_nom.get().strip(),
            "ar_lieu_cin": self.m_ar_lieu_cin.get().strip(),
            "ar_salle":    self.m_ar_salle.get().strip(),
            "ar_date":     self.m_ar_date.get().strip(),
            "num_contrat": self.m_ar_num_contrat.get().strip(),
        }
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        client = res.get("nom_client") or ar_info["ar_nom"] or "brouillon"
        filename = (f"Contrat_AR_{client}_{res.get('date_evenement','')}.docx"
                    .replace("/", "-").replace(" ", "_"))
        output_path = os.path.join(OUTPUT_DIR, filename)
        try:
            contract_generator.generate_arabic_contract(res, ar_info, output_path)
            os.startfile(output_path)
            self.status_lbl.configure(text=f"✔  Contrat arabe généré.", fg=C["success"])
        except Exception as ex:
            self.status_lbl.configure(text=f"Erreur : {ex}")


# ═══════════════════════════════════════════════════════════════════
#  TAB — EXPORT EXCEL
# ═══════════════════════════════════════════════════════════════════
class ExportTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self._build()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        tk.Label(hdr, text="Export Calendrier Excel", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Card
        wrap = tk.Frame(self, bg=C["bg"])
        wrap.pack(padx=20, pady=20, anchor="nw")
        card = tk.Frame(wrap, bg=C["surface"])
        card.pack()
        tk.Frame(card, bg=C["success"], height=2).pack(fill="x")
        body = tk.Frame(card, bg=C["surface"])
        body.pack(padx=28, pady=24)

        tk.Label(body, text="Année", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
        self.year_var = tk.StringVar(value=str(date.today().year))
        RoundedEntry(body, textvariable=self.year_var, width=8).grid(
            row=1, column=0, sticky="w", padx=(0, 20), ipady=3)

        tk.Label(body, text="Mois (1-12)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=1, sticky="w")
        self.month_var = tk.StringVar(value=str(date.today().month))
        RoundedEntry(body, textvariable=self.month_var, width=6).grid(
            row=1, column=1, sticky="w", padx=(0, 24), ipady=3)

        export_btn = RoundedButton(body, text="📊  Exporter le calendrier Excel",
                                   command=self.export, style="success", padx=16, pady=7)
        import tkinter.font as _tf5; _f5 = _tf5.Font(font=FNT["bold"])
        export_btn.configure(width=_f5.measure("📊  Exporter le calendrier Excel")+34,
                             height=_f5.metrics("linespace")+16)
        export_btn.grid(row=1, column=2, sticky="w", padx=(0, 0))

    def export(self):
        try:
            year  = int(self.year_var.get())
            month = int(self.month_var.get())
        except ValueError:
            messagebox.showerror("Erreur", "Année/mois invalides.")
            return
        os.makedirs(EXPORT_DIR, exist_ok=True)
        default_name = f"Calendrier_Marassim_{year}_{month:02d}.xlsx"
        path = filedialog.asksaveasfilename(
            initialdir=EXPORT_DIR, initialfile=default_name,
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path: return
        excel_export.export_month_calendar(year, month, path)
        toast(os.path.basename(path), "success", "Calendrier exporté")


# ═══════════════════════════════════════════════════════════════════
#  TAB — AJOUTER EMPLOYÉ
# ═══════════════════════════════════════════════════════════════════
class EmployeeTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._build()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")

        # Big icon on the left
        tk.Label(hdr, text="👤", bg=C["surface"],
                 font=("Segoe UI", 32)).pack(side="left", padx=(16, 8), pady=10)

        # Title + subtitle stacked
        title_stack = tk.Frame(hdr, bg=C["surface"])
        title_stack.pack(side="left", pady=10)
        tk.Label(title_stack, text="Ajouter un employé", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(anchor="w")
        tk.Label(title_stack, text="Remplissez les informations ci-dessous",
                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"]).pack(anchor="w")

        # Big type badge on the right — updated live when radio changes
        self.hdr_type_badge = tk.Label(
            hdr, text="🌞  Ouvrier de JOUR",
            bg=C["primary"], fg="white",
            font=("Segoe UI", 13, "bold"),
            padx=18, pady=8,
        )
        self.hdr_type_badge.pack(side="right", padx=20, pady=10)

        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # ── Middle container: footer + scrollable form ──
        mid = tk.Frame(self, bg=C["bg"])
        mid.pack(fill="both", expand=True)

        # Footer
        footer = tk.Frame(mid, bg=C["surface"])
        tk.Frame(footer, bg=C["border"], height=1).pack(fill="x")
        btn_bar = tk.Frame(footer, bg=C["surface"])
        btn_bar.pack(fill="x", padx=20, pady=10)
        _rbtn(btn_bar, "✔  Enregistrer l'employé", command=self._save, style="primary", padx=16, pady=7)
        _rbtn(btn_bar, "🔄  Réinitialiser", command=self._reset, style="secondary", padx=16, pady=7)
        self.status_label = tk.Label(footer, text="", bg=C["surface"],
                                     fg=C["danger"], font=FNT["sm"])
        self.status_label.pack(padx=20, pady=(0, 8), anchor="w")
        footer.pack(fill="x", side="bottom")

        # Scrollable form
        outer, sf = _scrollable(mid)
        outer.pack(fill="both", expand=True)

        # ── Section: Informations de base ──
        wrapper1 = tk.Frame(sf, bg=C["bg"])
        wrapper1.pack(fill="x", padx=20, pady=(0, 10))
        sec1 = tk.Frame(wrapper1, bg=C["surface"])
        sec1.pack(fill="x")
        tk.Frame(sec1, bg=C["primary"], height=2).pack(fill="x")
        tk.Label(sec1, text="Informations de l'employé", bg=C["card_hdr"],
                 fg=C["primary"], font=FNT["bold"]).pack(
            fill="x", anchor="w", padx=16, pady=10)
        tk.Frame(sec1, bg=C["border"], height=1).pack(fill="x")
        
        g1 = tk.Frame(sec1, bg=C["surface"])
        g1.pack(fill="x", padx=16, pady=(0, 14))
        for c in range(4):
            g1.columnconfigure(c, weight=1)

        self.nom_entry    = _field(g1, 0, 0, "Nom", bg=C["surface"])
        self.prenom_entry = _field(g1, 0, 1, "Prénom", bg=C["surface"])
        self.role_entry   = _field(g1, 0, 2, "Rôle / Poste", bg=C["surface"])
        self.icone_entry  = _field(g1, 0, 3, "Icône (emoji)", bg=C["surface"])
        self.icone_entry.insert(0, "👤")

        # Type ouvrier — toggle buttons (JOUR / NUIT)
        type_fr = tk.Frame(g1, bg=C["surface"])
        type_fr.grid(row=2, column=0, columnspan=4, sticky="w", padx=4, pady=(12, 4))
        tk.Label(type_fr, text="Type d'ouvrier", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(anchor="w", pady=(0, 6))

        self.type_ouvrier_var = tk.StringVar(value="jour")

        btn_row = tk.Frame(type_fr, bg=C["surface"])
        btn_row.pack(anchor="w")

        self._btn_jour = tk.Button(
            btn_row,
            text="☀  Ouvrier de JOUR",
            bg=C["primary"], fg="white",
            activebackground=C["primary_d"], activeforeground="white",
            relief="flat", font=("Segoe UI", 11, "bold"),
            padx=22, pady=10, cursor="hand2",
            command=lambda: self._select_type("jour"),
        )
        self._btn_jour.pack(side="left", padx=(0, 10))

        self._btn_nuit = tk.Button(
            btn_row,
            text="🌙  Ouvrier de NUIT",
            bg=C["stripe"], fg=C["text"],
            activebackground="#8B5CF6", activeforeground="white",
            relief="flat", font=("Segoe UI", 11, "bold"),
            padx=22, pady=10, cursor="hand2",
            command=lambda: self._select_type("nuit"),
        )
        self._btn_nuit.pack(side="left")

        # Small badge below (kept for _update_type_style compatibility)
        self.type_badge = tk.Label(type_fr, text="", bg=C["surface"])
        # hidden — badge is now the buttons themselves

        # ── Section: Salaires ──
        wrapper2 = tk.Frame(sf, bg=C["bg"])
        wrapper2.pack(fill="x", padx=20, pady=(0, 10))
        sec2 = tk.Frame(wrapper2, bg=C["surface"])
        sec2.pack(fill="x")
        tk.Frame(sec2, bg=C["accent"], height=2).pack(fill="x")
        tk.Label(sec2, text="Grille de salaire", bg=C["card_hdr"],
                 fg=C["primary"], font=FNT["bold"]).pack(
            fill="x", anchor="w", padx=16, pady=10)
        tk.Frame(sec2, bg=C["border"], height=1).pack(fill="x")
        
        g2 = tk.Frame(sec2, bg=C["surface"])
        g2.pack(fill="x", padx=16, pady=(0, 14))
        for c in range(4):
            g2.columnconfigure(c, weight=1)

        self.salaire_journalier = _field(g2, 0, 0, "Salaire journalier (DT)", bg=C["surface"])
        self.salaire_journalier.insert(0, "0.000")
        self.prix_heure_supp    = _field(g2, 0, 1, "Prix heure supplémentaire (DT)", bg=C["surface"])
        self.prix_heure_supp.insert(0, "0.000")

        # Wrapper for midi/après-midi/soirée — hidden for ouvrier de jour
        self._periode_wrapper = tk.Frame(sec2, bg=C["surface"])
        self._periode_wrapper.pack(fill="x", padx=16, pady=(0, 4))
        for c in range(3):
            self._periode_wrapper.columnconfigure(c, weight=1)

        self.salaire_midi       = _field(self._periode_wrapper, 0, 0, "Salaire midi (DT)", bg=C["surface"])
        self.salaire_midi.insert(0, "0.000")
        self.salaire_apres_midi = _field(self._periode_wrapper, 0, 1, "Salaire après-midi (DT)", bg=C["surface"])
        self.salaire_apres_midi.insert(0, "0.000")
        self.salaire_soiree     = _field(self._periode_wrapper, 0, 2, "Salaire soirée (DT)", bg=C["surface"])
        self.salaire_soiree.insert(0, "0.000")

        # Salaire minimum removed — no longer used

        # Apply initial visibility based on default type (jour)
        self._update_type_style()

    def _save(self):
        self.status_label.configure(text="", fg=C["danger"])
        nom = self.nom_entry.get().strip()
        prenom = self.prenom_entry.get().strip()
        role = self.role_entry.get().strip()
        if not nom or not prenom or not role:
            self.status_label.configure(text="Nom, prénom et rôle sont obligatoires.")
            return
        try:
            data = {
                "nom": nom,
                "prenom": prenom,
                "role": role,
                "icone": self.icone_entry.get().strip() or "👤",
                "type_ouvrier": self.type_ouvrier_var.get(),
                "salaire_journalier": float(self.salaire_journalier.get().strip() or 0),
                "prix_heure_supp": float(self.prix_heure_supp.get().strip() or 0),
                "salaire_minimum": 0.0,
                "salaire_midi": float(self.salaire_midi.get().strip() or 0),
                "salaire_apres_midi": float(self.salaire_apres_midi.get().strip() or 0),
                "salaire_soiree": float(self.salaire_soiree.get().strip() or 0),
            }
        except ValueError as e:
            self.status_label.configure(text=f"Erreur de saisie : {e}")
            return
        emp_id = database.insert_employee(data)
        self.status_label.configure(
            text=f"✔  Employé {nom} {prenom} enregistré (ID: {emp_id}).",
            fg=C["success"])
        toast(f"{nom} {prenom}", "success", "Employé ajouté")
        self._reset()

    def _select_type(self, val):
        """Called when a type toggle button is clicked."""
        self.type_ouvrier_var.set(val)
        self._update_type_style()

    def _update_type_style(self):
        val = self.type_ouvrier_var.get()
        if val == "jour":
            self._btn_jour.configure(bg=C["primary"],  fg="white")
            self._btn_nuit.configure(bg=C["stripe"],   fg=C["text"])
            self.hdr_type_badge.configure(text="🌞  Ouvrier de JOUR", bg=C["primary"])
            self._periode_wrapper.pack_forget()
        else:
            self._btn_jour.configure(bg=C["stripe"],   fg=C["text"])
            self._btn_nuit.configure(bg="#8B5CF6",     fg="white")
            self.hdr_type_badge.configure(text="🌙  Ouvrier de NUIT", bg="#8B5CF6")
            self._periode_wrapper.pack(fill="x", padx=16, pady=(0, 4))

    def _reset(self):
        self.nom_entry.delete(0, "end")
        self.prenom_entry.delete(0, "end")
        self.role_entry.delete(0, "end")
        self.icone_entry.delete(0, "end"); self.icone_entry.insert(0, "👤")
        self.type_ouvrier_var.set("jour")
        self._update_type_style()
        for e in [self.salaire_journalier, self.prix_heure_supp,
                  self.salaire_midi, self.salaire_apres_midi, self.salaire_soiree]:
            e.delete(0, "end"); e.insert(0, "0.000")
        self.status_label.configure(text="")


# ═══════════════════════════════════════════════════════════════════
#  FICHE EMPLOYÉ — popup standalone
# ═══════════════════════════════════════════════════════════════════
def _open_employee_fiche(parent_widget, emp_id: int):
    """
    Fiche complète d'un employé :
      • Header : identité + grille tarifaire
      • Historique des pointages groupés par date (midi / après-midi / soirée)
        avec total journalier et total général dû
      • Section paiements : historique + formulaire de saisie du montant payé
        et calcul automatique du reste
    """
    from datetime import datetime as _dt
    from collections import OrderedDict

    fiche = database.get_employee_fiche_data(emp_id)
    if not fiche:
        messagebox.showerror("Erreur", f"Employé ID {emp_id} introuvable.")
        return

    emp = fiche["employee"]

    icone  = emp.get("icone") or "👤"
    nom    = emp.get("nom", "")
    prenom = emp.get("prenom", "")
    role   = emp.get("role", "")

    TAR = {
        "midi":             float(emp.get("salaire_midi")       or 0),
        "apres_midi":       float(emp.get("salaire_apres_midi") or 0),
        "soiree":           float(emp.get("salaire_soiree")     or 0),
        "journee_complete": float(emp.get("salaire_journalier") or 0),
    }
    PER_LABEL = {
        "midi":             "Midi",
        "apres_midi":       "Apres-midi",
        "soiree":           "Soiree",
        "journee_complete": "Journee complete",
    }
    PER_ICON = {
        "midi":             "Midi",
        "apres_midi":       "Apres-midi",
        "soiree":           "Soiree",
        "journee_complete": "Journee complete",
    }
    PER_COLOR = {
        "midi":             "#F59E0B",
        "apres_midi":       "#0EA5E9",
        "soiree":           "#8B5CF6",
        "journee_complete": "#16A34A",
    }
    # Ordered list of periods for consistent display
    PER_ORDER = ["midi", "apres_midi", "soiree", "journee_complete"]

    def _iso_to_fr(iso):
        try:
            return _dt.strptime(iso, "%Y-%m-%d").strftime("%d/%m/%Y")
        except Exception:
            return iso or ""

    def _fr_to_iso(fr):
        try:
            d, m, y = fr.strip().split("/")
            return f"{y}-{m}-{d}"
        except Exception:
            return ""

    # ── Window ──────────────────────────────────────────────────
    win = tk.Toplevel(parent_widget)
    win.title(f"Fiche — {nom} {prenom}")
    win.minsize(900, 640)
    win.configure(bg=C["bg"])
    win.grab_set()
    # Screen-relative: 82% wide, 85% tall
    win.update_idletasks()
    _sw = win.winfo_screenwidth()
    _sh = win.winfo_screenheight()
    _w = max(900, int(_sw * 0.82))
    _h = max(640, int(_sh * 0.85))
    win.geometry(f"{_w}x{_h}+{(_sw - _w)//2}+{(_sh - _h)//2}")

    # ── Top header band ─────────────────────────────────────────
    hdr = tk.Frame(win, bg=C["sidebar"])
    hdr.pack(fill="x")
    tk.Label(hdr, text=f"{icone}  {nom} {prenom}",
             bg=C["sidebar"], fg="white",
             font=("Segoe UI", 14, "bold")).pack(side="left", padx=20, pady=14)
    tk.Label(hdr, text=f"  —  {role}",
             bg=C["sidebar"], fg="#94A3B8",
             font=("Segoe UI", 10)).pack(side="left", pady=14)

    # Tarif chips (right side of header)
    chip_fr = tk.Frame(hdr, bg=C["sidebar"])
    chip_fr.pack(side="right", padx=20, pady=10)
    for key in PER_ORDER:
        t   = TAR.get(key, 0)
        clr = PER_COLOR[key]
        lbl = PER_LABEL[key]
        ch  = tk.Frame(chip_fr, bg=clr, padx=8, pady=3)
        ch.pack(side="left", padx=3)
        tk.Label(ch, text=f"{lbl}: {t:.3f} DT",
                 bg=clr, fg="white",
                 font=("Segoe UI", 8, "bold")).pack()

    tk.Frame(win, bg=C["border"], height=1).pack(fill="x")

    # ── Main layout: top=pointages, bottom=paiements ────────────
    # We use a vertical PanedWindow feel via pack with weights
    body = tk.Frame(win, bg=C["bg"])
    body.pack(fill="both", expand=True)
    body.columnconfigure(0, weight=3)
    body.columnconfigure(1, weight=2)
    body.rowconfigure(0, weight=1)

    # ════════════════════════════════════════════════════════════
    #  LEFT — Historique pointages groupés par date
    # ════════════════════════════════════════════════════════════
    left = tk.Frame(body, bg=C["bg"])
    left.grid(row=0, column=0, sticky="nsew", padx=(16, 8), pady=16)
    left.rowconfigure(1, weight=1)
    left.columnconfigure(0, weight=1)

    # Header
    hdr_left = tk.Frame(left, bg=C["surface"])
    hdr_left.grid(row=0, column=0, columnspan=2, sticky="ew")
    tk.Frame(hdr_left, bg=C["primary"], height=2).pack(fill="x")
    lhdr_in = tk.Frame(hdr_left, bg=C["card_hdr"])
    lhdr_in.pack(fill="x")
    tk.Label(lhdr_in, text="Historique des pointages",
             bg=C["card_hdr"], fg=C["primary"],
             font=FNT["bold"]).pack(side="left", padx=14, pady=8)
    tk.Label(lhdr_in,
             text="(chaque ligne = 1 session : date | midi | apres-midi | soiree | total jour)",
             bg=C["card_hdr"], fg=C["text_m"],
             font=("Segoe UI", 8)).pack(side="left", pady=8)
    tk.Frame(hdr_left, bg=C["border"], height=1).pack(fill="x")

    # Treeview pointages — columns: date + one col per period + total jour
    pt_cols = ("date", "midi", "apres_midi", "soiree", "journee_complete", "total_jour", "notes")
    pt_tree = ttk.Treeview(left, columns=pt_cols, show="headings", selectmode="browse")

    pt_heads = {
        "date":             "Date",
        "midi":             "Midi (DT)",
        "apres_midi":       "Apres-midi (DT)",
        "soiree":           "Soiree (DT)",
        "journee_complete": "Journee (DT)",
        "total_jour":       "TOTAL JOUR",
        "notes":            "Notes",
    }
    pt_widths = {
        "date":             95,
        "midi":             85,
        "apres_midi":       100,
        "soiree":           85,
        "journee_complete": 95,
        "total_jour":       90,
        "notes":            0,
    }
    for c in pt_cols:
        pt_tree.heading(c, text=pt_heads[c])
        pt_tree.column(
            c,
            width=pt_widths[c] if pt_widths[c] else 160,
            stretch=(pt_widths[c] == 0),
            anchor="e" if c not in ("date", "notes") else "w",
        )

    # Row colour tags
    pt_tree.tag_configure("day_row",  background=C["surface"])
    pt_tree.tag_configure("day_alt",  background=C["stripe"])
    pt_tree.tag_configure("tot_row",  background=_lighten(C["primary"], 0.88),
                          font=FNT["bold"])

    pt_vsb = ttk.Scrollbar(left, orient="vertical", command=pt_tree.yview)
    pt_tree.configure(yscrollcommand=pt_vsb.set)
    pt_vsb.grid(row=1, column=1, sticky="ns")
    pt_tree.grid(row=1, column=0, sticky="nsew")

    # Summary strip below pointages tree
    pt_sum = tk.Frame(left, bg=C["surface"])
    pt_sum.grid(row=2, column=0, columnspan=2, sticky="ew")
    tk.Frame(pt_sum, bg=C["border"], height=1).pack(fill="x")
    pt_sum_in = tk.Frame(pt_sum, bg=C["surface"])
    pt_sum_in.pack(fill="x", padx=12, pady=8)

    lbl_nb_pt  = tk.Label(pt_sum_in, text="", bg=C["surface"],
                          fg=C["text_m"], font=FNT["sm"])
    lbl_nb_pt.pack(side="left", padx=(0, 20))
    lbl_tot_du = tk.Label(pt_sum_in, text="", bg=C["surface"],
                          fg=C["primary"], font=("Segoe UI", 11, "bold"))
    lbl_tot_du.pack(side="left")

    # ════════════════════════════════════════════════════════════
    #  RIGHT — Paiements + formulaire
    # ════════════════════════════════════════════════════════════
    right = tk.Frame(body, bg=C["bg"])
    right.grid(row=0, column=1, sticky="nsew", padx=(0, 16), pady=16)
    right.rowconfigure(1, weight=1)
    right.columnconfigure(0, weight=1)

    # Section title
    rhdr = tk.Frame(right, bg=C["surface"])
    rhdr.grid(row=0, column=0, sticky="ew")
    tk.Frame(rhdr, bg=C["success"], height=2).pack(fill="x")
    tk.Label(rhdr, text="Paiements",
             bg=C["card_hdr"], fg=C["success"],
             font=FNT["bold"]).pack(anchor="w", padx=14, pady=8)
    tk.Frame(rhdr, bg=C["border"], height=1).pack(fill="x")

    # Treeview paiements
    pay_cols = ("date", "montant", "notes")
    pay_tree = ttk.Treeview(right, columns=pay_cols, show="headings", height=8)
    pay_heads = {"date": "Date", "montant": "Montant (DT)", "notes": "Notes"}
    pay_widths = {"date": 95, "montant": 110, "notes": 0}
    for c in pay_cols:
        pay_tree.heading(c, text=pay_heads[c])
        pay_tree.column(c, width=pay_widths[c] if pay_widths[c] else 160,
                        stretch=(pay_widths[c] == 0),
                        anchor="e" if c == "montant" else "w")
    pay_vsb = ttk.Scrollbar(right, orient="vertical", command=pay_tree.yview)
    pay_tree.configure(yscrollcommand=pay_vsb.set)
    pay_vsb.grid(row=1, column=1, sticky="ns")
    pay_tree.grid(row=1, column=0, sticky="nsew")

    # Totals strip
    tot_fr = tk.Frame(right, bg=C["surface"])
    tot_fr.grid(row=2, column=0, columnspan=2, sticky="ew")
    tk.Frame(tot_fr, bg=C["border"], height=1).pack(fill="x")
    tot_in = tk.Frame(tot_fr, bg=C["surface"])
    tot_in.pack(fill="x", padx=12, pady=8)

    is_jour = (emp.get("type_ouvrier") or "jour") == "jour"

    lbl_total_du  = tk.Label(tot_in, text="", bg=C["surface"],
                              fg=C["text"], font=FNT["sm"])
    lbl_total_du.grid(row=0, column=0, sticky="w", padx=(0, 16))

    lbl_total_pay = tk.Label(tot_in, text="", bg=C["surface"],
                              fg=C["success"], font=FNT["sm"])
    lbl_total_pay.grid(row=1, column=0, sticky="w", padx=(0, 16))
    lbl_reste     = tk.Label(tot_in, text="", bg=C["surface"],
                              fg=C["danger"], font=("Segoe UI", 10, "bold"))
    lbl_reste.grid(row=2, column=0, sticky="w")

    # ── New payment form ─────────────────────────────────────────
    form_fr = tk.Frame(right, bg=C["card_hdr"])
    form_fr.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(8, 0))
    tk.Frame(form_fr, bg=C["accent"], height=2).pack(fill="x")
    tk.Label(form_fr, text="💰  Nouveau paiement",
             bg=C["card_hdr"], fg=C["accent"],
             font=FNT["bold"]).pack(anchor="w", padx=14, pady=(8, 4))
    tk.Frame(form_fr, bg=C["border"], height=1).pack(fill="x")

    form_body = tk.Frame(form_fr, bg=C["card_hdr"])
    form_body.pack(fill="x", padx=14, pady=10)

    # Row 0 : montant + date
    tk.Label(form_body, text="Montant (DT)", bg=C["card_hdr"],
             fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=0, sticky="w")
    pay_montant_e = RoundedEntry(form_body, width=12, bg_canvas=C["card_hdr"])
    pay_montant_e.grid(row=1, column=0, sticky="ew", padx=(0, 10), ipady=3)

    tk.Label(form_body, text="Date paiement", bg=C["card_hdr"],
             fg=C["text_m"], font=FNT["sm"]).grid(row=0, column=1, sticky="w")
    pay_date_e = RoundedEntry(form_body, width=12, bg_canvas=C["card_hdr"])
    pay_date_e.insert(0, date.today().strftime("%d/%m/%Y"))
    pay_date_e.grid(row=1, column=1, sticky="ew", padx=(0, 10), ipady=3)

    form_body.columnconfigure(0, weight=1)
    form_body.columnconfigure(1, weight=1)

    # Row 2 : notes
    tk.Label(form_body, text="Notes (optionnel)", bg=C["card_hdr"],
             fg=C["text_m"], font=FNT["sm"]).grid(row=2, column=0, columnspan=2,
                                                   sticky="w", pady=(8, 2))
    pay_notes_e = RoundedEntry(form_body, bg_canvas=C["card_hdr"])
    pay_notes_e.grid(row=3, column=0, columnspan=2, sticky="ew", ipady=3)

    pay_status = tk.Label(form_fr, text="", bg=C["card_hdr"],
                          fg=C["danger"], font=FNT["sm"])
    pay_status.pack(anchor="w", padx=14, pady=(4, 0))

    btn_row = tk.Frame(form_fr, bg=C["card_hdr"])
    btn_row.pack(fill="x", padx=14, pady=(6, 12))

    # ── Helpers : refresh UI ──────────────────────────────────────
    def _refresh():
        """Reload fiche data and repopulate both trees + totals."""
        data  = database.get_employee_fiche_data(emp_id)
        _pts  = data["pointages"]
        _pays = data["paiements"]
        _du   = data["total_du"]
        _paye = data["total_paye"]
        _rst  = data["reste"]

        # ── Pointages tree — group by date ──────────────────────
        for item in pt_tree.get_children():
            pt_tree.delete(item)

        # Build dict: date_iso → {periode: tarif, notes: [...]}
        by_date = OrderedDict()
        for p in sorted(_pts, key=lambda x: x.get("date_pointage", "")):
            d = p.get("date_pointage", "")
            if d not in by_date:
                by_date[d] = {"notes": [], "periods": {}}
            per   = p.get("periode") or ""
            tarif = float(p.get("tarif") or 0)
            by_date[d]["periods"][per] = tarif
            if p.get("notes"):
                by_date[d]["notes"].append(p["notes"])

        for idx, (d_iso, info) in enumerate(by_date.items()):
            periods = info["periods"]
            t_midi  = periods.get("midi", 0)
            t_am    = periods.get("apres_midi", 0)
            t_soir  = periods.get("soiree", 0)
            t_jour  = periods.get("journee_complete", 0)
            total_j = t_midi + t_am + t_soir + t_jour
            notes_str = " | ".join(info["notes"])
            tag = "day_alt" if idx % 2 else "day_row"
            pt_tree.insert("", "end", tags=(tag,), values=(
                _iso_to_fr(d_iso),
                f"{t_midi:.3f}"  if t_midi  else "—",
                f"{t_am:.3f}"    if t_am    else "—",
                f"{t_soir:.3f}"  if t_soir  else "—",
                f"{t_jour:.3f}"  if t_jour  else "—",
                f"{total_j:.3f}",
                notes_str,
            ))

        nb_days = len(by_date)
        nb_pts  = len(_pts)
        lbl_nb_pt.configure(text=f"{nb_pts} pointage(s) — {nb_days} jour(s)")
        lbl_tot_du.configure(text=f"  Total du :  {_du:.3f} DT")

        # ── Paiements tree ──────────────────────────────────────
        for item in pay_tree.get_children():
            pay_tree.delete(item)
        for p in _pays:
            pay_tree.insert("", "end", values=(
                _iso_to_fr(p.get("date_paiement", "")),
                f"{float(p.get('montant_total') or 0):.3f}",
                p.get("notes") or "",
            ))

        # ── Totals ──────────────────────────────────────────────
        net_a_payer = _du
        lbl_total_du.configure(text=f"Total du         :  {_du:.3f} DT")
        lbl_total_pay.configure(text=f"Montant payé     :  {_paye:.3f} DT")
        _rst_net = max(net_a_payer - _paye, 0.0)
        reste_clr = C["danger"] if _rst_net > 0.001 else C["success"]
        lbl_reste.configure(text=f"Reste            :  {_rst_net:.3f} DT",
                            fg=reste_clr)

        # Pre-fill montant with current reste
        pay_montant_e.delete(0, "end")
        pay_montant_e.insert(0, f"{_rst_net:.3f}" if _rst_net > 0 else "0.000")

    def _save_payment():
        pay_status.configure(text="", fg=C["danger"])
        raw_m = pay_montant_e.get().strip()
        raw_d = pay_date_e.get().strip()
        if not raw_m:
            pay_status.configure(text="Saisissez un montant.")
            return
        try:
            montant = float(raw_m)
            if montant <= 0:
                raise ValueError
        except ValueError:
            pay_status.configure(text="Montant invalide (nombre > 0).")
            return
        d_iso = _fr_to_iso(raw_d)
        if not d_iso:
            pay_status.configure(text="Date invalide (JJ/MM/AAAA).")
            return
        data_cur = database.get_employee_fiche_data(emp_id)
        data_pay = {
            "employee_id":         emp_id,
            "date_paiement":       d_iso,
            "periode_debut":       "2000-01-01",
            "periode_fin":         date.today().isoformat(),
            "montant_base":        data_cur["total_du"],
            "montant_heures_supp": 0,
            "montant_total":       montant,
            "notes":               pay_notes_e.get().strip() or None,
        }
        database.insert_paiement_employee(data_pay)
        pay_notes_e.delete(0, "end")
        pay_status.configure(
            text=f"Paiement de {montant:.3f} DT enregistre.",
            fg=C["success"])
        _refresh()

    # Wire save button
    tk.Button(btn_row, text="Enregistrer le paiement",
              bg=C["success"], fg="white",
              relief="flat", font=FNT["bold"],
              padx=14, pady=6, cursor="hand2",
              command=_save_payment).pack(side="left", padx=(0, 8))
    tk.Button(btn_row, text="Actualiser",
              bg=C["border"], fg=C["text"],
              relief="flat", font=FNT["bold"],
              padx=12, pady=6, cursor="hand2",
              command=_refresh).pack(side="left")

    # Close button (bottom)
    bot = tk.Frame(win, bg=C["surface"])
    bot.pack(fill="x", side="bottom")
    tk.Frame(bot, bg=C["border"], height=1).pack(fill="x")
    tk.Button(bot, text="Fermer",
              bg=C["sidebar"], fg="white",
              relief="flat", font=FNT["bold"],
              padx=20, pady=8, cursor="hand2",
              command=win.destroy).pack(side="right", padx=16, pady=8)

    # Initial population
    _refresh()


def _lighten(hex_color: str, factor: float = 0.85) -> str:
    """Return a lighter (pastel) version of a hex colour by blending toward white."""
    try:
        h = hex_color.lstrip("#")
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        r2 = int(r + (255 - r) * factor)
        g2 = int(g + (255 - g) * factor)
        b2 = int(b + (255 - b) * factor)
        return f"#{r2:02X}{g2:02X}{b2:02X}"
    except Exception:
        return "#F8F8F8"


# ═══════════════════════════════════════════════════════════════════
#  TAB — LISTE DES EMPLOYÉS
# ═══════════════════════════════════════════════════════════════════
class EmployeeListTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._build()
        self.refresh()

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        tk.Label(hdr, text="Liste des employés", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Toolbar
        tb = tk.Frame(self, bg=C["bg"])
        tb.pack(fill="x", padx=20, pady=10)
        _rbtn(tb, "Actualiser", command=self.refresh, style="secondary", padx=16, pady=7)
        _rbtn(tb, "Ouvrir Fiche", command=self._open_fiche, style="primary", padx=16, pady=7)
        _rbtn(tb, "Desactiver", command=self._deactivate, style="sm_danger", padx=10, pady=5)

        # Tree
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        cols = ("id", "icone", "nom", "prenom", "role", "salaire_j", "prix_hs", "actif")
        self.tree = ttk.Treeview(tree_wrap, columns=cols, show="headings", height=15)
        heads = {"id":"ID", "icone":"", "nom":"Nom", "prenom":"Prénom", "role":"Rôle",
                 "salaire_j":"Sal. jour (DT)", "prix_hs":"Heure supp (DT)", "actif":"Actif"}
        widths = {"id":40, "icone":40, "nom":140, "prenom":140, "role":150,
                  "salaire_j":110, "prix_hs":110, "actif":60}
        for c in cols:
            self.tree.heading(c, text=heads[c])
            self.tree.column(c, width=widths[c], anchor="center")
        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<Double-1>", lambda e: self._open_fiche())

        self.status_label = tk.Label(self, text="", bg=C["bg"],
                                     fg=C["text_m"], font=FNT["sm"])
        self.status_label.pack(padx=20, pady=(0, 10), anchor="w")

    def refresh(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        employees = database.get_all_employees(actif_only=False)
        for e in employees:
            actif_txt = "✓" if e.get("actif") else "✗"
            self.tree.insert("", "end", iid=str(e["id"]), values=(
                e["id"], e.get("icone","👤"), e["nom"], e["prenom"], e.get("role",""),
                f"{e.get('salaire_journalier',0):.3f}",
                f"{e.get('prix_heure_supp',0):.3f}",
                actif_txt
            ))
        self.status_label.configure(text=f"{len(employees)} employé(s).", fg=C["primary"])

    def _open_fiche(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Sélection requise",
                                   "Sélectionnez un employé dans la liste.")
            return
        emp_id = int(sel[0])
        _open_employee_fiche(self, emp_id)

    def _deactivate(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Sélection requise", "Sélectionnez un employé.")
            return
        emp_id = int(sel[0])
        emp = database.get_employee(emp_id)
        if not emp:
            return
        if messagebox.askyesno("Confirmer", f"Désactiver {emp['nom']} {emp['prenom']} ?"):
            database.delete_employee(emp_id)
            self.refresh()


# ═══════════════════════════════════════════════════════════════════
#  TAB — POINTAGE
# ═══════════════════════════════════════════════════════════════════
class PointageTab(tk.Frame):
    """
    Workflow :
      0. Filtrer les employés : Tous / Ouvrier de Jour / Ouvrier de Nuit
      1. Sélectionner une ou plusieurs Périodes selon le type
      2. Sélectionner un ou plusieurs Employés
      3. Saisir Date + Heures supplémentaires + Notes
      4. Enregistrer
    """

    # Périodes communes
    PERIODES_JOUR = [
        ("midi",             "Midi",             "#F59E0B"),
        ("apres_midi",       "Apres-midi",       "#0EA5E9"),
        ("journee_complete", "Journee complete", "#16A34A"),
    ]
    PERIODES_NUIT = [
        ("midi",             "Midi",             "#F59E0B"),
        ("apres_midi",       "Apres-midi",       "#0EA5E9"),
        ("soiree",           "Soiree",           "#8B5CF6"),
    ]
    PERIODES_TOUS = [
        ("midi",             "Midi",             "#F59E0B"),
        ("apres_midi",       "Apres-midi",       "#0EA5E9"),
        ("soiree",           "Soiree",           "#8B5CF6"),
        ("journee_complete", "Journee complete", "#16A34A"),
    ]

    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._emp_data    = []
        self._emp_frames  = {}
        self._sel_emps    = set()
        self._sel_per     = set()
        self._per_btns    = {}
        self._hs_entries  = {}   # emp_id → tk.Entry (per-employee heures supp)
        self._filter_type = tk.StringVar(value="tous")  # "tous"|"jour"|"nuit"
        self._build()

    # ── Construction ──────────────────────────────────────────────
    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        h_in = tk.Frame(hdr, bg=C["surface"])
        h_in.pack(fill="x", padx=20, pady=14)
        tk.Label(h_in, text="Pointage des employes", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(side="left")
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # Main layout
        mid = tk.Frame(self, bg=C["bg"])
        mid.pack(fill="both", expand=True)

        # Footer
        footer = tk.Frame(mid, bg=C["surface"])
        tk.Frame(footer, bg=C["border"], height=1).pack(fill="x")
        btn_bar = tk.Frame(footer, bg=C["surface"])
        btn_bar.pack(fill="x", padx=20, pady=10)
        _rbtn(btn_bar, "Enregistrer le pointage", command=self._save_pointage,
              style="primary", padx=18, pady=8)
        _rbtn(btn_bar, "Reinitialiser", command=self._reset,
              style="secondary", padx=16, pady=7)
        _rbtn(btn_bar, "Historique", command=self._view_history,
              style="secondary", padx=16, pady=7)
        self.status_label = tk.Label(footer, text="", bg=C["surface"],
                                     fg=C["danger"], font=FNT["sm"])
        self.status_label.pack(padx=20, pady=(0, 8), anchor="w")
        footer.pack(fill="x", side="bottom")

        # Scrollable body
        outer, sf = _scrollable(mid)
        outer.pack(fill="both", expand=True)

        # ── STEP 0 — Filtre type ouvrier ────────────────────────
        self._step_card(sf, "Etape 0 — Filtrer par type d'ouvrier", "#64748B")

        filt_wrap = tk.Frame(sf, bg=C["bg"])
        filt_wrap.pack(fill="x", padx=20, pady=(0, 4))
        filt_inner = tk.Frame(filt_wrap, bg=C["surface"])
        filt_inner.pack(fill="x")
        filt_row = tk.Frame(filt_inner, bg=C["surface"])
        filt_row.pack(fill="x", padx=16, pady=12)

        FILTER_OPTS = [
            ("tous",  "Tous les employes",  "#64748B"),
            ("jour",  "Ouvrier de Jour",    C["primary"]),
            ("nuit",  "Ouvrier de Nuit",    "#8B5CF6"),
        ]
        self._filter_btns = {}
        for val, lbl, clr in FILTER_OPTS:
            is_sel = (val == "tous")
            btn = tk.Button(
                filt_row, text=lbl,
                bg=clr if is_sel else C["stripe"],
                fg="white" if is_sel else C["text"],
                activebackground=clr, activeforeground="white",
                relief="flat", font=FNT["bold"],
                padx=18, pady=8, cursor="hand2",
                command=lambda v=val, c=clr: self._apply_type_filter(v, c),
            )
            btn.pack(side="left", padx=(0, 10))
            self._filter_btns[val] = (btn, clr)

        self.filter_hint = tk.Label(filt_inner,
                                    text="Affichage : Tous les employes",
                                    bg=C["surface"], fg=C["text_m"], font=FNT["sm"])
        self.filter_hint.pack(anchor="w", padx=16, pady=(0, 10))

        # ── STEP 1 — Période(s) ─────────────────────────────────
        self._step1_card = self._step_card(
            sf, "Étape 1 — Sélectionner la / les période(s)", C["primary"])

        per_wrap = tk.Frame(sf, bg=C["bg"])
        per_wrap.pack(fill="x", padx=20, pady=(0, 4))
        self._per_step_wrap = per_wrap
        self._per_outer = tk.Frame(per_wrap, bg=C["surface"])
        self._per_outer.pack(fill="x")
        self._pills_frame = tk.Frame(self._per_outer, bg=C["surface"])
        self._pills_frame.pack(fill="x", padx=16, pady=14)

        self.per_hint = tk.Label(self._per_outer, text="Aucune periode selectionnee",
                                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"])
        self.per_hint.pack(anchor="w", padx=16, pady=(0, 10))

        self._rebuild_period_pills()  # initial render with PERIODES_TOUS

        # ── STEP 2 — Employé(s) ─────────────────────────────────
        self._step2_anchor = self._step_card(sf, "Etape 2 — Selectionner le / les employe(s)", C["accent"])

        emp_outer = tk.Frame(sf, bg=C["bg"])
        emp_outer.pack(fill="x", padx=20, pady=(0, 4))
        emp_card = tk.Frame(emp_outer, bg=C["surface"])
        emp_card.pack(fill="x")

        emp_top = tk.Frame(emp_card, bg=C["surface"])
        emp_top.pack(fill="x", padx=16, pady=(10, 6))
        self.emp_hint = tk.Label(emp_top, text="Aucun employe selectionne",
                                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"])
        self.emp_hint.pack(side="left")
        tk.Button(emp_top, text="Recharger", bg=C["surface"], fg=C["primary"],
                  relief="flat", font=FNT["sm"], cursor="hand2",
                  command=self._reload_employees).pack(side="right")

        self.emp_grid = tk.Frame(emp_card, bg=C["surface"])
        self.emp_grid.pack(fill="x", padx=16, pady=(0, 14))
        tk.Frame(emp_card, bg=C["border"], height=1).pack(fill="x", padx=16)
        tk.Label(emp_card,
                 text="Cliquez sur un employe pour le selectionner / deselectionner.",
                 bg=C["surface"], fg=C["text_m"], font=FNT["sm"]).pack(
            anchor="w", padx=16, pady=(6, 10))

        # ── STEP 3 — Date + heures supp + notes ─────────────────
        self._step_card(sf, "Etape 3 — Date, heures supplementaires et notes", C["success"])

        det_outer = tk.Frame(sf, bg=C["bg"])
        det_outer.pack(fill="x", padx=20, pady=(0, 16))
        det_card = tk.Frame(det_outer, bg=C["surface"])
        det_card.pack(fill="x")

        det_body = tk.Frame(det_card, bg=C["surface"])
        det_body.pack(fill="x", padx=16, pady=14)
        det_body.columnconfigure(2, weight=1)

        # Date
        tk.Label(det_body, text="Date (JJ/MM/AAAA)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(
            row=0, column=0, sticky="w", pady=(0, 4))
        self.date_entry = RoundedEntry(det_body, width=14)
        self.date_entry.insert(0, date.today().strftime("%d/%m/%Y"))
        self.date_entry.grid(row=1, column=0, sticky="w", padx=(0, 24), ipady=3)

        # Heures supplémentaires — moved to individual employee cards
        # (no global HS field here)

        # Notes
        tk.Label(det_body, text="Notes (optionnel)", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).grid(
            row=0, column=2, sticky="w", pady=(0, 4))
        self.notes_text = tk.Text(det_body, height=3, font=FNT["default"],
                                  relief="solid", bd=1, bg=C["surface"],
                                  insertbackground=C["text"])
        self.notes_text.grid(row=1, column=2, sticky="ew", ipady=2)

        self._reload_employees()

    # ── Step card header helper ───────────────────────────────────
    def _step_card(self, parent, title, color):
        """Carte d'étape : ombre douce, en-tête blanc, barre d'accent."""
        outer = tk.Frame(parent, bg=C["bg"])
        outer.pack(fill="x", padx=22, pady=(14, 0))
        shadow = tk.Frame(outer, bg=C["shadow"])
        shadow.pack(fill="x", padx=(1, 0), pady=(2, 0))
        shell = tk.Frame(shadow, bg=C["border"])
        shell.pack(fill="x", padx=(0, 1), pady=(0, 2))
        wr = tk.Frame(shell, bg=C["surface"])
        wr.pack(fill="x", padx=1, pady=1)

        hdr = tk.Frame(wr, bg=C["surface"])
        hdr.pack(fill="x", padx=18, pady=(14, 0))
        tk.Frame(hdr, bg=color, width=3, height=16).pack(side="left", padx=(0, 10))
        tk.Label(hdr, text=title, bg=C["surface"], fg=C["text"],
                 font=FNT["section"]).pack(side="left")
        tk.Frame(wr, bg=C["border"], height=1).pack(fill="x", padx=18, pady=(12, 0))
        return wr

    # ── Type filter ───────────────────────────────────────────────
    def _apply_type_filter(self, val, clr):
        self._filter_type.set(val)
        # Update filter button styles
        for v, (btn, c) in self._filter_btns.items():
            if v == val:
                btn.configure(bg=c, fg="white")
            else:
                btn.configure(bg=C["stripe"], fg=C["text"])
        labels = {"tous": "Tous les employes",
                  "jour": "Ouvrier de Jour",
                  "nuit": "Ouvrier de Nuit"}
        self.filter_hint.configure(text=f"Affichage : {labels[val]}")
        # Reset period selection and rebuild pills for the selected type
        self._sel_per.clear()
        self._rebuild_period_pills()

        if val == "jour":
            # Hide step 1 and auto-select journee_complete
            self._step1_card.pack_forget()
            self._per_step_wrap.pack_forget()
            self._sel_per.add("journee_complete")
            self._update_per_hint()
        else:
            # Show step 1 + per_step_wrap (re-pack both in correct order)
            self._step1_card.pack_forget()
            self._per_step_wrap.pack_forget()
            self._step1_card.pack(fill="x", padx=20, pady=(12, 0),
                                  before=self._step2_anchor)
            self._per_step_wrap.pack(fill="x", padx=20, pady=(0, 4),
                                     before=self._step2_anchor)

        # Reload employee grid filtered by type
        self._reload_employees()

    def _rebuild_period_pills(self):
        """Destroy and re-create period pills based on current filter type."""
        for w in self._pills_frame.winfo_children():
            w.destroy()
        self._per_btns.clear()
        self._sel_per.clear()

        ft = self._filter_type.get()
        if ft == "jour":
            periodes = self.PERIODES_JOUR
        elif ft == "nuit":
            periodes = self.PERIODES_NUIT
        else:
            periodes = self.PERIODES_TOUS

        for key, label, color in periodes:
            btn = tk.Button(
                self._pills_frame, text=label,
                bg=C["stripe"], fg=C["text"],
                activebackground=color, activeforeground="white",
                relief="flat", font=FNT["bold"],
                padx=18, pady=8, cursor="hand2",
                command=lambda k=key: self._toggle_periode(k),
            )
            btn.pack(side="left", padx=(0, 10))
            self._per_btns[key] = (btn, color)
        self._update_per_hint()

    # ── Période toggle ────────────────────────────────────────────
    def _toggle_periode(self, key):
        btn, color = self._per_btns[key]
        if key in self._sel_per:
            self._sel_per.discard(key)
            btn.configure(bg=C["stripe"], fg=C["text"])
        else:
            self._sel_per.add(key)
            btn.configure(bg=color, fg="white")
        self._update_per_hint()

    def _update_per_hint(self):
        all_per = self.PERIODES_TOUS
        if not self._sel_per:
            self.per_hint.configure(text="Aucune periode selectionnee",
                                    fg=C["text_m"])
        else:
            labels = [lbl for k, lbl, _ in all_per if k in self._sel_per]
            self.per_hint.configure(
                text="Selectionne : " + "  -  ".join(labels),
                fg=C["success"])

    # ── Employee cards ────────────────────────────────────────────
    def _reload_employees(self):
        for w in self.emp_grid.winfo_children():
            w.destroy()
        self._emp_frames.clear()
        self._hs_entries.clear()
        self._sel_emps.clear()
        all_emps = database.get_all_employees(actif_only=True)
        ft = self._filter_type.get()
        if ft == "jour":
            self._emp_data = [e for e in all_emps
                              if (e.get("type_ouvrier") or "jour") == "jour"]
        elif ft == "nuit":
            self._emp_data = [e for e in all_emps
                              if (e.get("type_ouvrier") or "jour") == "nuit"]
        else:
            self._emp_data = all_emps
        self._render_emp_grid()
        self._update_emp_hint()

    def _render_emp_grid(self):
        COLS = 5
        TYPE_BADGE = {"jour": ("#2563EB", "JOUR"), "nuit": ("#8B5CF6", "NUIT")}
        for idx, emp in enumerate(self._emp_data):
            row, col = divmod(idx, COLS)
            card = tk.Frame(self.emp_grid, bg=C["stripe"],
                            relief="flat", bd=0, padx=8, pady=8)
            card.grid(row=row, column=col, padx=6, pady=6, sticky="nsew")
            self.emp_grid.columnconfigure(col, weight=1)

            icone  = emp.get("icone") or "👤"
            nom    = emp.get("nom", "")
            prenom = emp.get("prenom", "")
            role   = emp.get("role", "")
            eid    = emp["id"]
            t_type = emp.get("type_ouvrier") or "jour"
            badge_clr, badge_txt = TYPE_BADGE.get(t_type, ("#2563EB", "JOUR"))

            lbl_icon = tk.Label(card, text=icone, bg=C["stripe"],
                                font=("Segoe UI", 22), cursor="hand2")
            lbl_icon.pack()
            lbl_name = tk.Label(card, text=f"{nom}\n{prenom}", bg=C["stripe"],
                                fg=C["text"], font=("Segoe UI", 9, "bold"),
                                justify="center", cursor="hand2")
            lbl_name.pack()
            lbl_role = tk.Label(card, text=role, bg=C["stripe"],
                                fg=C["text_m"], font=("Segoe UI", 8),
                                justify="center", cursor="hand2")
            lbl_role.pack()
            # Type badge
            lbl_type = tk.Label(card, text=badge_txt,
                                bg=badge_clr, fg="white",
                                font=("Segoe UI", 7, "bold"), padx=6, pady=1)
            lbl_type.pack(pady=(2, 0))

            # ── Per-employee heures supp entry (hidden until selected) ──
            hs_frame = tk.Frame(card, bg=C["stripe"])
            hs_frame.pack(pady=(6, 0))
            tk.Label(hs_frame, text="H. Supp (DT)",
                     bg=C["stripe"], fg="white",
                     font=("Segoe UI", 7)).pack()
            hs_entry = tk.Entry(
                hs_frame, width=7, justify="center",
                relief="flat", bg="white", fg=C["text"],
                font=("Segoe UI", 9, "bold"),
                highlightthickness=1,
                highlightbackground=C["border"],
                highlightcolor=C["primary"],
                insertbackground=C["text"],
            )
            hs_entry.insert(0, "0")
            hs_entry.pack(ipady=3)
            hs_frame.pack_forget()   # hidden by default
            self._hs_entries[eid] = (hs_entry, hs_frame)

            detail_btn = tk.Button(
                card, text="Fiche",
                bg=C["stripe"], fg=C["primary"],
                relief="flat", font=("Segoe UI", 8),
                cursor="hand2", padx=4, pady=1,
                activebackground=C["border"],
                command=lambda e=emp: self._open_employee_detail(e),
            )
            detail_btn.pack(pady=(4, 0))

            self._emp_frames[eid] = card

            def _click(e, emp_id=eid):
                self._toggle_employee(emp_id)
            for w in [card, lbl_icon, lbl_name, lbl_role, lbl_type]:
                w.bind("<Button-1>", _click)

    def _toggle_employee(self, emp_id):
        card = self._emp_frames.get(emp_id)
        if not card:
            return
        hs_entry, hs_frame = self._hs_entries.get(emp_id, (None, None))
        if emp_id in self._sel_emps:
            self._sel_emps.discard(emp_id)
            # Restore unselected style
            card.configure(bg=C["stripe"])
            for w in card.winfo_children():
                if isinstance(w, tk.Label):
                    w.configure(bg=C["stripe"],
                                fg=C["text_m"] if w.cget("font") and "8" in str(w.cget("font"))
                                else C["text"])
                elif isinstance(w, tk.Button):
                    w.configure(bg=C["stripe"], fg=C["primary"])
                elif isinstance(w, tk.Frame):
                    w.configure(bg=C["stripe"])
                    for ww in w.winfo_children():
                        if isinstance(ww, tk.Label):
                            ww.configure(bg=C["stripe"])
            # Hide HS entry
            if hs_frame:
                hs_frame.pack_forget()
        else:
            self._sel_emps.add(emp_id)
            # Highlight selected
            card.configure(bg=C["sidebar_sel"])
            for w in card.winfo_children():
                if isinstance(w, tk.Label):
                    w.configure(bg=C["sidebar_sel"], fg="white")
                elif isinstance(w, tk.Button):
                    w.configure(bg=C["sidebar_sel"], fg="white",
                                activebackground=C["primary_d"])
                elif isinstance(w, tk.Frame):
                    w.configure(bg=C["sidebar_sel"])
                    for ww in w.winfo_children():
                        if isinstance(ww, tk.Label):
                            ww.configure(bg=C["sidebar_sel"], fg="white")
            # Show HS entry (pack before detail_btn = last child before it)
            if hs_frame:
                hs_frame.configure(bg=C["sidebar_sel"])
                hs_frame.pack(pady=(6, 0), before=card.winfo_children()[-1])
        self._update_emp_hint()

    def _update_emp_hint(self):
        if not self._sel_emps:
            self.emp_hint.configure(text="Aucun employé sélectionné", fg=C["text_m"])
        else:
            names = [
                f"{e['nom']} {e['prenom']}"
                for e in self._emp_data if e["id"] in self._sel_emps
            ]
            self.emp_hint.configure(
                text=f"{len(names)} sélectionné(s) : " + ", ".join(names),
                fg=C["success"])

    # ── Save ──────────────────────────────────────────────────────
    def _save_pointage(self):
        self.status_label.configure(text="", fg=C["danger"])

        if not self._sel_per:
            self.status_label.configure(text="Selectionnez au moins une periode.")
            return
        if not self._sel_emps:
            self.status_label.configure(text="Selectionnez au moins un employe.")
            return

        date_str = self.date_entry.get().strip()
        if not date_str:
            self.status_label.configure(text="Date requise.")
            return
        try:
            d, m, y = date_str.strip().split("/")
            date_iso = f"{y}-{m}-{d}"
        except Exception:
            self.status_label.configure(text="Format de date invalide (JJ/MM/AAAA).")
            return

        notes = self.notes_text.get("1.0", "end").strip() or None
        saved = []

        # Track nuit employees to insert into charges_omar after pointage
        nuit_totals = {}

        for emp_id in self._sel_emps:
            emp = next((e for e in self._emp_data if e["id"] == emp_id), None)
            if not emp:
                continue

            # Read per-employee heures supp
            hs_entry, _ = self._hs_entries.get(emp_id, (None, None))
            try:
                heures_supp = float(hs_entry.get().strip() or 0) if hs_entry else 0.0
                if heures_supp < 0:
                    raise ValueError
            except ValueError:
                self.status_label.configure(
                    text=f"Heures supp. invalides pour {emp.get('nom','')} (nombre >= 0).")
                return
                continue
            prix_hs   = float(emp.get("prix_heure_supp") or 0)
            is_nuit   = (emp.get("type_ouvrier") or "jour") == "nuit"
            TAR_EMP   = {
                "midi":             float(emp.get("salaire_midi")       or 0),
                "apres_midi":       float(emp.get("salaire_apres_midi") or 0),
                "soiree":           float(emp.get("salaire_soiree")     or 0),
                "journee_complete": float(emp.get("salaire_journalier") or 0),
            }
            for periode in sorted(self._sel_per):
                data = {
                    "employee_id":             emp_id,
                    "date_pointage":           date_iso,
                    "heure_arrivee":           None,
                    "heure_depart":            None,
                    "heures_supplementaires":  heures_supp,
                    "periode":                 periode,
                    "notes":                   notes,
                }
                database.insert_pointage(data)
                hs_txt = f" + {heures_supp}h supp." if heures_supp > 0 else ""
                saved.append(f"{emp['nom']} {emp['prenom']} / {periode}{hs_txt}")

                # Accumulate for nuit → charges_omar
                if is_nuit:
                    tarif  = TAR_EMP.get(periode, 0.0)
                    hs_amt = heures_supp * prix_hs
                    if emp_id not in nuit_totals:
                        nuit_totals[emp_id] = {"montant": 0.0, "emp": emp}
                    nuit_totals[emp_id]["montant"] += tarif + hs_amt

        # Insert ONE charges_omar row for the entire nuit session (all employees merged)
        omar_entries = []
        if nuit_totals:
            total_nuit = sum(info["montant"] for info in nuit_totals.values())
            if total_nuit > 0:
                # Build designation: notes if provided, else list employee names
                if notes:
                    desig = notes
                else:
                    noms = [
                        f"{info['emp']['nom']} {info['emp']['prenom']}"
                        for info in nuit_totals.values()
                        if info["montant"] > 0
                    ]
                    desig = f"Pointage {' / '.join(noms)} — {date_str}"
                database.insert_charge_omar(date_iso, desig, total_nuit)
                omar_entries.append(f"{desig} : {total_nuit:.3f} DT")

        nb = len(saved)
        msg = f"{nb} entree(s) creee(s) :\n" + "\n".join(f"  - {s}" for s in saved)
        if omar_entries:
            msg += "\n\nAjoute au Journal Caisse Omar :\n" + "\n".join(
                f"  - {x}" for x in omar_entries)

        self.status_label.configure(
            text=f"{nb} pointage(s) enregistre(s) pour le {date_str}."
                 + (f" | {len(omar_entries)} entree(s) Caisse Omar." if omar_entries else ""),
            fg=C["success"])
        toast(msg, "success", "Pointage enregistré")
        self._reset()

    def _reset(self):
        # Deselect all periods
        for key, (btn, color) in self._per_btns.items():
            btn.configure(bg=C["stripe"], fg=C["text"])
        self._sel_per.clear()
        self._update_per_hint()
        # Deselect all employees
        for eid, card in self._emp_frames.items():
            card.configure(bg=C["stripe"])
            for w in card.winfo_children():
                if isinstance(w, tk.Label):
                    w.configure(bg=C["stripe"],
                                fg=C["text_m"] if "8" in str(w.cget("font"))
                                else C["text"])
                elif isinstance(w, tk.Button):
                    w.configure(bg=C["stripe"], fg=C["primary"],
                                activebackground=C["border"])
        self._sel_emps.clear()
        self._update_emp_hint()
        # Reset date, heures supp entries, notes
        self.date_entry.delete(0, "end")
        self.date_entry.insert(0, date.today().strftime("%d/%m/%Y"))
        for hs_entry, hs_frame in self._hs_entries.values():
            hs_entry.delete(0, "end")
            hs_entry.insert(0, "0")
            hs_frame.pack_forget()
        self.notes_text.delete("1.0", "end")
        self.status_label.configure(text="")

    def _open_employee_detail(self, emp):
        """Popup : pointages + total à payer + enregistrement du paiement."""
        from datetime import datetime as _dt
        eid  = emp["id"]
        nom  = f"{emp.get('icone','👤')}  {emp.get('nom','')} {emp.get('prenom','')}"
        role = emp.get("role", "")

        # ── Tarifs ──────────────────────────────────────────────
        TAR = {
            "midi":             float(emp.get("salaire_midi")         or 0),
            "apres_midi":       float(emp.get("salaire_apres_midi")   or 0),
            "soiree":           float(emp.get("salaire_soiree")       or 0),
            "journee_complete": float(emp.get("salaire_journalier")   or 0),
        }
        PER_LABEL = {
            "midi":             "☀ Midi",
            "apres_midi":       "🌤 Après-midi",
            "soiree":           "🌙 Soirée",
            "journee_complete": "📅 Journée complète",
        }

        # ── Window ──────────────────────────────────────────────
        win = tk.Toplevel(self)
        win.title(f"Détail — {emp.get('nom','')} {emp.get('prenom','')}")
        win.configure(bg=C["bg"])
        win.grab_set()
        # Screen-relative: 65% wide, 75% tall
        win.update_idletasks()
        _sw = win.winfo_screenwidth()
        _sh = win.winfo_screenheight()
        _w = max(760, int(_sw * 0.65))
        _h = max(560, int(_sh * 0.75))
        win.geometry(f"{_w}x{_h}+{(_sw - _w)//2}+{(_sh - _h)//2}")

        # Header
        hdr = tk.Frame(win, bg=C["sidebar"], pady=0)
        hdr.pack(fill="x")
        tk.Label(hdr, text=nom, bg=C["sidebar"], fg="white",
                 font=("Segoe UI", 13, "bold")).pack(side="left", padx=20, pady=14)
        tk.Label(hdr, text=role, bg=C["sidebar"], fg="#94A3B8",
                 font=("Segoe UI", 10)).pack(side="left", padx=(0, 20), pady=14)

        # ── Period filter ────────────────────────────────────────
        filt_fr = tk.Frame(win, bg=C["surface"])
        filt_fr.pack(fill="x")
        filt_in = tk.Frame(filt_fr, bg=C["surface"])
        filt_in.pack(fill="x", padx=20, pady=10)

        tk.Label(filt_in, text="Période du", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(side="left")
        d_debut_e = RoundedEntry(filt_in, width=13)
        d_debut_e.pack(side="left", padx=(6, 0), ipady=3)

        tk.Label(filt_in, text="  au", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(side="left")
        d_fin_e = RoundedEntry(filt_in, width=13)
        d_fin_e.pack(side="left", padx=(6, 16), ipady=3)

        tk.Frame(win, bg=C["border"], height=1).pack(fill="x")

        # ── Treeview ─────────────────────────────────────────────
        tree_wrap = tk.Frame(win, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=10)

        cols_t = ("date", "periode", "heures_supp", "tarif", "notes")
        tree = ttk.Treeview(tree_wrap, columns=cols_t, show="headings")
        heads_t = {"date": "Date", "periode": "Periode",
                   "heures_supp": "H. Supp.", "tarif": "Tarif (DT)", "notes": "Notes"}
        widths_t = {"date": 105, "periode": 140, "heures_supp": 80, "tarif": 110, "notes": 0}
        for c in cols_t:
            tree.heading(c, text=heads_t[c])
            tree.column(c, width=widths_t[c] if widths_t[c] else 300,
                        stretch=(widths_t[c] == 0),
                        anchor="e" if c in ("heures_supp", "tarif") else "w")
        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        tree.pack(fill="both", expand=True)

        # ── Summary bar ──────────────────────────────────────────
        sum_fr = tk.Frame(win, bg=C["surface"])
        sum_fr.pack(fill="x")
        tk.Frame(sum_fr, bg=C["border"], height=1).pack(fill="x")
        sum_in = tk.Frame(sum_fr, bg=C["surface"])
        sum_in.pack(fill="x", padx=20, pady=10)

        lbl_nb  = tk.Label(sum_in, text="0 pointage(s)", bg=C["surface"],
                           fg=C["text_m"], font=FNT["sm"])
        lbl_nb.pack(side="left", padx=(0, 24))
        lbl_tot = tk.Label(sum_in, text="Total calculé : 0.000 DT",
                           bg=C["surface"], fg=C["primary"],
                           font=("Segoe UI", 11, "bold"))
        lbl_tot.pack(side="left")

        # ── Payment section ──────────────────────────────────────
        pay_fr = tk.Frame(win, bg=C["card_hdr"])
        pay_fr.pack(fill="x")
        tk.Frame(pay_fr, bg=C["success"], height=2).pack(fill="x")
        pay_in = tk.Frame(pay_fr, bg=C["card_hdr"])
        pay_in.pack(fill="x", padx=20, pady=12)

        tk.Label(pay_in, text="💰  Montant à payer (DT)",
                 bg=C["card_hdr"], fg=C["text_m"], font=FNT["sm"]).pack(
            side="left")
        pay_entry = RoundedEntry(pay_in, width=14, bg_canvas=C["card_hdr"])
        pay_entry.pack(side="left", padx=(8, 16), ipady=4)

        tk.Label(pay_in, text="Date paiement",
                 bg=C["card_hdr"], fg=C["text_m"], font=FNT["sm"]).pack(side="left")
        pay_date_e = RoundedEntry(pay_in, width=13, bg_canvas=C["card_hdr"])
        pay_date_e.insert(0, date.today().strftime("%d/%m/%Y"))
        pay_date_e.pack(side="left", padx=(8, 16), ipady=4)

        tk.Label(pay_in, text="Notes",
                 bg=C["card_hdr"], fg=C["text_m"], font=FNT["sm"]).pack(side="left")
        pay_notes_e = RoundedEntry(pay_in, width=22, bg_canvas=C["card_hdr"])
        pay_notes_e.pack(side="left", padx=(8, 16), ipady=4)

        pay_status = tk.Label(pay_fr, text="", bg=C["card_hdr"],
                              fg=C["danger"], font=FNT["sm"])
        pay_status.pack(anchor="w", padx=20, pady=(0, 8))

        # ── Bottom buttons ───────────────────────────────────────
        bot = tk.Frame(win, bg=C["surface"])
        bot.pack(fill="x")
        tk.Frame(bot, bg=C["border"], height=1).pack(fill="x")
        bot_in = tk.Frame(bot, bg=C["surface"])
        bot_in.pack(fill="x", padx=20, pady=10)

        # ── Helper : populate tree & totals ─────────────────────
        _total_holder = [0.0]   # mutable to update from inner func
        _hs_holder    = [0.0]   # total heures supp amount

        def _fill(start_iso=None, end_iso=None):
            for item in tree.get_children():
                tree.delete(item)
            rows = database.get_pointage_by_employee(eid, start_iso, end_iso)
            total        = 0.0
            total_hs_amt = 0.0
            prix_hs      = float(emp.get("prix_heure_supp") or 0)
            for r in rows:
                per   = r.get("periode") or ""
                tarif = TAR.get(per, 0.0)
                hs    = float(r.get("heures_supplementaires") or 0)
                hs_amt = hs * prix_hs
                total        += tarif + hs_amt
                total_hs_amt += hs_amt
                try:
                    d_fmt = _dt.strptime(r["date_pointage"], "%Y-%m-%d").strftime("%d/%m/%Y")
                except Exception:
                    d_fmt = r.get("date_pointage", "")
                tarif_total = tarif + hs_amt
                hs_str = f"{hs:.1f}h" if hs > 0 else "—"
                tree.insert("", "end", values=(
                    d_fmt,
                    PER_LABEL.get(per, per),
                    hs_str,
                    f"{tarif_total:.3f}" if tarif_total else "—",
                    r.get("notes") or "",
                ))
            _total_holder[0] = total
            _hs_holder[0]    = total_hs_amt
            hs_info = f"  (dont {total_hs_amt:.3f} DT h.supp)" if total_hs_amt > 0 else ""
            lbl_nb.configure(text=f"{len(rows)} pointage(s)")
            lbl_tot.configure(text=f"Total calcule : {total:.3f} DT{hs_info}")
            pay_entry.delete(0, "end")
            pay_entry.insert(0, f"{total:.3f}")

        def _apply_filter():
            s = d_debut_e.get().strip()
            e2 = d_fin_e.get().strip()
            try:
                si = f"{s[6:10]}-{s[3:5]}-{s[0:2]}" if s else None
                ei = f"{e2[6:10]}-{e2[3:5]}-{e2[0:2]}" if e2 else None
            except Exception:
                si = ei = None
            _fill(si, ei)

        def _save_payment():
            pay_status.configure(text="", fg=C["danger"])
            raw_m = pay_entry.get().strip()
            raw_d = pay_date_e.get().strip()
            if not raw_m:
                pay_status.configure(text="Saisissez un montant.")
                return
            try:
                montant = float(raw_m)
                if montant <= 0:
                    raise ValueError
            except ValueError:
                pay_status.configure(text="Montant invalide.")
                return
            try:
                d_pay_iso = f"{raw_d[6:10]}-{raw_d[3:5]}-{raw_d[0:2]}"
            except Exception:
                pay_status.configure(text="Date invalide (JJ/MM/AAAA).")
                return
            # Determine period bounds from filter fields
            s = d_debut_e.get().strip()
            e2 = d_fin_e.get().strip()
            try:
                p_debut = f"{s[6:10]}-{s[3:5]}-{s[0:2]}" if s else "2000-01-01"
                p_fin   = f"{e2[6:10]}-{e2[3:5]}-{e2[0:2]}" if e2 else date.today().isoformat()
            except Exception:
                p_debut = "2000-01-01"
                p_fin   = date.today().isoformat()
            data_pay = {
                "employee_id":         eid,
                "date_paiement":       d_pay_iso,
                "periode_debut":       p_debut,
                "periode_fin":         p_fin,
                "montant_base":        _total_holder[0] - _hs_holder[0],
                "montant_heures_supp": _hs_holder[0],
                "montant_total":       montant,
                "notes":               pay_notes_e.get().strip() or None,
            }
            database.insert_paiement_employee(data_pay)
            pay_status.configure(
                text=f"Paiement de {montant:.3f} DT enregistre.",
                fg=C["success"])
            messagebox.showinfo("Paiement enregistre",
                                f"Paiement de {montant:.3f} DT enregistre.")

        # Wire buttons
        tk.Button(bot_in, text="🔍  Appliquer filtre",
                  bg=C["accent"], fg="white",
                  relief="flat", font=FNT["bold"], padx=14, pady=6,
                  cursor="hand2", command=_apply_filter).pack(side="left", padx=(0, 8))
        tk.Button(bot_in, text="💾  Enregistrer paiement",
                  bg=C["success"], fg="white",
                  relief="flat", font=FNT["bold"], padx=14, pady=6,
                  cursor="hand2", command=_save_payment).pack(side="left", padx=(0, 8))
        tk.Button(bot_in, text="Fermer",
                  bg=C["border"], fg=C["text"],
                  relief="flat", font=FNT["bold"], padx=14, pady=6,
                  cursor="hand2", command=win.destroy).pack(side="right")

        # Initial load (all pointages)
        _fill()

    def _view_history(self):
        """Fenetre modale : derniers 100 pointages avec type ouvrier et h.supp."""
        win = tk.Toplevel(self)
        win.title("Historique des pointages")
        win.geometry("900x480")
        win.configure(bg=C["bg"])
        win.grab_set()

        # Filter row
        top_fr = tk.Frame(win, bg=C["bg"])
        top_fr.pack(fill="x", padx=20, pady=(16, 4))
        tk.Label(top_fr, text="Historique des pointages", bg=C["bg"],
                 fg=C["primary"], font=FNT["title"]).pack(side="left")

        # Type filter pills
        pill_fr = tk.Frame(win, bg=C["surface"])
        pill_fr.pack(fill="x", padx=20, pady=(0, 6))
        pill_in = tk.Frame(pill_fr, bg=C["surface"])
        pill_in.pack(fill="x", padx=12, pady=8)
        tk.Label(pill_in, text="Filtrer :", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(side="left", padx=(0, 8))

        hist_filter = tk.StringVar(value="tous")
        hist_pills  = {}

        cols = ("date", "employe", "type", "periode", "heures_supp", "notes")
        tree = ttk.Treeview(win, columns=cols, show="headings")
        heads = {"date": "Date", "employe": "Employe", "type": "Type",
                 "periode": "Periode", "heures_supp": "H.Supp", "notes": "Notes"}
        widths = {"date": 95, "employe": 160, "type": 80,
                  "periode": 140, "heures_supp": 70, "notes": 0}
        for c in cols:
            tree.heading(c, text=heads[c])
            tree.column(c, width=widths[c] if widths[c] else 260,
                        stretch=(widths[c] == 0), anchor="w")
        tree.tag_configure("jour", background=_lighten(C["primary"]))
        tree.tag_configure("nuit", background=_lighten("#8B5CF6"))

        def _load_hist(ftype="tous"):
            for item in tree.get_children():
                tree.delete(item)
            from datetime import datetime as _dt2
            rows = database.get_all_pointage()
            for r in rows:
                emp_type = r.get("type_ouvrier") or "jour"
                if ftype != "tous" and emp_type != ftype:
                    continue
                try:
                    d_fmt = _dt2.strptime(r["date_pointage"], "%Y-%m-%d").strftime("%d/%m/%Y")
                except Exception:
                    d_fmt = r.get("date_pointage", "")
                nom  = f"{r.get('nom', '')} {r.get('prenom', '')}".strip()
                hs   = float(r.get("heures_supplementaires") or 0)
                hs_s = f"{hs:.1f}h" if hs > 0 else "—"
                tree.insert("", "end", tags=(emp_type,), values=(
                    d_fmt, nom,
                    "Jour" if emp_type == "jour" else "Nuit",
                    r.get("periode", ""),
                    hs_s,
                    r.get("notes", "") or ""))

        HOPTS = [("tous","Tous","#64748B"),("jour","Ouvrier Jour",C["primary"]),
                 ("nuit","Ouvrier Nuit","#8B5CF6")]
        for val, lbl, clr in HOPTS:
            is_sel = (val == "tous")
            b = tk.Button(pill_in, text=lbl,
                          bg=clr if is_sel else C["stripe"],
                          fg="white" if is_sel else C["text"],
                          relief="flat", font=FNT["sm"],
                          padx=12, pady=4, cursor="hand2")
            b.pack(side="left", padx=(0, 6))
            hist_pills[val] = (b, clr)

            def _set_filter(v=val, c=clr):
                hist_filter.set(v)
                for vv, (bb, cc) in hist_pills.items():
                    bb.configure(bg=cc if vv == v else C["stripe"],
                                 fg="white" if vv == v else C["text"])
                _load_hist(v)
            b.configure(command=_set_filter)

        vsb = ttk.Scrollbar(win, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y", padx=(0, 8), pady=8)
        tree.pack(fill="both", expand=True, padx=(20, 0), pady=(0, 8))

        tk.Button(win, text="Fermer", command=win.destroy,
                  bg=C["primary"], fg="white", relief="flat",
                  font=FNT["bold"], padx=20, pady=6).pack(pady=(0, 12))

        _load_hist()


# ═══════════════════════════════════════════════════════════════════
#  TAB — CENTRE DE RÉCEPTION
#  Affiche tous les acomptes Caisse Omar (en attente / accepté / refusé).
#  Omar accepte ou refuse ici ; seuls les "accepté" apparaissent
#  dans le Journal de Caisse Omar.
# ═══════════════════════════════════════════════════════════════════
class CentreReceptionTab(tk.Frame):

    # couleurs badge statut
    _STATUT_CFG = {
        "en_attente": {"label": "En attente", "bg": "#FEF3C7", "fg": "#92400E"},
        "accepte":    {"label": "Accepté ✔",  "bg": "#D1FAE5", "fg": "#065F46"},
        "refuse":     {"label": "Refusé ✘",   "bg": "#FEE2E2", "fg": "#991B1B"},
    }

    def __init__(self, parent, app):
        super().__init__(parent, bg=C["bg"])
        self.app = app
        self._build()
        self.refresh()

    # ─────────────────────────────────────────────────────────────
    def _build(self):
        # ── Header ───────────────────────────────────────────────
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        h_in = tk.Frame(hdr, bg=C["surface"])
        h_in.pack(fill="x", padx=20, pady=14)
        tk.Label(h_in, text="Centre de Réception — Caisse Omar",
                 bg=C["surface"], fg=C["primary"], font=FNT["title"]).pack(side="left")
        self.count_lbl = tk.Label(h_in, text="", bg=C["surface"],
                                  fg=C["text_m"], font=FNT["sm"])
        self.count_lbl.pack(side="left", padx=16)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # ── Filtre statut ─────────────────────────────────────────
        filt = tk.Frame(self, bg=C["surface"])
        filt.pack(fill="x")
        f_in = tk.Frame(filt, bg=C["surface"])
        f_in.pack(fill="x", padx=20, pady=10)
        tk.Label(f_in, text="Afficher :", bg=C["surface"],
                 fg=C["text_m"], font=FNT["sm"]).pack(side="left", padx=(0, 8))
        self.filtre_var = tk.StringVar(value="tous")
        for val, label in [("tous",       "Tous"),
                            ("en_attente", "En attente"),
                            ("accepte",    "Acceptés"),
                            ("refuse",     "Refusés")]:
            tk.Radiobutton(
                f_in, text=label, variable=self.filtre_var, value=val,
                bg=C["surface"], activebackground=C["surface"],
                fg=C["text"], selectcolor=C["surface"],
                font=FNT["sm"], cursor="hand2",
                command=self.refresh,
            ).pack(side="left", padx=(0, 12))
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")

        # ── Toolbar ───────────────────────────────────────────────
        tb = tk.Frame(self, bg=C["bg"])
        tb.pack(fill="x", padx=20, pady=(8, 4))
        _rbtn(tb, "🔄  Actualiser",         command=self.refresh,
              style="secondary", padx=16, pady=7)
        _rbtn(tb, "✔  Accepter",            command=self._accepter,
              style="success",   padx=16, pady=7)
        _rbtn(tb, "✘  Refuser",             command=self._refuser,
              style="sm_danger", padx=16, pady=7)
        _rbtn(tb, "↺  Remettre en attente", command=self._remettre,
              style="secondary", padx=16, pady=7)

        # ── Treeview ──────────────────────────────────────────────
        tree_wrap = tk.Frame(self, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True, padx=20, pady=(4, 12))

        cols = ("res_id", "num_bon", "client", "salle", "date_evt",
                "date_acompte", "montant_omar", "methode", "statut", "date_accept")
        self.tree = ttk.Treeview(tree_wrap, columns=cols,
                                 show="headings", selectmode="extended")
        heads = {
            "res_id":       "ID",
            "num_bon":      "N° Bon",
            "client":       "Client",
            "salle":        "Salle",
            "date_evt":     "Date évén.",
            "date_acompte": "Date acompte",
            "montant_omar": "Montant Omar (DT)",
            "methode":      "Méthode",
            "statut":       "Statut",
            "date_accept":  "Date acceptation",
        }
        widths = {
            "res_id": 45, "num_bon": 110, "client": 160, "salle": 90,
            "date_evt": 90, "date_acompte": 100, "montant_omar": 130,
            "methode": 90, "statut": 110, "date_accept": 110,
        }
        for c in cols:
            self.tree.heading(c, text=heads[c],
                              command=lambda _c=c: self._sort(_c))
            self.tree.column(c, width=widths[c], anchor="center",
                             stretch=(c == "client"))

        # row colour tags
        self.tree.tag_configure("en_attente", background="#FEF08A", foreground="#713F12")
        self.tree.tag_configure("accepte",    background="#BBF7D0", foreground="#14532D")
        self.tree.tag_configure("refuse",     background="#FECACA", foreground="#7F1D1D")
        self.tree.tag_configure("date_sep",   background=C["sidebar"],
                                foreground="white", font=FNT["bold"])

        vsb = ttk.Scrollbar(tree_wrap, orient="vertical",
                            command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)

        # ── Info banner ───────────────────────────────────────────
        info_fr = tk.Frame(self, bg=C["stripe"])
        info_fr.pack(fill="x", padx=20, pady=(0, 6))
        tk.Frame(info_fr, bg=C["primary"], width=4).pack(side="left", fill="y")
        self.info_lbl = tk.Label(
            info_fr,
            text="Sélectionnez une ou plusieurs lignes, puis choisissez Accepter ou Refuser.",
            bg=C["stripe"], fg=C["text_m"], font=FNT["sm"],
            anchor="w", padx=12, pady=7,
        )
        self.info_lbl.pack(side="left", fill="x", expand=True)

        # ── Status bar ────────────────────────────────────────────
        self.status_lbl = tk.Label(self, text="", bg=C["bg"],
                                   fg=C["danger"], font=FNT["sm"])
        self.status_lbl.pack(padx=20, pady=(0, 8), anchor="w")

        self._sort_col = "date_acompte"
        self._sort_rev = False

    # ─────────────────────────────────────────────────────────────
    def refresh(self):
        self._populate(database.get_acomptes_caisse_omar())

    def _populate(self, rows):
        for item in self.tree.get_children():
            self.tree.delete(item)

        filtre = self.filtre_var.get()
        filtered = [r for r in rows if filtre == "tous" or r["statut"] == filtre]

        def _key(r):
            return r.get(self._sort_col) or ""
        filtered.sort(key=_key, reverse=self._sort_rev)

        current_date = None
        for r in filtered:
            # Date separator when grouping by date_acompte
            row_date = r.get("date_acompte") or ""
            if row_date != current_date:
                current_date = row_date
                try:
                    from datetime import datetime as _dt3, date as _date3
                    d = _date3.fromisoformat(row_date)
                    day_names = ["LUNDI","MARDI","MERCREDI","JEUDI","VENDREDI","SAMEDI","DIMANCHE"]
                    sep_lbl = f"── {day_names[d.weekday()]}  {d.strftime('%d/%m/%Y')} ──"
                except Exception:
                    sep_lbl = f"── {row_date} ──"
                self.tree.insert("", "end", tags=("date_sep",), values=(
                    "", sep_lbl, "", "", "", "", "", "", "", "",
                ))

            st  = r["statut"]
            cfg = self._STATUT_CFG.get(st, self._STATUT_CFG["en_attente"])
            try:
                from datetime import datetime as _dt2
                date_evt_fmt = _dt2.strptime(
                    r["date_evenement"], "%Y-%m-%d").strftime("%d/%m/%Y")
            except Exception:
                date_evt_fmt = r["date_evenement"]

            iid = f"{r['res_id']}_{r['slot']}"
            self.tree.insert("", "end", iid=iid, tags=(st,), values=(
                r["res_id"],
                r["num_bon"],
                r["nom_client"],
                r["salle"],
                date_evt_fmt,
                r["date_acompte_fmt"],
                f"{r['montant_omar']:,.3f}",
                r["methode"] or "Espèce",
                cfg["label"],
                r["date_accept_fmt"],
            ))

        pending  = sum(1 for r in rows if r["statut"] == "en_attente")
        accepted = sum(1 for r in rows if r["statut"] == "accepte")
        refused  = sum(1 for r in rows if r["statut"] == "refuse")
        self.count_lbl.configure(
            text=(f"En attente : {pending}   |   "
                  f"Acceptés : {accepted}   |   Refusés : {refused}"))
        self.status_lbl.configure(text="")

    def _sort(self, col):
        self._sort_rev = (not self._sort_rev) if self._sort_col == col else False
        self._sort_col = col
        self.refresh()

    # ─────────────────────────────────────────────────────────────
    def _selected_items(self):
        items = []
        for iid in self.tree.selection():
            try:
                rid, slot = iid.split("_")
                items.append((int(rid), int(slot)))
            except Exception:
                pass
        return items

    def _apply_statut(self, new_statut):
        sel = self._selected_items()
        if not sel:
            self.status_lbl.configure(
                text="Sélectionnez au moins une ligne.", fg=C["danger"])
            return
        for res_id, slot in sel:
            database.update_omar_statut(res_id, slot, new_statut)
        label_map = {
            "accepte": "accepté(s)",
            "refuse":  "refusé(s)",
            None:      "remis en attente",
        }
        self.status_lbl.configure(
            text=f"✔  {len(sel)} acompte(s) {label_map.get(new_statut, new_statut)}.",
            fg=C["success"])
        self.refresh()

    def _accepter(self):  self._apply_statut("accepte")
    def _refuser(self):   self._apply_statut("refuse")
    def _remettre(self):  self._apply_statut(None)


# ═══════════════════════════════════════════════════════════════════
#  FENÊTRE DE GESTION DES ACCÈS (ADMIN ONLY)
# ═══════════════════════════════════════════════════════════════════
class AccessManagementWindow(tk.Toplevel):
    """Fenêtre de gestion des utilisateurs et de leurs accès aux modules."""
    
    def __init__(self, parent):
        super().__init__(parent)
        self.title("Gestion des accès")
        self.minsize(900, 600)
        self.configure(bg=C["bg"])
        self.grab_set()
        # Screen-relative: 80% wide, 80% tall
        self.update_idletasks()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        w = max(900, int(sw * 0.80))
        h = max(600, int(sh * 0.80))
        x = (sw - w) // 2
        y = (sh - h) // 2
        self.geometry(f"{w}x{h}+{x}+{y}")
        
        self.selected_user_id = None
        
        self._build_ui()
        self._refresh_user_list()
    
    def _build_ui(self):
        # ── Header ───────────────────────────────────────────────
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.pack(fill="x")
        tk.Label(hdr, text="Gestion des accès", bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x")
        
        # ── Main layout : list + details ─────────────────────────
        main = tk.Frame(self, bg=C["bg"])
        main.pack(fill="both", expand=True, padx=20, pady=16)
        main.columnconfigure(0, weight=1)
        main.columnconfigure(1, weight=2)
        main.rowconfigure(0, weight=1)
        
        # ── LEFT : Liste des utilisateurs ───────────────────────
        left = tk.Frame(main, bg=C["bg"])
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        
        left_hdr = tk.Frame(left, bg=C["surface"])
        left_hdr.pack(fill="x")
        tk.Frame(left_hdr, bg=C["primary"], height=2).pack(fill="x")
        tk.Label(left_hdr, text="Utilisateurs", bg=C["card_hdr"],
                 fg=C["primary"], font=FNT["bold"]).pack(
            anchor="w", padx=14, pady=8)
        tk.Frame(left_hdr, bg=C["border"], height=1).pack(fill="x")
        
        # Toolbar
        toolbar = tk.Frame(left, bg=C["bg"])
        toolbar.pack(fill="x", pady=(8, 4))
        _rbtn(toolbar, "Ajouter", command=self._add_user, style="sm", padx=10, pady=5)
        _rbtn(toolbar, "Modifier", command=self._edit_user, style="secondary", padx=10, pady=5)
        _rbtn(toolbar, "Supprimer", command=self._delete_user, style="sm_danger", padx=8, pady=4)
        
        # Treeview utilisateurs
        tree_wrap = tk.Frame(left, bg=C["bg"])
        tree_wrap.pack(fill="both", expand=True)
        
        cols = ("id", "username", "nom", "role")
        self.user_tree = ttk.Treeview(tree_wrap, columns=cols, show="headings")
        heads = {"id": "ID", "username": "Login", "nom": "Nom complet", "role": "Rôle"}
        widths = {"id": 40, "username": 100, "nom": 0, "role": 80}
        for c in cols:
            self.user_tree.heading(c, text=heads[c])
            self.user_tree.column(c, width=widths[c] if widths[c] else 180,
                                  stretch=(widths[c] == 0), anchor="w")
        
        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.user_tree.yview)
        self.user_tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.user_tree.pack(fill="both", expand=True)
        
        self.user_tree.bind("<<TreeviewSelect>>", self._on_user_select)
        
        # ── RIGHT : Détails et modules ───────────────────────────
        right = tk.Frame(main, bg=C["bg"])
        right.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        
        right_hdr = tk.Frame(right, bg=C["surface"])
        right_hdr.pack(fill="x")
        tk.Frame(right_hdr, bg=C["accent"], height=2).pack(fill="x")
        tk.Label(right_hdr, text="Modules accessibles", bg=C["card_hdr"],
                 fg=C["accent"], font=FNT["bold"]).pack(
            anchor="w", padx=14, pady=8)
        tk.Frame(right_hdr, bg=C["border"], height=1).pack(fill="x")
        
        self.detail_label = tk.Label(right, text="Sélectionnez un utilisateur",
                                     bg=C["stripe"], fg=C["text_m"],
                                     font=FNT["sm"], anchor="w", padx=14, pady=8)
        self.detail_label.pack(fill="x", pady=(8, 0))
        
        # Scrollable module checkboxes
        scroll_outer, scroll_inner = _scrollable(right)
        scroll_outer.pack(fill="both", expand=True, pady=(8, 0))
        
        self.module_frame = scroll_inner
        self.module_vars = {}
        
        # Boutons d'action
        btn_bar = tk.Frame(right, bg=C["bg"])
        btn_bar.pack(fill="x", pady=(8, 0))
        _rbtn(btn_bar, "Enregistrer les modifications", command=self._save_modules,
              style="primary", padx=14, pady=7)
        _rbtn(btn_bar, "Appliquer template rôle", command=self._apply_role_template,
              style="secondary", padx=14, pady=7)
        
        self.status_label = tk.Label(right, text="", bg=C["bg"],
                                     fg=C["danger"], font=FNT["sm"])
        self.status_label.pack(anchor="w", pady=(4, 0))
    
    def _refresh_user_list(self):
        """Rafraîchit la liste des utilisateurs."""
        for item in self.user_tree.get_children():
            self.user_tree.delete(item)
        
        users = access_control.get_all_users(actif_only=False)
        for u in users:
            nom_complet = f"{u.get('nom', '')} {u.get('prenom', '')}".strip()
            self.user_tree.insert("", "end", iid=str(u["id"]), values=(
                u["id"], u["username"], nom_complet, u["role"]
            ))
    
    def _on_user_select(self, event=None):
        """Quand un utilisateur est sélectionné."""
        sel = self.user_tree.selection()
        if not sel:
            return
        
        self.selected_user_id = int(sel[0])
        user = access_control.get_user(self.selected_user_id)
        if not user:
            return
        
        nom_complet = f"{user.get('nom', '')} {user.get('prenom', '')}".strip()
        self.detail_label.configure(
            text=f"👤 {nom_complet}  —  @{user['username']}  —  {user['role']}",
            fg=C["text"])
        
        # Charger les modules
        self._load_user_modules()
    
    def _load_user_modules(self):
        """Charge les modules de l'utilisateur sélectionné."""
        # Effacer les anciens checkboxes
        for w in self.module_frame.winfo_children():
            w.destroy()
        self.module_vars.clear()
        
        if not self.selected_user_id:
            return
        
        perms = access_control.get_all_user_module_permissions(self.selected_user_id)
        
        for module in access_control.ALL_MODULES:
            var = tk.IntVar(value=1 if perms.get(module, False) else 0)
            self.module_vars[module] = var

            cb = tk.Checkbutton(
                self.module_frame, text=module,
                variable=var,
                onvalue=1, offvalue=0,
                bg=C["bg"], fg=C["text"],
                activebackground=C["bg"],
                selectcolor="white", font=FNT["default"]
            )
            cb.pack(anchor="w", pady=2, padx=10)
    
    def _save_modules(self):
        """Enregistre les modifications des modules."""
        self.status_label.configure(text="", fg=C["danger"])
        
        if not self.selected_user_id:
            self.status_label.configure(text="Aucun utilisateur sélectionné")
            return
        
        selected_modules = [
            module for module, var in self.module_vars.items()
            if var.get() == 1
        ]
        
        access_control.set_user_modules(self.selected_user_id, selected_modules)
        
        self.status_label.configure(
            text=f"✔ Modifications enregistrées ({len(selected_modules)} modules)",
            fg=C["success"])
    
    def _apply_role_template(self):
        """Applique un template de permissions basé sur le rôle."""
        if not self.selected_user_id:
            return
        
        user = access_control.get_user(self.selected_user_id)
        if not user:
            return
        
        role = user.get('role', '')
        if messagebox.askyesno("Appliquer template",
                               f"Appliquer les permissions par défaut du rôle '{role}' ?\n"
                               "Cela remplacera les permissions actuelles."):
            access_control.apply_role_template(self.selected_user_id, role)
            self._load_user_modules()
            self.status_label.configure(
                text=f"✔ Template '{role}' appliqué", fg=C["success"])
    
    def _add_user(self):
        """Ajoute un nouvel utilisateur."""
        UserEditWindow(self, None, on_save=self._refresh_user_list)
    
    def _edit_user(self):
        """Modifie l'utilisateur sélectionné."""
        if not self.selected_user_id:
            messagebox.showwarning("Sélection requise",
                                   "Sélectionnez un utilisateur à modifier")
            return
        
        UserEditWindow(self, self.selected_user_id, on_save=self._refresh_user_list)
    
    def _delete_user(self):
        """Supprime (désactive) l'utilisateur sélectionné."""
        if not self.selected_user_id:
            messagebox.showwarning("Sélection requise",
                                   "Sélectionnez un utilisateur à supprimer")
            return
        
        user = access_control.get_user(self.selected_user_id)
        if not user:
            return
        
        if messagebox.askyesno("Confirmer la suppression",
                               f"Désactiver l'utilisateur @{user['username']} ?"):
            access_control.delete_user(self.selected_user_id)
            self._refresh_user_list()
            self.selected_user_id = None
            self.detail_label.configure(
                text="Sélectionnez un utilisateur", fg=C["text_m"])
            for w in self.module_frame.winfo_children():
                w.destroy()


# ═══════════════════════════════════════════════════════════════════
#  FENÊTRE D'ÉDITION D'UTILISATEUR
# ═══════════════════════════════════════════════════════════════════
class UserEditWindow(tk.Toplevel):
    """Fenêtre pour créer ou modifier un utilisateur."""
    
    def __init__(self, parent, user_id=None, on_save=None):
        super().__init__(parent)
        self.user_id = user_id
        self.on_save = on_save
        
        title = "Modifier l'utilisateur" if user_id else "Nouvel utilisateur"
        self.title(title)
        self.resizable(True, True)
        self.minsize(400, 440)
        self.configure(bg=C["bg"])
        self.grab_set()

        # Screen-relative size: ~35% wide, ~80% tall
        self.update_idletasks()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        w = max(460, min(560, int(sw * 0.35)))
        h = max(500, min(700, int(sh * 0.80)))
        x = (sw // 2) - (w // 2)
        y = max(0, (sh // 2) - (h // 2))
        self.geometry(f"{w}x{h}+{x}+{y}")

        self._build_ui()
        
        if user_id:
            self._load_user_data()
    
    def _build_ui(self):
        # Root grid — header fixed, scrollable body expands, footer fixed
        self.rowconfigure(0, weight=0)
        self.rowconfigure(1, weight=0)  # divider
        self.rowconfigure(2, weight=1)  # scrollable form
        self.rowconfigure(3, weight=0)  # footer buttons
        self.columnconfigure(0, weight=1)

        # ── Header ────────────────────────────────────────────────
        hdr = tk.Frame(self, bg=C["surface"])
        hdr.grid(row=0, column=0, sticky="ew")
        title = "Modifier l'utilisateur" if self.user_id else "Nouvel utilisateur"
        tk.Label(hdr, text=title, bg=C["surface"],
                 fg=C["primary"], font=FNT["title"]).pack(
            side="left", padx=20, pady=14)

        tk.Frame(self, bg=C["border"], height=1).grid(row=1, column=0, sticky="ew")

        # ── Scrollable form area ──────────────────────────────────
        scroll_outer = tk.Frame(self, bg=C["bg"])
        scroll_outer.grid(row=2, column=0, sticky="nsew")
        scroll_outer.rowconfigure(0, weight=1)
        scroll_outer.columnconfigure(0, weight=1)

        canvas = tk.Canvas(scroll_outer, bg=C["bg"],
                           highlightthickness=0, borderwidth=0)
        vsb = ttk.Scrollbar(scroll_outer, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.grid(row=0, column=1, sticky="ns")
        canvas.grid(row=0, column=0, sticky="nsew")

        form = tk.Frame(canvas, bg=C["bg"])
        win_id = canvas.create_window((0, 0), window=form, anchor="nw")

        def _resize_canvas(event):
            canvas.itemconfig(win_id, width=event.width)
        canvas.bind("<Configure>", _resize_canvas)

        def _update_scrollregion(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))
        form.bind("<Configure>", _update_scrollregion)

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
        canvas.bind("<MouseWheel>", _on_mousewheel)
        form.bind("<MouseWheel>", _on_mousewheel)

        # ── Form fields (inside canvas) ───────────────────────────
        inner = tk.Frame(form, bg=C["bg"])
        inner.pack(fill="both", expand=True, padx=30, pady=20)
        inner.columnconfigure(0, weight=1)

        def _lbl(text):
            tk.Label(inner, text=text, bg=C["bg"],
                     fg=C["text_m"], font=FNT["sm"]).pack(anchor="w")

        def _entry(show=None):
            e = RoundedEntry(inner, show=show) if show else RoundedEntry(inner)
            e.pack(fill="x", pady=(4, 14), ipady=4)
            e.bind("<MouseWheel>", _on_mousewheel)
            return e

        # Username
        _lbl("Nom d'utilisateur (login)")
        self.username_entry = _entry()

        # Nom
        _lbl("Nom")
        self.nom_entry = _entry()

        # Prénom
        _lbl("Prénom")
        self.prenom_entry = _entry()

        # Rôle
        _lbl("Rôle")
        self.role_var = tk.StringVar(value="employe")
        role_frame = tk.Frame(inner, bg=C["bg"])
        role_frame.pack(anchor="w", pady=(4, 14))
        for role, label in [("admin", "Administrateur"),
                             ("manager", "Manager"),
                             ("employe", "Employé")]:
            tk.Radiobutton(
                role_frame, text=label, variable=self.role_var, value=role,
                bg=C["bg"], fg=C["text"], selectcolor=C["bg"],
                activebackground=C["bg"], font=FNT["default"]
            ).pack(side="left", padx=(0, 16))

        # Password
        pw_label = ("Mot de passe" if not self.user_id
                    else "Nouveau mot de passe (laisser vide = inchangé)")
        _lbl(pw_label)
        self.password_entry = _entry(show="●")

        # ── Password strength indicator ───────────────────────────
        pw_strength_frame = tk.Frame(inner, bg=C["bg"])
        pw_strength_frame.pack(fill="x", pady=(0, 10))
        pw_strength_frame.columnconfigure(0, weight=1)

        # Bar track
        bar_track = tk.Frame(pw_strength_frame, bg=C["border"], height=6)
        bar_track.pack(fill="x")
        bar_track.pack_propagate(False)
        self._pw_bar = tk.Frame(bar_track, bg=C["border"], height=6)
        self._pw_bar.place(relx=0, rely=0, relheight=1, relwidth=0)

        # Label
        self._pw_strength_label = tk.Label(
            pw_strength_frame, text="", bg=C["bg"],
            fg=C["text_m"], font=FNT["sm"], anchor="w"
        )
        self._pw_strength_label.pack(anchor="w", pady=(3, 0))

        def _update_strength(*_):
            pw = self.password_entry.get()
            if not pw:
                self._pw_bar.place(relwidth=0)
                self._pw_strength_label.configure(text="", fg=C["text_m"])
                return
            # Score : length + complexity
            score = 0
            if len(pw) >= 6:   score += 1
            if len(pw) >= 10:  score += 1
            if any(c.isdigit() for c in pw):     score += 1
            if any(c.isupper() for c in pw):     score += 1
            if any(c in "!@#$%^&*()-_=+[]{}|;:',.<>?/`~" for c in pw): score += 1

            levels = [
                (1, 0.20, "#DC2626", "Très faible"),   # red
                (2, 0.40, "#D97706", "Faible"),         # orange
                (3, 0.60, "#CA8A04", "Moyen"),          # yellow
                (4, 0.80, "#16A34A", "Fort"),           # green
                (5, 1.00, "#15803D", "Très fort"),      # dark green
            ]
            for lvl, width, color, label in levels:
                if score <= lvl:
                    self._pw_bar.place(relwidth=width)
                    self._pw_bar.configure(bg=color)
                    self._pw_strength_label.configure(
                        text=f"Niveau : {label}", fg=color)
                    break

        self.password_entry.bind("<KeyRelease>", _update_strength)
        # Also bind the inner Entry directly (RoundedEntry proxies bind to inner entry)

        # Confirm password
        _lbl("Confirmer le mot de passe")
        self.confirm_entry = _entry(show="●")

        # Status
        self.status_label = tk.Label(inner, text="", bg=C["bg"],
                                     fg=C["danger"], font=FNT["sm"],
                                     wraplength=380, justify="left")
        self.status_label.pack(anchor="w", pady=(2, 8))

        # ── Footer buttons (fixed, outside scroll) ────────────────
        footer = tk.Frame(self, bg=C["surface"])
        footer.grid(row=3, column=0, sticky="ew")
        tk.Frame(footer, bg=C["border"], height=1).pack(fill="x")
        btn_bar = tk.Frame(footer, bg=C["surface"])
        btn_bar.pack(fill="x", padx=20, pady=12)
        tk.Button(btn_bar, text="Enregistrer", command=self._save,
                  bg=C["primary"], fg="white", activebackground=C["primary_d"],
                  relief="flat", font=FNT["bold"], cursor="hand2",
                  padx=18, pady=9).pack(side="left", padx=(0, 10))
        tk.Button(btn_bar, text="Annuler", command=self.destroy,
                  bg=C["border"], fg=C["text"], activebackground=C["stripe"],
                  relief="flat", font=FNT["bold"], cursor="hand2",
                  padx=18, pady=9).pack(side="left")
    
    def _load_user_data(self):
        """Charge les données de l'utilisateur existant."""
        user = access_control.get_user(self.user_id)
        if not user:
            return
        
        self.username_entry.insert(0, user.get("username", ""))
        self.nom_entry.insert(0, user.get("nom", ""))
        self.prenom_entry.insert(0, user.get("prenom", ""))
        self.role_var.set(user.get("role", "employe"))
    
    def _save(self):
        """Enregistre l'utilisateur."""
        self.status_label.configure(text="", fg=C["danger"])
        
        username = self.username_entry.get().strip()
        nom = self.nom_entry.get().strip()
        password = self.password_entry.get().strip()
        confirm = self.confirm_entry.get().strip()
        role = self.role_var.get()
        
        # Validation
        if not username or not nom:
            self.status_label.configure(text="Nom d'utilisateur et nom sont requis")
            return
        
        if not self.user_id and not password:
            self.status_label.configure(text="Mot de passe requis pour un nouvel utilisateur")
            return
        
        if password and password != confirm:
            self.status_label.configure(text="Les mots de passe ne correspondent pas")
            return
        
        if password and len(password) < 6:
            self.status_label.configure(
                text="Le mot de passe doit contenir au moins 6 caractères")
            return
        
        try:
            if self.user_id:
                # Mise à jour
                data = {
                    "username": username,
                    "nom": nom,
                    "prenom": self.prenom_entry.get().strip(),
                    "role": role,
                }
                if password:
                    data["password"] = password
                access_control.update_user(self.user_id, data)
            else:
                # Création
                data = {
                    "username": username,
                    "password": password,
                    "nom": nom,
                    "prenom": self.prenom_entry.get().strip(),
                    "role": role,
                }
                new_id = access_control.create_user(data)
                # Appliquer le template de rôle
                access_control.apply_role_template(new_id, role)
            
            if self.on_save:
                self.on_save()
            
            toast("Utilisateur enregistré.", "success", "Succès")
            self.destroy()
        
        except sqlite3.IntegrityError:
            self.status_label.configure(text="Ce nom d'utilisateur existe déjà")
        except Exception as e:
            self.status_label.configure(text=f"Erreur : {e}")


# ═══════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ═══════════════════════════════════════════════════════════════════
def run(user):
    app = MarassimApp(user)
    app.mainloop()
