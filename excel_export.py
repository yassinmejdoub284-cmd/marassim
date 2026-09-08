"""
excel_export.py — Génère le calendrier Excel (1ère page du système) avec
une colonne par salle (Almes, Chichkhane, Gazon/Rayhane), colorée avec la
couleur signature de la salle, et une note indiquant l'heure de début /
fin de chaque événement.
"""
import calendar
from datetime import date
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.comments import Comment

from rules import SALLE_COLORS
from database import get_reservations_in_range

HEADER_FILL = PatternFill("solid", fgColor="404040")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WEEKEND_FILL = PatternFill("solid", fgColor="F2F2F2")


def export_month_calendar(year: int, month: int, output_path: str):
    start = date(year, month, 1)
    last_day = calendar.monthrange(year, month)[1]
    end = date(year, month, last_day)

    reservations = get_reservations_in_range(start.isoformat(), end.isoformat())
    by_date = {}
    for r in reservations:
        by_date.setdefault(r["date_evenement"], []).append(r)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Calendrier {calendar.month_name[month]} {year}"

    headers = ["Date", "Jour", "Almes", "Chichkhane", "Rayhane"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN_BORDER

    jours_fr = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]

    row_idx = 2
    for day_num in range(1, last_day + 1):
        d = date(year, month, day_num)
        date_str = d.isoformat()
        weekday = d.weekday()

        ws.cell(row=row_idx, column=1, value=d.strftime("%d/%m/%Y")).border = THIN_BORDER
        ws.cell(row=row_idx, column=2, value=jours_fr[weekday]).border = THIN_BORDER

        is_weekend = weekday >= 5
        if is_weekend:
            ws.cell(row=row_idx, column=1).fill = WEEKEND_FILL
            ws.cell(row=row_idx, column=2).fill = WEEKEND_FILL

        col_map = {"Almes": 3, "Chichkhane": 4, "Rayhane": 5}
        day_reservations = by_date.get(date_str, [])
        for salle, col in col_map.items():
            cell = ws.cell(row=row_idx, column=col)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            matches = [r for r in day_reservations if r["salle"] == salle]
            if matches:
                texts = []
                for r in matches:
                    label = f"{r['heure_debut']}-{r['heure_fin']}"
                    if r.get("nom_client"):
                        label += f"\n{r['nom_client']}"
                    texts.append(label)
                cell.value = "\n---\n".join(texts)
                cell.fill = PatternFill("solid", fgColor=SALLE_COLORS[salle])
                note_lines = [
                    f"{r['type_evenement'] or ''} | {r['heure_debut']}-{r['heure_fin']} | "
                    f"{r['nom_client'] or ''} | {r.get('nb_places') or ''} places"
                    for r in matches
                ]
                cell.comment = Comment("\n".join(note_lines), "Marassim")
            # option violoniste dans le Gazon, affichée en plus si présente ce jour
            if salle == "Rayhane":
                violonistes = [
                    r for r in day_reservations
                    if r["salle"] in ("Almes", "Chichkhane") and r.get("with_violoniste")
                ]
                if violonistes:
                    extra = "\n".join(
                        f"Violoniste {v['violoniste_heure_debut']}-{v['violoniste_heure_fin']}"
                        for v in violonistes
                    )
                    cell.value = (cell.value + "\n" + extra) if cell.value else extra
                    if not matches:
                        cell.fill = PatternFill("solid", fgColor=SALLE_COLORS["Rayhane"])

        row_idx += 1

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 10
    for col_letter in ("C", "D", "E"):
        ws.column_dimensions[col_letter].width = 26
    ws.freeze_panes = "A2"

    # Légende
    legend_row = row_idx + 2
    ws.cell(row=legend_row, column=1, value="Légende :").font = Font(bold=True)
    for i, (salle, color) in enumerate(SALLE_COLORS.items()):
        r = legend_row + 1 + i
        ws.cell(row=r, column=1, value=salle).fill = PatternFill("solid", fgColor=color)
        ws.cell(row=r, column=1).border = THIN_BORDER

    wb.save(output_path)
    return output_path
