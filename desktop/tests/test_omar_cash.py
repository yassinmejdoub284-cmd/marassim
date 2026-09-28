import base64
import io
import json
import unittest
from decimal import Decimal
from pathlib import Path
import sys

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'server'))
import test_server
from app import APIError


class OmarCashTests(unittest.TestCase):
    setUp=test_server.ServerTests.setUp
    tearDown=test_server.ServerTests.tearDown
    call=test_server.ServerTests.call
    create=test_server.ServerTests.create
    reservation=test_server.ServerTests.reservation

    def journal(self,start,end):
        return self.call('GET',f'/journal?caisse=omar&start={start}&end={end}')

    def test_history_carry_december_january_empty_month_and_backdated_change(self):
        self.call('POST','/recettes-omar',{'date_recette':'2025-12-20','designation':'Ancienne recette','montant':7508.078})
        charge=self.call('POST','/charges-omar',{'date_encaissement':'2025-12-31','designation':'Ancienne charge','montant':100.001})
        self.call('POST','/recettes-omar',{'date_recette':'2026-01-01','designation':'Recette janvier','montant':500.001})
        self.call('POST','/charges-omar',{'date_encaissement':'2026-01-31','designation':'Charge janvier','montant':385.002})
        december=self.journal('2025-12-01','2025-12-31'); january=self.journal('2026-01-01','2026-01-31')
        self.assertEqual(december['closing'],january['opening'])
        self.assertEqual(Decimal(str(january['opening'])),Decimal('7408.077'))
        self.assertEqual(january['closing'],7523.076)
        february=self.journal('2026-02-01','2026-02-28')
        self.assertEqual(february['opening'],january['closing']);self.assertEqual(february['opening'],february['closing'])
        self.call('PUT',f"/charges-omar/{charge['id']}",{'revision':charge['revision'],'montant':101.001})
        self.assertEqual(self.journal('2026-02-01','2026-02-28')['opening'],7522.076)
        # Viewing/regenerating reports never creates carry receipts.
        self.assertEqual(len(self.call('GET','/recettes-omar')),2)

    def test_accepted_omar_amount_once_and_no_tawfik_or_day_workers(self):
        booking=self.create(dict(self.reservation(acompte1=2000,date_acompte1='2026-02-01',num_caisse_omar1='500'),forfait=3000))
        self.assertEqual(self.journal('2026-02-01','2026-02-28')['income'],0)
        self.call('POST','/reception',{'res_id':booking['id'],'slot':1,'statut':'accepte','revision':booking['revision']})
        self.call('POST','/charges',{'date_encaissement':'2026-02-02','designation':'Autre caisse','montant':150})
        for kind in ('jour','nuit'):
            emp=self.call('POST','/employees',{'nom':kind,'prenom':'Test','role':'Employé','type_ouvrier':kind,'actif':1})
            self.call('POST','/employee-payments',{'employee_id':emp['id'],'employeeRevision':emp['revision'],'date_paiement':'2026-02-03','periode_debut':'2026-02-01','periode_fin':'2026-02-28','montant_total':100})
        report=self.journal('2026-02-01','2026-02-28')
        self.assertEqual(report['income'],500);self.assertEqual(report['expense'],100);self.assertEqual(report['closing'],400)

    def test_pointage_details_are_snapshots_not_extra_cash_movements(self):
        emp=self.call('POST','/employees',{'nom':'Employé test','prenom':'Nuit','role':'Accueil','type_ouvrier':'nuit','salaire_soiree':80,'salaire_apres_midi':40,'actif':1})
        self.call('POST','/pointage/batch',{'entries':[{'employee_id':emp['id'],'date_pointage':'2026-06-05','periode':period} for period in ('apres_midi','soiree')],'notes':'Dépense soirée CHICHKHANE 15H ET ALMES 21H'})
        report=self.journal('2026-06-01','2026-06-30')
        self.assertEqual(report['expense'],120);self.assertEqual(len(report['sorties']),1)
        self.assertEqual([d['amount'] for d in report['sorties'][0]['details']],[40,80])
        refreshed=self.call('GET',f"/employees/{emp['id']}")
        self.call('PUT',f"/employees/{emp['id']}",{'revision':refreshed['revision'],'salaire_soiree':500})
        self.assertEqual(self.journal('2026-06-01','2026-06-30')['sorties'][0]['details'][1]['amount'],80)

    def test_soiree_group_contains_only_night_employees_and_its_chosen_events(self):
        evening=self.create(dict(self.reservation(),date_evenement='2026-06-06'))
        afternoon=self.create(dict(self.reservation(),date_evenement='2026-06-06',salle='Chichkhane',heure_debut='15:00',heure_fin='18:00'))
        night=self.call('POST','/employees',{'nom':'Service nuit','prenom':'Test','role':'Accueil','type_ouvrier':'nuit','salaire_soiree':80,'actif':1})
        day=self.call('POST','/employees',{'nom':'Service jour','prenom':'Test','role':'Entretien','type_ouvrier':'jour','salaire_journalier':500,'actif':1})
        self.call('POST','/pointage/batch',{'entries':[{'employee_id':night['id'],'date_pointage':'2026-06-06','periode':'soiree'},{'employee_id':day['id'],'date_pointage':'2026-06-06','periode':'journee_complete'}],'eventIds':[evening['id']]})
        data=self.journal('2026-06-01','2026-06-30'); row=data['sorties'][0]
        self.assertEqual(data['expense'],80);self.assertEqual(len(row['details']),1)
        self.assertIn('ALMES 21H',row['designation']);self.assertNotIn('CHICHKHANE',row['designation'])
        self.assertIn('Service nuit',row['details'][0]['name']);self.assertNotIn('Service jour',json.dumps(row))

    def test_excel_pdf_match_statement_and_user_text_is_not_a_formula(self):
        self.call('POST','/recettes-omar',{'date_recette':'2026-05-30','designation':'Recette antérieure','montant':7508.078})
        self.call('POST','/charges-omar',{'date_encaissement':'2026-06-05','designation':'=DANGER()','montant':555})
        query='caisse=omar&start=2026-06-01&end=2026-06-30'
        from openpyxl import load_workbook
        file=self.call('GET','/exports/journal?'+query)
        workbook=load_workbook(io.BytesIO(base64.b64decode(file['base64'])))
        ws=workbook.active
        self.assertEqual([c.value for c in ws[4]],['Date','Désignation','Débit','Crédit','Solde'])
        self.assertEqual(ws['E5'].value,7508.078);self.assertEqual(ws['E8'].value,6953.078)
        self.assertEqual(ws['B6'].data_type,'s');self.assertEqual(ws['B6'].value,'=DANGER()')
        pdf=self.call('GET','/exports/journal?'+query+'&format=pdf')
        from pypdf import PdfReader
        text=' '.join(p.extract_text() for p in PdfReader(io.BytesIO(base64.b64decode(pdf['base64']))).pages)
        self.assertIn('SOLDE DE DÉPART',text);self.assertIn('7 508,078',text);self.assertIn('6 953,078',text)
        self.assertIn('REPORT AU MOIS SUIVANT',text)

    def test_module_exports_use_permissions_and_omit_passwords(self):
        file=self.call('GET','/exports/table?resource=users')
        from openpyxl import load_workbook
        ws=load_workbook(io.BytesIO(base64.b64decode(file['base64']))).active
        text=' '.join(str(c.value) for row in ws for c in row)
        self.assertNotIn('password',text);self.assertNotIn('scrypt',text)
        user=self.call('POST','/users',{'username':'limited','password':'test-password-123','nom':'Lecture','role':'employe','modules':['Charges']})
        token=self.app.dispatch('POST','/api/login',{}, {'username':'limited','password':'test-password-123'})['token']
        self.assertTrue(self.call('GET','/exports/table?resource=charges',token=token)['base64'])
        for resource in ('users','charges-omar','reservations','../../audit_log'):
            with self.assertRaises(APIError):self.call('GET','/exports/table?resource='+resource,token=token)


if __name__=='__main__':unittest.main()
