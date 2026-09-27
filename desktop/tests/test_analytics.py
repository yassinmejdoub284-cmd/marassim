import base64
import io
import sqlite3
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'server'))
import test_server
from app import APIError
from backups import restore_backup
from openpyxl import load_workbook


class AnalyticsTests(unittest.TestCase):
    setUp = test_server.ServerTests.setUp
    tearDown = test_server.ServerTests.tearDown
    call = test_server.ServerTests.call
    create = test_server.ServerTests.create
    reservation = test_server.ServerTests.reservation

    def plan(self, **values):
        data = dict(due_date='2027-01-31',designation='Loyer',direction='out',amount=100,category='Loyer',frequency='once',state='planned')
        data.update(values)
        return self.call('POST','/forecast-items',data)

    def test_contract_j15_partial_payments_and_exclusions(self):
        data = self.reservation(acompte1=100.111,acompte2=200.222,acompte3=300.333,date_reste='2027-01-14')
        data['forfait'] = 1000.777
        row = self.create(data)
        temp = self.reservation(is_temporaire=1); temp['date_evenement'] = '2027-01-20'
        self.create(temp)
        full = self.reservation(acompte1=1000); full['date_evenement'] = '2027-01-21'; self.create(full)
        result = self.call('GET','/forecast?start=2026-12-31&end=2027-01-31&opening=0&overdue=0')
        self.assertEqual(len(result['events']),1)
        self.assertEqual(result['events'][0]['due'],'2026-12-31')
        self.assertEqual(result['summary']['income'],400.111)
        self.assertEqual(result['summary']['closing'],400.111)
        scaled = self.call('GET','/forecast?start=2026-12-31&end=2027-01-31&rate=80&overdue=0')
        self.assertEqual(scaled['summary']['income'],320.089)
        self.assertEqual(scaled['events'][0]['nominal'],400.111)
        self.assertEqual(self.call('GET',f"/reservations/{row['id']}")['acompte1'],100.111)

    def test_overdue_option_carries_balances_once(self):
        self.create()
        result = self.call('GET','/forecast?start=2027-01-10&end=2027-03-01')
        self.assertEqual(result['events'][0]['date'],'2027-01-10')
        self.assertEqual(result['events'][0]['due'],'2026-12-31')
        self.assertTrue(result['events'][0]['late'])
        self.assertEqual(result['summary']['income'],1000)
        self.assertEqual(self.call('GET','/forecast?start=2027-01-10&end=2027-03-01&overdue=0')['events'],[])

    def test_month_end_recurrence_and_settled_do_not_post_cash(self):
        row = self.plan(frequency='monthly',repeat_until='2027-03-31')
        self.plan(designation='Réalisé',state='settled',amount=200)
        cancel = self.plan(designation='Annulé',amount=300)
        self.call('DELETE',f"/forecast-items/{cancel['id']}",{'revision':cancel['revision']})
        result = self.call('GET','/forecast?start=2027-01-01&end=2027-03-31')
        self.assertEqual([e['date'] for e in result['events']],['2027-01-31','2027-02-28','2027-03-31'])
        self.assertEqual(result['summary']['expense'],300)
        self.call('PUT',f"/forecast-items/{row['id']}",{'revision':row['revision'],'state':'settled'})
        self.assertEqual(self.call('GET','/forecast?start=2027-01-01&end=2027-03-31')['summary']['expense'],0)
        self.assertEqual(self.call('GET','/charges'),[])
        self.assertEqual(self.call('GET','/employee-payments'),[])
        self.assertEqual(self.call('GET','/charges-omar'),[])

    def test_daily_dip_alert_even_when_month_closes_positive(self):
        self.plan(due_date='2027-01-02',amount=200)
        self.plan(due_date='2027-01-20',amount=300,direction='in')
        result = self.call('GET','/forecast?start=2027-01-01&end=2027-01-31&opening=100&floor=50')
        self.assertEqual(result['summary']['closing'],200)
        self.assertEqual(result['summary']['lowest'],-100)
        self.assertEqual(result['summary']['negativeDate'],'2027-01-02')
        self.assertEqual(result['summary']['floorDate'],'2027-01-02')

    def test_idempotent_plan_and_concurrent_revisions(self):
        body = dict(due_date='2027-01-31',designation='Unique',direction='out',amount=100)
        key = str(uuid.uuid4())
        row = self.call('POST','/forecast-items',body,key)
        self.assertEqual(row,self.call('POST','/forecast-items',body,key))
        def change(n):
            try: return self.call('PUT',f"/forecast-items/{row['id']}",{'amount':n+200,'revision':row['revision']})
            except APIError as e: return e.status
        with ThreadPoolExecutor(4) as pool: result = list(pool.map(change,range(4)))
        self.assertEqual(sum(isinstance(r,dict) for r in result),1)
        self.assertEqual(result.count(409),3)
        self.assertEqual(len(self.call('GET','/forecast-items')),1)
        settings = self.call('GET','/forecast')['settings']
        update = {k:settings[k] for k in ('opening_date','opening_balance','safety_floor','collection_rate','revision')}
        update['opening_balance'] = -10.123
        self.call('PUT','/forecast/settings',update)
        with self.assertRaises(APIError) as error: self.call('PUT','/forecast/settings',update)
        self.assertEqual(error.exception.status,409)
        self.assertEqual(self.call('GET','/forecast')['parameters']['opening'],-10.123)

    def test_reports_event_vs_payment_dates_and_hall_cash(self):
        self.create(self.reservation(acompte1=100,date_acompte1='2026-12-15',methode_acompte1='Espèce',num_caisse1='1'))
        data = self.reservation(acompte1=200,date_acompte1='2027-01-01'); data.update(salle='Rayhane',date_evenement='2027-02-15',heure_debut='16:00',heure_fin='20:00',with_sono=1); self.create(data)
        self.call('POST','/charges',{'date_encaissement':'2027-01-15','designation':'Toutes salles','montant':50})
        result = self.call('GET','/reports?start=2027-01-01&end=2027-01-31&hall=Almes')
        self.assertEqual(result['summary']['booked'],1000)
        self.assertEqual(result['summary']['paid'],100)
        self.assertEqual(result['summary']['deposits'],0)
        self.assertEqual(result['summary']['outstanding'],900)
        self.assertEqual(result['cash'][0]['expense'],50)
        result = self.call('GET','/reports?start=2027-01-01&end=2027-01-31')
        self.assertEqual(result['summary']['deposits'],200)
        self.assertEqual(result['previousPeriod'],{'start':'2026-12-01','end':'2026-12-31'})

    def test_report_employees_keeps_existing_tariff_and_global_balance(self):
        emp = self.call('POST','/employees',{'nom':'Employé','prenom':'Test','role':'Service','type_ouvrier':'jour','salaire_journalier':50,'prix_heure_supp':5})
        for day in ('2026-12-31','2027-01-15'):
            self.call('POST','/pointage/batch',{'entries':[{'employee_id':emp['id'],'date_pointage':day,'periode':'journee_complete','heures_supplementaires':2}]})
        result = self.call('GET','/reports?start=2027-01-01&end=2027-01-31')['employees'][0]
        self.assertEqual(result['days'],1)
        self.assertEqual(result['earned'],60)
        self.assertEqual(result['globalRemaining'],120)

    def test_permissions_are_explicit_and_independent(self):
        account = self.call('POST','/users',{'username':'limited','password':'test-password-123','nom':'Limité','role':'employe','modules':['Calendrier']})
        login = self.call('POST','/login',{'username':'limited','password':'test-password-123'})
        for path in ['/reports','/forecast','/forecast-items','/exports/reports','/exports/forecast']:
            with self.assertRaises(APIError) as error: self.call('GET',path,token=login['token'])
            self.assertEqual(error.exception.status,403)
        row = next(r for r in self.call('GET','/users') if r['id'] == account['id'])
        self.call('PUT',f"/users/{row['id']}",{**row,'modules':['Rapports avancés']})
        login = self.call('POST','/login',{'username':'limited','password':'test-password-123'})
        self.assertEqual(login['user']['modules'],['Rapports avancés'])
        self.assertIn('summary',self.call('GET','/reports',token=login['token']))
        with self.assertRaises(APIError): self.call('GET','/forecast',token=login['token'])

    def test_exports_and_backup_include_new_shared_data(self):
        self.plan(designation='=HYPERLINK("bad")')
        result = self.call('GET','/exports/forecast?start=2027-01-01&end=2027-03-31')
        workbook = load_workbook(io.BytesIO(base64.b64decode(result['base64'])))
        self.assertEqual(workbook['Mouvements prévus']['C2'].data_type,'s')
        self.assertIn('Trésorerie prévue',workbook.sheetnames)
        report = self.call('GET','/exports/reports?start=2027-01-01&end=2027-03-31')
        self.assertEqual(len(load_workbook(io.BytesIO(base64.b64decode(report['base64']))).sheetnames),9)
        snapshot = self.app.backups.create()
        restore_backup(self.app.backups.root/snapshot['name'],self.app.backups.key_path,self.root/'restored')
        conn = sqlite3.connect(self.root/'restored/marassim.db')
        try:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM forecast_items').fetchone()[0],1)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM forecast_settings').fetchone()[0],1)
        finally:
            conn.close()

    def test_invalid_horizons_and_plans_rejected(self):
        for change in [{'amount':0},{'direction':'bad'},{'frequency':'monthly','repeat_until':'2026-12-31'},{'state':'bad'}]:
            with self.assertRaises((ValueError,APIError)): self.plan(**change)
        for path in ['/forecast?start=2027-02-01&end=2027-01-01','/forecast?rate=101','/reports?start=2027-01-01&end=2030-01-01']:
            with self.assertRaises(ValueError): self.call('GET',path)

if __name__ == '__main__': unittest.main()
