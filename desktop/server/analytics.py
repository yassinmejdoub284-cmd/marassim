"""Read-only operational reports and explicitly planned cash movements, in DT."""
import calendar
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

REPORT_MODULES = ['Rapports avancés', 'Cash-flow prévu']
HALLS = ['Almes', 'Chichkhane', 'Rayhane']
SCHEMA = '''
CREATE TABLE IF NOT EXISTS forecast_items (
 id INTEGER PRIMARY KEY AUTOINCREMENT, due_date TEXT NOT NULL,
 designation TEXT NOT NULL, direction TEXT NOT NULL, amount REAL NOT NULL,
 category TEXT NOT NULL DEFAULT 'Autre', frequency TEXT NOT NULL DEFAULT 'once',
 repeat_until TEXT, state TEXT NOT NULL DEFAULT 'planned',
 revision INTEGER NOT NULL DEFAULT 1, created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS forecast_settings (
 id INTEGER PRIMARY KEY CHECK(id=1), opening_date TEXT NOT NULL,
 opening_balance REAL NOT NULL DEFAULT 0, safety_floor REAL NOT NULL DEFAULT 0,
 collection_rate INTEGER NOT NULL DEFAULT 100, revision INTEGER NOT NULL DEFAULT 1
);
'''

def amount(value):
    try:
        result = Decimal(str(value or 0))
        if not result.is_finite() or abs(result) > Decimal('1000000000000'):
            raise ValueError('Montant hors limites.')
        return result.quantize(Decimal('.001'), rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError):
        raise ValueError('Montant invalide.')

def number(value):
    return float(amount(value))

def days(start, end, limit=731):
    a, b = date.fromisoformat(start), date.fromisoformat(end)
    if start != a.isoformat() or end != b.isoformat():
        raise ValueError('Utilisez une date au format AAAA-MM-JJ.')
    if a > b or (b-a).days >= limit:
        raise ValueError('Choisissez une période de deux ans maximum, dans le bon ordre.')
    return a, b

def remaining(row):
    return max(Decimal(0), amount(row.get('forfait')) - sum((amount(row.get(f'acompte{i}')) for i in range(1, 4)), Decimal(0)))

def due(row):
    # The selected contractual rule remains authoritative, regardless of date_reste.
    return (date.fromisoformat(row['date_evenement']) - timedelta(days=15)).isoformat()

def month_after(day, offset):
    index = day.year * 12 + day.month - 1 + offset
    year, month = divmod(index, 12)
    return date(year, month+1, min(day.day, calendar.monthrange(year, month+1)[1]))

def validate_item(row):
    if date.fromisoformat(row['due_date']).isoformat() != row['due_date']:
        raise ValueError('Date invalide : utilisez AAAA-MM-JJ.')
    if not isinstance(row.get('designation'), str) or not row['designation'].strip():
        raise ValueError('Désignation obligatoire.')
    if amount(row.get('amount')) <= 0:
        raise ValueError('Le montant doit être positif.')
    if row.get('direction') not in ('in', 'out') or row.get('frequency') not in ('once', 'monthly') or row.get('state') not in ('planned', 'settled', 'cancelled'):
        raise ValueError('Sens, répétition ou état invalide.')
    if row['frequency'] == 'monthly':
        if not row.get('repeat_until'):
            raise ValueError('Indiquez la fin de la répétition mensuelle.')
        start, end = date.fromisoformat(row['due_date']), date.fromisoformat(row['repeat_until'])
        if end < start or (end-start).days > 3653:
            raise ValueError('La répétition doit durer dix ans maximum, dans le bon ordre.')

def validate_settings(data):
    expected = {'opening_date', 'opening_balance', 'safety_floor', 'collection_rate', 'revision'}
    if set(data) != expected:
        raise ValueError('Hypothèses incomplètes ou champs inconnus.')
    date.fromisoformat(data['opening_date'])
    opening, floor = number(data['opening_balance']), number(data['safety_floor'])
    rate = amount(data['collection_rate'])
    if floor < 0 or not rate.is_finite() or rate != rate.to_integral_value() or not 0 <= rate <= 100:
        raise ValueError('Seuil positif et taux de recouvrement entier entre 0 et 100 requis.')
    return {'opening_date': data['opening_date'], 'opening_balance': opening, 'safety_floor': floor, 'collection_rate': int(rate)}

def report(conn, database, query):
    start = query.get('start', date.today().replace(day=1).isoformat())
    end = query.get('end', date.today().isoformat())
    a, b = days(start, end)
    hall = query.get('hall', 'all')
    if hall not in HALLS + ['all']:
        raise ValueError('Salle inconnue.')
    all_rows = [dict(r) for r in conn.execute('SELECT * FROM reservations')]
    scoped = [r for r in all_rows if hall == 'all' or r['salle'] == hall]
    confirmed = [r for r in scoped if not r.get('is_temporaire')]
    def valid_event(r):
        try:
            return date.fromisoformat(r['date_evenement']).isoformat() == r['date_evenement']
        except (ValueError, TypeError):
            return False
    invalid_events = sum(not valid_event(r) for r in scoped)
    selected = [r for r in confirmed if valid_event(r) and start <= r['date_evenement'] <= end]
    temporary = [r for r in scoped if r.get('is_temporaire') and valid_event(r) and start <= r['date_evenement'] <= end]
    previous_end = a - timedelta(days=1)
    previous_start = previous_end - (b-a)
    previous = [r for r in confirmed if valid_event(r) and previous_start.isoformat() <= r['date_evenement'] <= previous_end.isoformat()]
    paid = lambda r: sum((amount(r.get(f'acompte{i}')) for i in range(1,4)), Decimal(0))
    booked = sum((amount(r.get('forfait')) for r in selected), Decimal(0))
    deposits, methods = [], defaultdict(lambda: {'count': 0, 'amount': Decimal(0)})
    unknown_dates = 0
    for r in confirmed:
        for slot in range(1, 4):
            value, day = amount(r.get(f'acompte{slot}')), r.get(f'date_acompte{slot}')
            if value and not day:
                unknown_dates += 1
            if value and day and start <= day <= end:
                method = r.get(f'methode_acompte{slot}') or 'Non renseigné'
                methods[method]['count'] += 1
                methods[method]['amount'] += value
                deposits.append({'date': day, 'amount': value})
    by_hall, by_type = [], []
    for label, target in [(h, by_hall) for h in HALLS] + [(t, by_type) for t in sorted({r.get('type_evenement') or 'Non renseigné' for r in selected})]:
        rows = [r for r in selected if (r['salle'] if target is by_hall else r.get('type_evenement') or 'Non renseigné') == label]
        total = sum((amount(r.get('forfait')) for r in rows), Decimal(0))
        item = {'name': label, 'count': len(rows), 'booked': number(total)}
        if target is by_hall:
            item.update(paid=number(sum((paid(r) for r in rows), Decimal(0))), outstanding=number(sum((remaining(r) for r in rows), Decimal(0))), average=number(total/len(rows)) if rows else 0)
        target.append(item)
    monthly = []
    cursor = a.replace(day=1)
    while cursor <= b:
        key = cursor.isoformat()[:7]
        rows = [r for r in selected if r['date_evenement'].startswith(key)]
        monthly.append({'month': key, 'count': len(rows), 'booked': number(sum((amount(r.get('forfait')) for r in rows), Decimal(0))), 'deposits': number(sum((p['amount'] for p in deposits if p['date'].startswith(key)), Decimal(0)))})
        cursor = month_after(cursor, 1)
    outstanding = []
    for r in selected:
        balance = remaining(r)
        if balance > 0:
            deadline = due(r)
            outstanding.append({'id': r['id'], 'bon': r['num_bon'], 'client': r['nom_client'], 'hall': r['salle'], 'event': r['date_evenement'], 'due': deadline, 'remaining': number(balance), 'overdueDays': max(0, (date.today()-date.fromisoformat(deadline)).days)})
    outstanding.sort(key=lambda r: (r['due'], -r['remaining']))
    cash = []
    for label, function, expense_key in [('Tawfik', database.get_journal_caisse, 'charges'), ('Omar', database.get_journal_caisse_omar, 'sorties')]:
        journal = function(start, end)
        income = sum((amount(r['montant']) for r in journal['recettes']), Decimal(0))
        expense = sum((amount(r['montant']) for r in journal[expense_key]), Decimal(0))
        cash.append({'name': label, 'income': number(income), 'expense': number(expense), 'net': number(income-expense)})
    employees = []
    for employee in conn.execute('SELECT * FROM employees ORDER BY nom,prenom'):
        fiche = database.get_employee_fiche_data(employee['id'])
        pointages = [p for p in fiche['pointages'] if start <= p['date_pointage'] <= end]
        payments = [p for p in fiche['paiements'] if start <= p['date_paiement'] <= end]
        if pointages or payments:
            employees.append({'name': f"{employee['nom']} {employee['prenom'] or ''}".strip(), 'type': employee['type_ouvrier'], 'days': len({p['date_pointage'] for p in pointages}), 'periods': len(pointages), 'hours': number(sum((amount(p.get('heures_supplementaires')) for p in pointages), Decimal(0))), 'earned': number(sum((amount(p['tarif']) for p in pointages), Decimal(0))), 'paid': number(sum((amount(p['montant_total']) for p in payments), Decimal(0))), 'globalRemaining': number(fiche['reste'])})
    prev_booked = sum((amount(r.get('forfait')) for r in previous), Decimal(0))
    warnings = ([f'{unknown_dates} acompte(s) sans date : exclus des encaissements de la période.'] if unknown_dates else []) + ([f'{invalid_events} réservation(s) avec date invalide : exclues des totaux par événement.'] if invalid_events else [])
    return {'period': {'start': start, 'end': end, 'hall': hall}, 'asOf': date.today().isoformat(), 'summary': {'count': len(selected), 'temporaryCount': len(temporary), 'booked': number(booked), 'paid': number(sum((paid(r) for r in selected), Decimal(0))), 'outstanding': number(sum((remaining(r) for r in selected), Decimal(0))), 'deposits': number(sum((p['amount'] for p in deposits), Decimal(0))), 'overdue': number(sum((amount(r['remaining']) for r in outstanding if r['due'] < date.today().isoformat()), Decimal(0))), 'previousBooked': number(prev_booked), 'previousCount': len(previous), 'changePercent': round(float((booked-prev_booked)/prev_booked*100), 1) if prev_booked else None}, 'previousPeriod': {'start': previous_start.isoformat(), 'end': previous_end.isoformat()}, 'halls': by_hall, 'types': by_type, 'monthly': monthly, 'outstanding': outstanding, 'methods': [{'name': k, 'count': v['count'], 'amount': number(v['amount'])} for k,v in methods.items()], 'cash': cash, 'employees': employees, 'warnings': warnings}

def forecast(conn, query):
    settings = dict(conn.execute('SELECT * FROM forecast_settings WHERE id=1').fetchone())
    start = query.get('start', settings['opening_date'])
    end = query.get('end', (date.fromisoformat(start)+timedelta(days=179)).isoformat())
    a, b = days(start, end)
    opening = amount(query.get('opening', settings['opening_balance']))
    floor = amount(query.get('floor', settings['safety_floor']))
    rate = amount(query.get('rate', settings['collection_rate']))
    if floor < 0 or not 0 <= rate <= 100:
        raise ValueError('Seuil positif et taux entre 0 et 100 requis.')
    group = query.get('group', 'month')
    if group not in ('month', 'week'):
        raise ValueError('Regroupement inconnu.')
    include_overdue = query.get('overdue', '1') == '1'
    events, overdue, unscheduled = [], Decimal(0), 0
    for row in conn.execute('SELECT * FROM reservations WHERE COALESCE(is_temporaire,0)=0'):
        r = dict(row)
        balance = remaining(r)
        if not balance:
            continue
        try:
            deadline = due(r)
        except (ValueError, TypeError):
            unscheduled += 1
            continue
        if deadline < start:
            overdue += balance
            if not include_overdue:
                continue
        day = max(start, deadline)
        if day <= end:
            events.append({'id': f"reservation:{r['id']}", 'date': day, 'due': deadline, 'designation': f"Solde {r['num_bon'] or r['id']} · {r['nom_client']}", 'direction': 'in', 'amount': number(balance*rate/100), 'nominal': number(balance), 'source': 'reservation', 'hall': r['salle'], 'late': deadline < start})
    items = [dict(r) for r in conn.execute('SELECT * FROM forecast_items ORDER BY due_date,id')]
    for item in items:
        if item['state'] != 'planned':
            continue
        first = date.fromisoformat(item['due_date'])
        last = date.fromisoformat(item['repeat_until']) if item['frequency'] == 'monthly' else first
        offset = max(0, (a.year-first.year)*12+a.month-first.month-1) if item['frequency'] == 'monthly' else 0
        while True:
            occurrence = month_after(first, offset) if item['frequency'] == 'monthly' else first
            if occurrence > min(b, last):
                break
            if occurrence >= a:
                events.append({'id': f"plan:{item['id']}:{occurrence}", 'date': occurrence.isoformat(), 'due': occurrence.isoformat(), 'designation': item['designation'], 'direction': item['direction'], 'amount': number(item['amount']), 'nominal': number(item['amount']), 'source': 'planned', 'category': item['category'], 'late': False})
            if item['frequency'] != 'monthly':
                break
            offset += 1
    events.sort(key=lambda e: (e['date'], e['direction'], e['id']))
    daily = defaultdict(lambda: {'income': Decimal(0), 'expense': Decimal(0)})
    for event in events:
        daily[event['date']]['income' if event['direction'] == 'in' else 'expense'] += amount(event['amount'])
    balance, low, low_date, negative_date, floor_date = opening, opening, start, start if opening < 0 else None, start if opening < floor else None
    buckets = {}
    cursor = a
    while cursor <= b:
        key = cursor.isoformat()[:7] if group == 'month' else (cursor-timedelta(days=cursor.weekday())).isoformat()
        bucket = buckets.setdefault(key, {'label': key, 'start': cursor.isoformat(), 'end': cursor.isoformat(), 'opening': number(balance), 'income': Decimal(0), 'expense': Decimal(0)})
        movement = daily[cursor.isoformat()]
        balance += movement['income']-movement['expense']
        bucket['income'] += movement['income']; bucket['expense'] += movement['expense']
        bucket['closing'] = number(balance); bucket['end'] = cursor.isoformat()
        if balance < low:
            low, low_date = balance, cursor.isoformat()
        if balance < 0 and negative_date is None:
            negative_date = cursor.isoformat()
        if balance < floor and floor_date is None:
            floor_date = cursor.isoformat()
        cursor += timedelta(days=1)
    periods = [{**row, 'income': number(row['income']), 'expense': number(row['expense']), 'net': number(row['income']-row['expense'])} for row in buckets.values()]
    return {'settings': settings, 'parameters': {'start': start, 'end': end, 'opening': number(opening), 'floor': number(floor), 'rate': number(rate), 'group': group, 'overdue': include_overdue}, 'summary': {'income': number(sum((amount(p['income']) for p in periods), Decimal(0))), 'expense': number(sum((amount(p['expense']) for p in periods), Decimal(0))), 'closing': number(balance), 'lowest': number(low), 'lowestDate': low_date, 'negativeDate': negative_date, 'floorDate': floor_date, 'overdue': number(overdue)}, 'periods': periods, 'events': events, 'items': items, 'warnings': [f'{unscheduled} réservation(s) avec date invalide : non projetées.'] if unscheduled else []}

def export_workbook(data, kind, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    def sheet(name, fields, rows):
        ws = wb.create_sheet(name)
        ws.append([label for _,label in fields])
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True); cell.fill = PatternFill('solid', fgColor='387B6E')
        for row in rows:
            ws.append([row.get(key, '') for key,_ in fields])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # User-provided descriptions never become Excel formulas.
                if isinstance(cell.value, (float, int)):
                    cell.number_format = '#,##0.000'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if ws.max_row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F0F6F3')
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for i, (key,_) in enumerate(fields, 1):
            ws.column_dimensions[get_column_letter(i)].width = 42 if key in ('designation','client','name','label') else 22
    if kind == 'reports':
        summary = data['summary']
        sheet('Synthèse', [('label','Indicateur'),('value','Valeur')], [{'label': k, 'value': v} for k,v in {'Période du': data['period']['start'], 'Période au': data['period']['end'], 'Salle': data['period']['hall'], 'Situation des soldes au': data['asOf'], 'Bons confirmés': summary['count'], 'Forfaits (DT)': summary['booked'], 'Acomptes des bons (DT)': summary['paid'], 'Soldes actuels (DT)': summary['outstanding'], 'Acomptes datés de la période (DT)': summary['deposits'], 'Forfaits période précédente (DT)': summary['previousBooked']}.items()])
        sheet('Salles', [('name','Salle'),('count','Bons'),('booked','Forfaits (DT)'),('paid','Acomptes (DT)'),('outstanding','Reste (DT)'),('average','Forfait moyen (DT)')], data['halls'])
        sheet('Événements', [('name','Type'),('count','Bons'),('booked','Forfaits (DT)')], data['types'])
        sheet('Mensuel', [('month','Mois'),('count','Bons'),('booked','Forfaits (DT)'),('deposits','Acomptes datés (DT)')], data['monthly'])
        sheet('Impayés', [('bon','Bon'),('client','Client'),('hall','Salle'),('event','Événement'),('due','Échéance J-15'),('remaining','Reste (DT)'),('overdueDays','Jours de retard')], data['outstanding'])
        sheet('Caisses toutes salles', [('name','Caisse'),('income','Recettes (DT)'),('expense','Sorties (DT)'),('net','Solde des mouvements (DT)')], data['cash'])
        sheet('Employés', [('name','Employé'),('type','Type'),('days','Jours'),('hours','Heures supp.'),('earned','Tarifs actuels période (DT)'),('paid','Paiements période (DT)'),('globalRemaining','Reste global (DT)')], data['employees'])
        sheet('Modes de paiement', [('name','Mode'),('count','Acomptes'),('amount','Montant daté (DT)')], data['methods'])
        notes = ['Soldes des bons : situation actuelle, pas un solde historique.', 'Caisses : journaux existants, toutes salles, sans filtre de salle.', 'Employés : tarifs actuels ; le reste global couvre tout l’historique.'] + data['warnings']
    else:
        sheet('Hypothèses', [('label','Hypothèse'),('value','Valeur')], [{'label': k,'value': v} for k,v in {'Début de période': data['parameters']['start'], 'Fin de période': data['parameters']['end'], 'Solde initial (DT)': data['parameters']['opening'], 'Seuil (DT)': data['parameters']['floor'], 'Recouvrement soldes (%)': data['parameters']['rate'], 'Soldes échus inclus': str(data['parameters']['overdue']), 'Solde final prévu (DT)': data['summary']['closing']}.items()])
        sheet('Trésorerie prévue', [('start','Du'),('end','Au'),('opening','Ouverture (DT)'),('income','Entrées prévues (DT)'),('expense','Sorties prévues (DT)'),('net','Variation (DT)'),('closing','Solde prévu (DT)')], data['periods'])
        sheet('Mouvements prévus', [('date','Date projetée'),('due','Échéance d’origine'),('designation','Désignation'),('direction','Sens in/out'),('amount','Montant projeté (DT)'),('nominal','Montant nominal (DT)'),('source','Source'),('hall','Salle')], data['events'])
        notes = ['Prévision : soldes actuels des réservations confirmées à J-15 + mouvements planifiés.', 'Les dépenses futures sont à planifier. Les charges et acomptes déjà enregistrés ne sont pas projetés une seconde fois.', 'Les soldes échus inclus sont regroupés au début de période ; cette date est une hypothèse de recouvrement.', 'Marquer un plan réalisé ne crée aucun paiement ou mouvement dans les caisses.'] + data['warnings']
    sheet('Lecture du rapport', [('label','Informations')], [{'label': line} for line in notes])
    wb.save(path)
