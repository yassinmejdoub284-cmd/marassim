"""Omar cash statement: existing receipts/outgoings, with an automatic carry."""
import json
from datetime import date, datetime
from decimal import Decimal
from html import escape

from analytics import amount, number


def journal(conn, database, start, end):
    a, b = date.fromisoformat(start), date.fromisoformat(end)
    if a.isoformat() != start or b.isoformat() != end or a > b:
        raise ValueError('Période de caisse invalide.')
    # Read the same authoritative legacy cash movements once, including history.
    source = database.get_journal_caisse_omar(date.min.isoformat(), end)
    opening = Decimal(0)
    income, expense, warnings = [], [], []
    details = {}
    if conn.execute("SELECT 1 FROM sqlite_master WHERE name='omar_charge_details'").fetchone():
        for row in conn.execute('SELECT c.id,c.date_encaissement,c.designation,c.montant,d.details FROM charges_omar c JOIN omar_charge_details d ON d.charge_id=c.id WHERE c.date_encaissement BETWEEN ? AND ?', (start, end)):
            # A charge carries its creation-time detail, never today's employee rates.
            details[row['id']] = json.loads(row['details'])
    for entries, direction, target in ((source['recettes'], 1, income), (source['sorties'], -1, expense)):
        for row in entries:
            try:
                day = date.fromisoformat(row['date'])
                if day.isoformat() != row['date']:
                    raise ValueError()
                value = amount(row['montant'])
            except (ValueError, TypeError):
                warnings.append('Un mouvement avec une date ou un montant invalide a été exclu.')
                continue
            if row['date'] < start:
                opening += direction * value
            else:
                item = {**row, 'montant': number(value)}
                if item.get('facture'):
                    item['designation'] += ' · Facture ' + str(item['facture'])
                if direction == -1 and item.get('type') == 'charge':
                    item['details'] = details.get(item.get('charge_id'), [])
                    if item['details'] and sum((amount(d['amount']) for d in item['details']), Decimal(0)) != value:
                        warnings.append('Une charge de pointage a été modifiée ; son détail décrit la session à sa création.')
                target.append(item)
    total_income = sum((amount(r['montant']) for r in income), Decimal(0))
    total_expense = sum((amount(r['montant']) for r in expense), Decimal(0))
    income.sort(key=lambda r: (r['date'], r.get('res_id', 0), r.get('slot', 0), r.get('extra_id', 0)))
    expense.sort(key=lambda r: (r['date'], 0 if r.get('type')=='charge' else 1, r.get('charge_id', 0)))
    closing = opening + total_income - total_expense
    return {'recettes': income, 'sorties': expense, 'period': {'start': start, 'end': end},
            'opening': number(opening), 'income': number(total_income), 'expense': number(total_expense),
            'closing': number(closing), 'net': number(total_income-total_expense), 'warnings': warnings}


def fmt(day):
    return datetime.strptime(day, '%Y-%m-%d').strftime('%d/%m/%Y')


def statement_rows(data):
    yield [fmt(data['period']['start']), 'SOLDE DE DÉPART', '', '', data['opening']], 'balance'
    # The paper reference groups receipts, then chronological charges and details.
    for row in data['recettes']:
        yield [fmt(row['date']), row['designation'], row['montant'], '', ''], 'income'
    for row in data['sorties']:
        yield [fmt(row['date']), row['designation'], '', row['montant'], ''], 'expense'
        for n, detail in enumerate(row.get('details', []), 1):
            yield ['', f"{n} — {detail['name']} ({detail['period']} · {detail['amount']:.3f} DT)", '', '', ''], 'detail'
    yield ['', 'TOTAL DES MOUVEMENTS', data['income'], data['expense'], ''], 'total'
    yield [fmt(data['period']['end']), 'SOLDE DE FIN · REPORT AU MOIS SUIVANT', '', '', data['closing']], 'balance'


def export_excel(data, path):
    # Keep the existing application writer/runtime; only Omar's requested layout changes.
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.drawing.image import Image
    from app import LEGACY
    wb = Workbook(); ws = wb.active; ws.title = 'Caisse Omar'
    logo = LEGACY / 'desktop/assets/marassim-logo.png'
    if logo.exists():
        mark = Image(str(logo)); mark.width = 62; mark.height = 59; ws.add_image(mark, 'A1')
    ws.merge_cells('B1:D1'); ws['B1'] = 'MARASSIM · COMPLEXE DE RÉCEPTION'
    ws['E1'] = 'Caisse Omar · ' + data['period']['start'][:7]
    ws['E1'].font = Font(name='Times New Roman', size=10, italic=True)
    ws['E1'].alignment = Alignment(horizontal='right')
    ws.merge_cells('B2:E2'); ws['B2'] = ''
    ws.merge_cells('A3:E3'); ws['A3'] = f"ÉTAT DE CAISSE du {fmt(data['period']['start'])} JUSQU’AU {fmt(data['period']['end'])}"
    for r in (1, 2, 3):
        ws.cell(r, 2 if r < 3 else 1).font = Font(name='Times New Roman', size=12, bold=True)
        ws.cell(r, 2 if r < 3 else 1).alignment = Alignment(horizontal='center')
        ws.row_dimensions[r].height = 24
    ws.append(['Date', 'Désignation', 'Débit', 'Crédit', 'Solde'])
    line = Side(style='thin', color='000000')
    for values, kind in statement_rows(data):
        ws.append(values)
        row = ws.max_row
        ws.row_dimensions[row].height = max(19, 15 * (1 + len(str(values[1])) // 95))
        for cell in ws[row]:
            cell.font = Font(name='Times New Roman', size=10, bold=kind in ('total', 'balance'))
            cell.border = Border(left=line, right=line, top=line, bottom=line)
            cell.alignment = Alignment(vertical='top', horizontal='right' if cell.column >= 3 else 'left', wrap_text=True)
            if isinstance(cell.value, str): cell.data_type = 's'
            if cell.column >= 3: cell.number_format = '#,##0.000'
    for cell in ws[4]:
        cell.font = Font(name='Times New Roman', size=11, bold=True)
        cell.border = Border(left=line, right=line, top=line, bottom=line)
        cell.alignment = Alignment(horizontal='center')
    for key, width in zip('ABCDE', (15, 100, 17, 17, 19)):
        ws.column_dimensions[key].width = width
    ws.freeze_panes = 'C6'; ws.print_title_rows = '1:4'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=.3, right=.3, top=.35, bottom=.35)
    ws.print_area = ws.dimensions
    ws.oddFooter.center.text = 'Page &P / &N'; ws.oddHeader.right.text = 'Caisse Omar'
    for warning in data['warnings']:
        ws.append(['', warning])
        ws.cell(ws.max_row, 2).data_type = 's'
    wb.save(path)


def export_pdf(data, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from app import LEGACY
    style = ParagraphStyle('cash', fontName='Times-Roman', fontSize=10, leading=12)
    title = ParagraphStyle('title', parent=style, fontName='Times-Bold', fontSize=14, leading=18, alignment=1)
    body = []
    logo = LEGACY / 'desktop/assets/marassim-logo.png'
    heading = [Image(str(logo), width=38, height=36) if logo.exists() else '', Paragraph('MARASSIM · COMPLEXE DE RÉCEPTION', style), Paragraph('Caisse Omar · ' + data['period']['start'][:7], style)]
    header = Table([heading], colWidths=[45, 620, 115]); header.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
    body += [header, Spacer(1, 6), Paragraph(f"ÉTAT DE CAISSE du {fmt(data['period']['start'])} JUSQU’AU {fmt(data['period']['end'])}", title), Spacer(1, 10)]
    rows = [['Date', 'Désignation', 'Débit', 'Crédit', 'Solde']]
    emphasis = []
    for values, kind in statement_rows(data):
        display = [f'{value:,.3f}'.replace(',', ' ').replace('.', ',') if isinstance(value, (int, float)) else str(value) for value in values]
        rows.append([display[0], Paragraph(escape(display[1]), style), *display[2:]])
        if kind in ('total', 'balance'): emphasis.append(('FONTNAME', (0, len(rows)-1), (-1, len(rows)-1), 'Times-Bold'))
    table = Table(rows, colWidths=[65, 440, 85, 85, 105], repeatRows=1, hAlign='CENTER')
    table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .4, colors.black), ('FONTNAME', (0, 0), (-1, -1), 'Times-Roman'),
                              ('FONTNAME', (0, 0), (-1, 0), 'Times-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 10),
                              ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('ALIGN', (2, 1), (-1, -1), 'RIGHT'),
                              ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4), *emphasis]))
    body.append(table)
    for warning in data['warnings']: body.append(Paragraph(escape(warning), style))
    def footer(canvas, document):
        canvas.setFont('Times-Roman', 9)
        canvas.drawRightString(816, 17, f'Page {document.page}')
        canvas.drawString(26, 17, 'Caisse Omar · ' + data['period']['start'][:7])
    SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=26, rightMargin=26, topMargin=24, bottomMargin=29).build(body, onFirstPage=footer, onLaterPages=footer)
