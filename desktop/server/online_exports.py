"""Read-only module exports; each export uses the same permission as its screen."""
from app import APIError

FIELDS = {
    'charges': [('date_encaissement', 'Date'), ('designation', 'Désignation'), ('montant', 'Montant (DT)')],
    'charges-omar': [('date_encaissement', 'Date'), ('designation', 'Désignation'), ('montant', 'Montant (DT)')],
    'recettes-omar': [('date_recette', 'Date'), ('designation', 'Désignation'), ('montant', 'Montant (DT)')],
    'reservations': [('num_bon', 'Bon'), ('nom_client', 'Client'), ('salle', 'Salle'), ('date_evenement', 'Événement'), ('heure_debut', 'Début'), ('heure_fin', 'Fin'), ('forfait', 'Forfait (DT)'), ('acompte1', 'Acompte 1 (DT)'), ('acompte2', 'Acompte 2 (DT)'), ('acompte3', 'Acompte 3 (DT)'), ('tel1', 'Téléphone'), ('is_temporaire', 'Option')],
    'employees': [('nom', 'Nom'), ('prenom', 'Prénom'), ('role', 'Poste'), ('type_ouvrier', 'Type'), ('salaire_journalier', 'Tarif journalier (DT)'), ('salaire_midi', 'Midi (DT)'), ('salaire_apres_midi', 'Après-midi (DT)'), ('salaire_soiree', 'Soirée (DT)'), ('prix_heure_supp', 'Heure supplémentaire (DT)'), ('actif', 'Actif')],
    'pointage': [('employee', 'Employé'), ('date_pointage', 'Date'), ('periode', 'Période'), ('heure_arrivee', 'Arrivée'), ('heure_depart', 'Départ'), ('heures_supplementaires', 'Heures supplémentaires'), ('notes', 'Notes')],
    'employee-payments': [('employee', 'Employé'), ('date_paiement', 'Date'), ('periode_debut', 'Période du'), ('periode_fin', 'Période au'), ('montant_total', 'Total payé (DT)'), ('notes', 'Notes')],
    'rules': [('nom', 'Règle'), ('salle', 'Salle'), ('date_debut', 'Du'), ('date_fin', 'Au'), ('creneaux', 'Créneaux'), ('actif', 'Active'), ('notes', 'Notes')],
    'reception': [('nom_client', 'Client'), ('num_bon', 'Bon'), ('salle', 'Salle'), ('date_evenement', 'Événement'), ('date_acompte', 'Date acompte'), ('montant_omar', 'Montant Omar (DT)'), ('statut', 'État')],
    'users': [('username', 'Identifiant'), ('nom', 'Nom'), ('prenom', 'Prénom'), ('role', 'Rôle'), ('actif', 'Actif'), ('modules', 'Modules autorisés')],
}


def export_table(application, conn, user, resource, query, path):
    if resource not in FIELDS:
        raise APIError('Module d’export inconnu.', 404)
    rows = application.business_route(conn, user, 'GET', '/' + resource, {}, {})
    selected = query.get('ids')
    if selected is not None:
        ids = {int(value) for value in selected.split(',') if value}
        rows = [row for row in rows if row.get('id', row.get('res_id')) in ids]
    if resource=='reception' and 'slots' in query:
        slots={tuple(map(int,value.split(':'))) for value in query['slots'].split(',') if value}
        rows=[row for row in rows if (row['res_id'],row['slot']) in slots]
    employees = {}
    if resource in ('pointage', 'employee-payments'):
        employees = {r['id']: f"{r['nom']} {r['prenom'] or ''}".strip() for r in conn.execute('SELECT id,nom,prenom FROM employees')}
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = 'Consultation'
    fields = FIELDS[resource]
    ws.append([label for _, label in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='387B6E')
    for row in rows:
        values = {**row, 'employee': employees.get(row.get('employee_id'), f"Employé #{row.get('employee_id')}")}
        ws.append([' / '.join(map(str, values.get(key, []))) if isinstance(values.get(key), list) else values.get(key, '') for key, _ in fields])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, (int, float)): cell.number_format = '#,##0.000'
    for n, (key, _) in enumerate(fields, 1):
        ws.column_dimensions[get_column_letter(n)].width = 50 if key in ('designation', 'notes', 'modules') else 24
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = '1:1'
    wb.save(path)
