"""Excel layouts copied unchanged from the existing Tkinter application.
Regenerate with scripts/extract-legacy-exports.py after an intentional legacy change.
"""
from datetime import datetime
from types import SimpleNamespace
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def _write_tawfik(self, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Journal de Caisse"
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = self._last_data["header"]
    c.font = Font(bold=True, size=12, color="1A2B45")
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 24
    hdrs = ["Date","Désignation","FN°","N° Caisse","Recette (DT)","Charge (DT)"]
    fill_hdr = PatternFill("solid", fgColor="1A2B45")
    for ci, h in enumerate(hdrs, 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    ri = 3
    for row in self._last_data["recettes"]:
        fn_d = f"FN°{row['fn']}" if row["fn"] and not row["fn"].startswith("FN") else row["fn"]
        for ci, v in enumerate([self._fmt_date(row["date"]), row["designation"],
                                fn_d or "", row["num_caisse"] or "", row["montant"], ""], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill = PatternFill("solid", fgColor="ECFDF5")
            cell.border = border
            if ci >= 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
        ri += 1
    for row in self._last_data["charges"]:
        for ci, v in enumerate([row["date_fmt"], row["designation"], "", "", "", row["montant"]], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill = PatternFill("solid", fgColor="FFF1F2")
            cell.border = border
            if ci >= 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
        ri += 1
    tr = sum(r["montant"] for r in self._last_data["recettes"])
    tc = sum(r["montant"] for r in self._last_data["charges"])
    for ci, v in enumerate(["","","","TOTAUX", tr, tc], 1):
        cell = ws.cell(row=ri, column=ci, value=v)
        cell.fill = PatternFill("solid", fgColor="EEF2FF")
        cell.font = Font(bold=True)
        cell.border = border
        if ci >= 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
    ri += 1
    for ci, v in enumerate(["","","","SOLDE NET", tr-tc, ""], 1):
        cell = ws.cell(row=ri, column=ci, value=v)
        cell.fill = PatternFill("solid", fgColor="FFFBEB")
        cell.font = Font(bold=True, size=11)
        cell.border = border
        if ci == 5: cell.number_format = '#,##0.000'; cell.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 16
    ws.freeze_panes = "A3"
    wb.save(path)

def _write_omar(self, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Journal Caisse Omar"

    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title row
    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = self._last_data["header"]
    c.font  = Font(bold=True, size=12, color="1A2B45")
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 24

    # Header row
    hdrs = ["Date", "Designation", "Recette (DT)", "Sortie (DT)"]
    fill_hdr = PatternFill("solid", fgColor="1A2B45")
    for ci, h in enumerate(hdrs, 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = fill_hdr
        cell.alignment = Alignment(horizontal="center")
        cell.border    = border

    ri = 3
    # Recettes
    for row in self._last_data["recettes"]:
        for ci, v in enumerate([
            self._fmt_date(row["date"]),
            row["designation"],
            row["montant"],
            "",
        ], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill   = PatternFill("solid", fgColor="ECFDF5")
            cell.border = border
            if ci >= 3:
                cell.number_format = "#,##0.000"
                cell.alignment = Alignment(horizontal="right")
        ri += 1

    # Sorties
    for row in self._last_data["sorties"]:
        fg_color = "FFF1F2" if row["type"] == "charge" else "FFF0F9"
        for ci, v in enumerate([
            row["date_fmt"],
            row["designation"],
            "",
            row["montant"],
        ], 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill   = PatternFill("solid", fgColor=fg_color)
            cell.border = border
            if ci >= 3:
                cell.number_format = "#,##0.000"
                cell.alignment = Alignment(horizontal="right")
        ri += 1

    # Totals
    tr = sum(r["montant"] for r in self._last_data["recettes"])
    ts = sum(r["montant"] for r in self._last_data["sorties"])
    for ci, v in enumerate(["", "TOTAUX", tr, ts], 1):
        cell = ws.cell(row=ri, column=ci, value=v)
        cell.fill   = PatternFill("solid", fgColor="EEF2FF")
        cell.font   = Font(bold=True)
        cell.border = border
        if ci >= 3:
            cell.number_format = "#,##0.000"
            cell.alignment = Alignment(horizontal="right")
    ri += 1
    for ci, v in enumerate(["", "SOLDE NET", tr - ts, ""], 1):
        cell = ws.cell(row=ri, column=ci, value=v)
        cell.fill   = PatternFill("solid", fgColor="FFFBEB")
        cell.font   = Font(bold=True, size=11)
        cell.border = border
        if ci == 3:
            cell.number_format = "#,##0.000"
            cell.alignment = Alignment(horizontal="right")

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.freeze_panes = "A3"

    wb.save(path)

def export_journal(database, omar, start, end, path, number='?'):
    def fmt(value):
        return datetime.strptime(value, '%Y-%m-%d').strftime('%d/%m/%Y') if value else ''
    days = ['LUNDI', 'MARDI', 'MERCREDI', 'JEUDI', 'VENDREDI', 'SAMEDI', 'DIMANCHE']
    day = lambda value: days[datetime.strptime(value, '%Y-%m-%d').weekday()]
    data = (database.get_journal_caisse_omar if omar else database.get_journal_caisse)(start, end)
    header = f"Journal de Caisse{' Omar' if omar else ''} N°{number}  —  {day(start)} {fmt(start)}  au  {day(end)} {fmt(end)}"
    data.update(header=header, debut=start, fin=end)
    stub = SimpleNamespace(_last_data=data, _fmt_date=fmt)
    (_write_omar if omar else _write_tawfik)(stub, path)
