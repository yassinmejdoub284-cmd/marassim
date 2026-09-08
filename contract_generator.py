"""
contract_generator.py — Génère le Bon de Réservation (.docx) à partir du
template `template/Bon_Recu_Marassim_Template.docx` en remplaçant les
placeholders {{XXX}} par les données de la réservation.
"""
import os
import re
import docx

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              "template", "Bon_Recu_Marassim_Template.docx")


def _fmt_date(d):
    if not d:
        return ""
    try:
        from datetime import datetime
        return datetime.strptime(d, "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        return d


def _fmt_montant(v):
    if v in (None, "", 0):
        return ""
    try:
        return f"{float(v):.3f}"
    except Exception:
        return str(v)


def _fmt_acompte_line(montant, date_contrat_iso, methode, num_ref,
                      banque=None, date_cheque_iso=None):
    """Formate une ligne d'acompte selon la méthode :

    Espèce  → '4000.000 DT ESP LE 04/08/2026 NON RECUPERABLE EN CAS DE DESISTEMENT'
    Chèque  → '4000.000 DT CHQ (005465, BIAT, 04/08/2026) RECU LE 04/08/2026 NON RECUPERABLE EN CAS DE DESISTEMENT'
    Virement→ même format que chèque avec VIR
    Effet   → même format que chèque avec EFF
    """
    if not montant:
        return ""
    m = _fmt_montant(montant)
    d_contrat = _fmt_date(date_contrat_iso)
    methode_up = (methode or "Espèce").strip().upper()

    if methode_up.startswith("ESP") or methode_up == "":
        return (f"{m} DT ESP LE {d_contrat} "
                f"NON RECUPERABLE EN CAS DE DESISTEMENT")

    # Chèque / Virement / Effet — détails entre parenthèses
    if methode_up.startswith("CHQ") or methode_up.startswith("CHÈ") or methode_up.startswith("CHE"):
        label = "CHQ"
    elif methode_up.startswith("VIR"):
        label = "VIR"
    elif methode_up.startswith("EFF"):
        label = "EFF"
    else:
        label = methode_up[:3]

    parts = []
    if num_ref:
        parts.append(num_ref)
    if banque:
        parts.append(banque)
    if date_cheque_iso:
        parts.append(_fmt_date(date_cheque_iso))
    details = f" ({', '.join(parts)})" if parts else ""
    return (f"{m} DT {label}{details} RECU LE {d_contrat} "
            f"NON RECUPERABLE EN CAS DE DESISTEMENT")


def build_placeholder_map(res: dict) -> dict:
    """Construit le dictionnaire {{PLACEHOLDER}} -> valeur pour une réservation."""
    from datetime import date as _date
    res = {k: ("" if v is None else v) for k, v in res.items()}
    option_violoniste = ""
    if res.get("with_violoniste"):
        option_violoniste = (f"+ 1h gratuite Violoniste (Espace Rayhane) "
                              f"de {res.get('violoniste_heure_debut','')} à "
                              f"{res.get('violoniste_heure_fin','')}")

    solde_acomptes = sum(float(res.get(k) or 0) for k in ("acompte1", "acompte2", "acompte3"))
    forfait = float(res.get("forfait") or 0)
    reste = forfait - solde_acomptes

    # Utiliser le statut stocké (ex: "SOLDE LE 04/08/2026") sinon calculer
    stored_statut = str(res.get("statut_paiement") or "").strip()
    if stored_statut:
        statut_paiement = stored_statut
    elif reste <= 0.001:
        statut_paiement = f"SOLDE LE {_date.today().strftime('%d/%m/%Y')}"
    else:
        statut_paiement = f"Reste à payer : {reste:.3f} DT"

    # DATE_SIGNATURE = date_signature stockée ou aujourd'hui
    date_sig = res.get("date_signature") or _date.today().isoformat()

    return {
        "{{NUM_BON}}": str(res.get("num_bon") or res.get("id") or ""),
        "{{DATE_EVENEMENT}}": _fmt_date(res.get("date_evenement")),
        "{{SALLE}}": res.get("salle", ""),
        "{{NB_PLACES}}": str(res.get("nb_places") or ""),
        "{{NOM_CLIENT}}": res.get("nom_client", ""),
        "{{CIN}}": res.get("cin", ""),
        "{{DATE_CIN}}": _fmt_date(res.get("date_cin")),
        "{{TEL_1}}": res.get("tel1", ""),
        "{{TEL_2}}": res.get("tel2", ""),
        "{{LIEU_CONTRAT}}": res.get("lieu_contrat", ""),
        "{{DATE_CONTRAT}}": _fmt_date(res.get("date_contrat")),
        "{{HEURE_CONTRAT}}": res.get("heure_contrat", ""),
        "{{OPTION_VIOLONISTE}}": option_violoniste,
        "{{DOSSIER_TRAITE_PAR}}": res.get("dossier_traite_par", ""),
        "{{TYPE_EVENEMENT}}": res.get("type_evenement", ""),
        "{{HEURE_DEBUT}}": res.get("heure_debut", ""),
        "{{HEURE_FIN}}": res.get("heure_fin", ""),
        "{{FORFAIT}}": _fmt_montant(res.get("forfait")),
        "{{STATUT_PAIEMENT}}": statut_paiement,
        "{{ACOMPTE_1}}": _fmt_acompte_line(
            res.get("acompte1"),
            res.get("date_contrat") or res.get("date_acompte1"),
            res.get("methode_acompte1") or res.get("methode_paiement"),
            res.get("fn1"),
            banque=res.get("banque"),
            date_cheque_iso=res.get("date_cheque_virement")),
        "{{DATE_ACOMPTE_1}}": "",
        "{{FN_1}}": "",
        "{{ACOMPTE_2}}": _fmt_acompte_line(
            res.get("acompte2"),
            res.get("date_contrat") or res.get("date_acompte2"),
            res.get("methode_acompte2"),
            res.get("fn2"),
            banque=res.get("banque"),
            date_cheque_iso=res.get("date_cheque_virement")),
        "{{DATE_ACOMPTE_2}}": "",
        "{{FN_2}}": "",
        "{{ACOMPTE_3}}": _fmt_acompte_line(
            res.get("acompte3"),
            res.get("date_contrat") or res.get("date_acompte3"),
            res.get("methode_acompte3"),
            res.get("fn3"),
            banque=res.get("banque"),
            date_cheque_iso=res.get("date_cheque_virement")),
        "{{DATE_ACOMPTE_3}}": "",
        "{{FN_3}}": "",
        "{{DATE_SIGNATURE}}": _fmt_date(date_sig),
    }


_EMPTY_ACOMPTE_RE = re.compile(
    r"^\s*[•●\-]?\s*DT\s+(ESP|CHQ|VIR|EFF)\s+LE\s+NON\s+RECUPERABLE", re.IGNORECASE
)
# Static tail the template appends after {{ACOMPTE_N}} — strip it once the
# placeholder has already been expanded to the full formatted string.
_ACOMPTE_TAIL_RE = re.compile(
    r"\s*DT\s+ESP\s+LE\s+NON\s+RECUPERABLE\s+EN\s+CAS\s+DE\s+DESISTEMENT\s*\.\s*\[FN:\s*\]",
    re.IGNORECASE
)

def _replace_in_paragraph(paragraph, mapping):
    """Remplace les placeholders dans un paragraphe, même s'ils sont
    fragmentés sur plusieurs `runs` (comportement classique de Word)."""
    full_text = "".join(run.text for run in paragraph.runs)
    if not full_text or "{{" not in full_text:
        return
    new_text = full_text
    for key, val in mapping.items():
        new_text = new_text.replace(key, str(val))
    if new_text == full_text:
        return

    # Strip the static template tail that follows {{ACOMPTE_N}} when the
    # placeholder has already been expanded to a complete formatted string.
    # e.g. "2500.000 DT CHQ ... NON RECUPERABLE  DT ESP LE  NON RECUPERABLE. [FN: ]"
    if "NON RECUPERABLE EN CAS DE DESISTEMENT" in new_text.upper():
        new_text = _ACOMPTE_TAIL_RE.sub("", new_text)

    # Effacer les lignes d'acompte vides : bullet dont ACOMPTE_N est vide
    # → après remplacement on obtient "• DT ESP LE  NON RECUPERABLE..." sans montant
    if _EMPTY_ACOMPTE_RE.search(new_text):
        new_text = ""
    # Aussi effacer les lignes entièrement vides
    elif re.search(r"^\s*[•●\-]?\s*$", new_text.strip()):
        new_text = ""

    if not paragraph.runs:
        return
    # Réécrit tout dans le premier run, vide les suivants
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def generate_contract(res: dict, output_path: str):
    """Génère le fichier .docx rempli à `output_path`."""
    mapping = build_placeholder_map(res)
    document = docx.Document(TEMPLATE_PATH)

    for p in document.paragraphs:
        _replace_in_paragraph(p, mapping)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, mapping)

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    document.save(output_path)
    return output_path


# ═══════════════════════════════════════════════════════════════════════════
#  CONTRAT ARABE  —  عقد تسويغ قاعة فراح
# ═══════════════════════════════════════════════════════════════════════════

ARABIC_TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "template", "Contrat_Arabe_Template.docx"
)

# Mapping salle français → arabe
SALLE_AR = {
    "Almes":           "ألماس",
    "Chichkhane":      "شيشخان",
    "Rayhane":         "ريحان",
    "Toute la complexe": "كامل الفضاء",
}

# Mapping jour de semaine (weekday()) → arabe
JOURS_AR = ["الإثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]

# Heures en toutes lettres arabes (pour affichage heure événement)
HEURES_AR = {
    "00:00": "منتصف الليل", "01:00": "الواحدة صباحا", "02:00": "الثانية صباحا",
    "03:00": "الثالثة صباحا", "04:00": "الرابعة صباحا", "05:00": "الخامسة صباحا",
    "06:00": "السادسة صباحا", "07:00": "السابعة صباحا", "08:00": "الثامنة صباحا",
    "09:00": "التاسعة صباحا", "10:00": "العاشرة صباحا", "11:00": "الحادية عشرة صباحا",
    "12:00": "الظهر", "13:00": "الواحدة مساءا", "14:00": "الثانية مساءا",
    "15:00": "الثالثة مساءا", "16:00": "الرابعة مساءا", "17:00": "الخامسة مساءا",
    "18:00": "السادسة مساءا", "19:00": "السابعة مساءا", "20:00": "الثامنة مساءا",
    "21:00": "التاسعة ليلا", "22:00": "العاشرة ليلا", "23:00": "الحادية عشرة ليلا",
}


def _heure_ar(hhmm: str) -> str:
    return HEURES_AR.get(hhmm, hhmm)


def _fmt_date_ar(d: str) -> str:
    """YYYY-MM-DD  →  DD/MM/YYYY"""
    if not d:
        return ""
    try:
        from datetime import datetime
        return datetime.strptime(d, "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        return d


def _jour_ar(d: str) -> str:
    """YYYY-MM-DD → nom du jour en arabe"""
    try:
        from datetime import datetime
        return JOURS_AR[datetime.strptime(d, "%Y-%m-%d").weekday()]
    except Exception:
        return ""


def build_arabic_placeholder_map(res: dict, ar_info: dict) -> dict:
    """
    Construit le dict {{PLACEHOLDER}} → valeur pour le contrat arabe.

    ar_info doit contenir :
      - ar_nom       : الاسم واللقب  (peut être différent de nom_client)
      - ar_lieu_cin  : مكان بطاقة التعريف
      - ar_salle     : nom salle en arabe (pré-rempli, éditable)
      - ar_date      : date de génération JJ/MM/AAAA
      - num_contrat  : numéro du contrat (ex: 677/2026)
    """
    from datetime import date as _date
    res = {k: ("" if v is None else v) for k, v in res.items()}

    salle_fr  = res.get("salle", "")
    salle_ar  = ar_info.get("ar_salle") or SALLE_AR.get(salle_fr, salle_fr)
    date_evt  = res.get("date_evenement", "")
    hd        = res.get("heure_debut", "")
    hf        = res.get("heure_fin", "")

    forfait   = float(res.get("forfait") or 0)
    acomptes  = sum(float(res.get(f"acompte{i}") or 0) for i in range(1, 4))
    reste     = forfait - acomptes

    date_gen  = ar_info.get("ar_date") or _date.today().strftime("%d/%m/%Y")
    # Convertir date_gen JJ/MM/AAAA → ISO pour _jour_ar
    try:
        d_parts = date_gen.split("/")
        date_gen_iso = f"{d_parts[2]}-{d_parts[1]}-{d_parts[0]}"
    except Exception:
        date_gen_iso = _date.today().isoformat()

    nb_places = str(res.get("nb_places") or "")
    type_evt  = res.get("type_evenement", "") or "حفل"

    return {
        "{{NUM_CONTRAT}}":   ar_info.get("num_contrat", ""),
        "{{NOM_AR}}":        ar_info.get("ar_nom", "") or res.get("nom_client", ""),
        "{{CIN}}":           res.get("cin", ""),
        "{{DATE_CIN}}":      _fmt_date_ar(res.get("date_cin")),
        "{{LIEU_CIN}}":      ar_info.get("ar_lieu_cin", "") or res.get("lieu_cin", ""),
        "{{ADRESSE}}":       res.get("lieu_contrat", ""),
        "{{SALLE_AR}}":      salle_ar,
        "{{JOUR_AR}}":       _jour_ar(date_evt),
        "{{DATE_EVENEMENT}}": _fmt_date_ar(date_evt),
        "{{HEURE_DEBUT_AR}}": _heure_ar(hd),
        "{{HEURE_FIN_AR}}":  _heure_ar(hf),
        "{{TYPE_EVENEMENT}}": type_evt,
        "{{FORFAIT}}":       f"{forfait:,.3f}" if forfait else "",
        "{{ACOMPTE_PAYE}}":  f"{acomptes:,.3f}" if acomptes else "",
        "{{RESTE}}":         f"{reste:,.3f}"   if forfait else "",
        "{{NB_PLACES}}":     nb_places,
        "{{DATE_SIGNATURE}}": date_gen,
        "{{TEL_1}}":         res.get("tel1", ""),
        "{{TEL_2}}":         res.get("tel2", ""),
    }


def build_arabic_template():
    """
    Crée le fichier Contrat_Arabe_Template.docx dans le dossier template/.
    Appelé une seule fois (ou si le fichier est absent).
    Le document est en arabe, aligné à droite, police Amiri / Traditional Arabic.
    """
    import docx as _docx
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import copy

    doc = _docx.Document()

    # ── Mise en page : A4, marges réduites ──────────────────────────────
    section = doc.sections[0]
    section.page_width  = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin  = section.bottom_margin = Cm(1.8)

    AR_FONT   = "Traditional Arabic"
    FR_FONT   = "Times New Roman"
    NAVY      = RGBColor(0x1A, 0x3C, 0x5E)
    BLACK     = RGBColor(0, 0, 0)

    def _rtl_para(para):
        """Force RTL + bidi sur un paragraphe."""
        pPr = para._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "right")
        pPr.append(jc)

    def add_para(text="", size=13, bold=False, color=BLACK,
                 align=WD_ALIGN_PARAGRAPH.RIGHT, font=AR_FONT, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        _rtl_para(p)
        p.paragraph_format.space_after  = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(text)
        run.font.name  = font
        run.font.size  = Pt(size)
        run.font.bold  = bold
        run.font.color.rgb = color
        # Force Arabic font in rPr/rFonts
        rFonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        rFonts.set(qn("w:cs"),    font)
        rFonts.set(qn("w:ascii"), font)
        rFonts.set(qn("w:hAnsi"), font)
        return p

    def add_line(parts, size=12, space_after=4):
        """
        parts = list of (text, bold, color).
        Tous dans le même paragraphe RTL.
        """
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _rtl_para(p)
        p.paragraph_format.space_after  = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        for text, bold, color in parts:
            run = p.add_run(text)
            run.font.name  = AR_FONT
            run.font.size  = Pt(size)
            run.font.bold  = bold
            run.font.color.rgb = color
            rFonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            rFonts.set(qn("w:cs"),    AR_FONT)
            rFonts.set(qn("w:ascii"), AR_FONT)
            rFonts.set(qn("w:hAnsi"), AR_FONT)
        return p

    # ══════════════ EN-TÊTE ══════════════
    add_para("عقد تسويغ قاعة فراح", size=18, bold=True, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para("─" * 55, size=9, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    # Numéro contrat
    add_line([("عـــــدد  ", False, BLACK), ("{{NUM_CONTRAT}}", True, NAVY)],
             size=13, space_after=10)

    # ── Ligne séparatrice
    add_para("اتفق الطرفان :", size=13, bold=True, color=NAVY, space_after=6)

    # ── Partie 1
    add_para("الطرف الأوّل : « المسوّغ »", size=12, bold=True, space_after=2)
    add_para("مركب مراسيم — طريق تنيور كلم 9.5 الشيحية صفاقس 3041",
             size=11, space_after=8)

    # ── Partie 2
    add_para("الطرف الثّاني : « المتسوّغ »", size=12, bold=True, space_after=4)

    add_line([
        ("السيّد/ة  ", False, BLACK),
        ("{{NOM_AR}}", True, NAVY),
        ("   صاحب/ة بطاقة التعريف الوطنية عدد  ", False, BLACK),
        ("{{CIN}}", True, NAVY),
        ("  صادرة في  ", False, BLACK),
        ("{{DATE_CIN}}", True, NAVY),
        ("  بـ  ", False, BLACK),
        ("{{LIEU_CIN}}", True, NAVY),
    ], size=12, space_after=4)

    add_line([
        ("العنوان :  ", False, BLACK),
        ("{{ADRESSE}}", True, NAVY),
        ("    الهاتف :  ", False, BLACK),
        ("{{TEL_1}}", True, NAVY),
    ], size=12, space_after=10)

    # ── Objet du contrat
    add_para("على تسويغ المركب :", size=13, bold=True, color=NAVY, space_after=4)

    add_line([
        ("فضاء  ", False, BLACK),
        ("{{SALLE_AR}}", True, NAVY),
        ("   يوم  ", False, BLACK),
        ("{{JOUR_AR}}", True, NAVY),
        ("  الموافق  ", False, BLACK),
        ("{{DATE_EVENEMENT}}", True, NAVY),
    ], size=12, space_after=4)

    add_line([
        ("من الساعة  ", False, BLACK),
        ("{{HEURE_DEBUT_AR}}", True, NAVY),
        ("  إلى الساعة  ", False, BLACK),
        ("{{HEURE_FIN_AR}}", True, NAVY),
        ("  لإقامة  ", False, BLACK),
        ("{{TYPE_EVENEMENT}}", True, NAVY),
        ("  .  ", False, BLACK),
    ], size=12, space_after=4)

    add_line([
        ("بمبلغ جملي قدره  ", False, BLACK),
        ("{{FORFAIT}}", True, NAVY),
        ("  دينار — دُفع منه عند تحرير هذا العقد  ", False, BLACK),
        ("{{ACOMPTE_PAYE}}", True, NAVY),
        ("  دينار — والباقي  ", False, BLACK),
        ("{{RESTE}}", True, NAVY),
        ("  دينار.  ", False, BLACK),
    ], size=12, space_after=4)

    add_line([
        ("عدد المقاعد المخصّصة :  ", False, BLACK),
        ("{{NB_PLACES}}", True, NAVY),
        ("  مقعدًا.  ", False, BLACK),
    ], size=12, space_after=10)

    # ── Séparateur
    add_para("تمّ الاتفاق والإمضاء بين الطرفين على الشروط التالية :", size=13,
             bold=True, color=NAVY, space_after=6)

    # ── Clauses
    clauses = [
        "1) إحضار رخصة إقامة الحفل من بلديّة الشّيحية قبل 48 ساعة من الموعد المذكور.",
        "2) يلتزم المتسوّغ بمنع تداول الخمور وشربه طيلة الحفل داخل القاعة وفي محيطها.",
        "3) يمنع استعمال الشموع الزيتية على الطاولات ما عدا الشموع الكهربائية.",
        "4) يتعهّد المتسوّغ بإرجاع القاعة والتجهيزات والمعدّات في حالتها الأصلية مع النظافة ورفع الفضلات قبل المغادرة.",
        "5) يمنع التدخين داخل القاعة ويمنع التعليق أو المسّ من الجدران واللافتات الحائطية أو إدخال أي تغيير على معدّات القاعة.",
        "6) الممضي أسفل هذا العقد وصاحب الحفل يضمنان كلّ تلف أو تعييب بالتجهيزات الموجودة.",
        "7) يتحمّل المتسوّغ جميع المسؤوليات القانونية عن أي طارئ أو عيب أو تلف أثناء الحفل فيما يخصّ القاعة ومعدّاتها والمدعوين والعمّال والأطعمة.",
        "8) توفّر إدارة القاعة الأعوان المسؤولين على الحراسة الداخلية والخارجية وأعوان الخدمات.",
        "9) في حالة إلغاء الحفل لظرف خارق للعادة يمكن الانتفاع لمدّة لا تفوق 6 أشهر مع الالتزام بتغيير الاتفاق كتابيًا قبل 3 أشهر.",
        "10) كلّ مخالفة لأي بند من هذا العقد ينجرّ عنها المتابعات العدلية لدى المحاكم المختصة بصفاقس.",
        "11) يصادق الطرفان على جميع البنود ويتعهّدان بالالتزام بكلّ ما ورد فيه، وكلّ تغيير يجب أن يكون مصادقًا عليه كتابيًا من الطرفين.",
        "12) حُرِّر هذا العقد في نظيرين أصليّين بتاريخ الإمضاء.",
    ]
    for cl in clauses:
        add_para(cl, size=11, space_after=3)

    # ── Signatures
    add_para(" ", size=8, space_after=4)
    sig_tbl = doc.add_table(rows=1, cols=2)
    sig_tbl.style = "Table Grid"
    # Masquer les bordures
    for cell in sig_tbl.rows[0].cells:
        for border in ("top", "left", "bottom", "right"):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = OxmlElement("w:tcBorders")
            bd = OxmlElement(f"w:{border}")
            bd.set(qn("w:val"), "none")
            tcBorders.append(bd)
            tcPr.append(tcBorders)

    left_cell  = sig_tbl.rows[0].cells[1]   # arabe = droite → col 1
    right_cell = sig_tbl.rows[0].cells[0]   # entreprise → col 0

    def _sig_cell(cell, title, sub):
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p1.add_run(title)
        r.font.name = AR_FONT; r.font.size = Pt(12); r.font.bold = True
        rFonts = r._r.get_or_add_rPr().get_or_add_rFonts()
        rFonts.set(qn("w:cs"), AR_FONT)
        if sub:  # Only add subtitle if not empty
            p2 = cell.add_paragraph(sub)
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.runs[0]
            r2.font.name = FR_FONT; r2.font.size = Pt(9)
        cell.add_paragraph(" ").paragraph_format.space_after = Pt(30)

    _sig_cell(left_cell,  "المتسوّغ", "")
    _sig_cell(right_cell, "مركب مراسيم", "Complexe Marassim — Rte Teniour km 9,5 Chihia Sfax")

    # ── Date de signature
    add_para(" ", size=6, space_after=2)
    add_line([
        ("حُرِّر بتاريخ  ", False, BLACK),
        ("{{DATE_SIGNATURE}}", True, NAVY),
    ], size=12, space_after=4)

    # ── Pied de page entreprise
    add_para("─" * 55, size=9, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(
        "Complexe Marassim  |  Rte Teniour km 9,5 Chihia Sfax BP 3041\n"
        "M : +216 98 764 444   T : +216 74 848 333   F : +216 74 848 444   E : contact@marassim.tn",
        size=9, font=FR_FONT, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, space_after=0)

    os.makedirs(os.path.dirname(ARABIC_TEMPLATE_PATH), exist_ok=True)
    doc.save(ARABIC_TEMPLATE_PATH)
    return ARABIC_TEMPLATE_PATH


def generate_arabic_contract(res: dict, ar_info: dict, output_path: str) -> str:
    """
    Génère le contrat arabe .docx à output_path.
    Crée le template s'il n'existe pas encore.
    """
    if not os.path.exists(ARABIC_TEMPLATE_PATH):
        build_arabic_template()

    mapping  = build_arabic_placeholder_map(res, ar_info)
    document = docx.Document(ARABIC_TEMPLATE_PATH)

    for p in document.paragraphs:
        _replace_in_paragraph(p, mapping)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, mapping)

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    document.save(output_path)
    return output_path
